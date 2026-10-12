"""Offline tests for reviews.py. No network, no keys.

Run (from the skill folder):
  python3 -m unittest scripts/test_reviews.py
  uv run --with wordcloud --with matplotlib --with openpyxl python -m unittest scripts/test_reviews.py
The second form also runs the full offline pipeline (word clouds and workbook).
"""

import importlib.util
import json
import shutil
import sys
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).parent))
import reviews as rv  # noqa: E402

SKILL = Path(__file__).resolve().parent.parent
FIX = SKILL / "tests" / "fixtures"
HAS_WC = importlib.util.find_spec("wordcloud") is not None and importlib.util.find_spec("matplotlib") is not None
HAS_XL = importlib.util.find_spec("openpyxl") is not None


def raw(platform):
    return json.loads((FIX / "raw" / f"{platform}.json").read_text(encoding="utf-8"))


def norm(platform, salt="s"):
    return [r for r in (rv.normalise(platform, it, "https://profile.example/", salt) for it in raw(platform)) if r]


class Normalise(unittest.TestCase):
    def test_every_record_has_the_common_schema(self):
        for platform in ("trustpilot", "google_maps", "productreview", "g2", "capterra", "reddit"):
            for rec in norm(platform):
                self.assertEqual(tuple(rec), rv.SCHEMA_FIELDS, platform)
                self.assertEqual(rec["platform"], platform)

    def test_trustpilot(self):
        a, b = norm("trustpilot")
        self.assertEqual((a["rating"], a["date"], a["reply_date"]), (4.0, "2026-09-30", "2026-10-02"))
        self.assertEqual(a["owner_reply"], "Thanks, glad we could help.")
        self.assertEqual(a["review_url"], "https://www.trustpilot.com/reviews/tp0001")
        self.assertEqual((b["rating"], b["owner_reply"], b["date"]), (1.0, "", "2026-08-01"))

    def test_google_maps_keeps_star_only_and_drops_ads(self):
        rows = norm("google_maps")
        self.assertEqual(len(rows), 2)  # the advertisement is gone
        self.assertEqual(rows[0]["owner_reply"], "Thank you!")
        self.assertEqual((rows[1]["text"], rows[1]["rating"]), ("", 4.0))  # star-only review kept

    def test_productreview_brand_reply_and_listing_row_dropped(self):
        rows = norm("productreview")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["owner_reply"], "Sorry to hear this, please contact us.")
        self.assertEqual(rows[0]["profile_url"], "https://www.productreview.com.au/listings/acme-widget")

    def test_g2_uses_likes_and_dislikes_and_skips_summary(self):
        rows = norm("g2")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["text"], "Likes: Fast setup. Dislikes: Pricey add-ons.")
        self.assertEqual((rows[0]["date"], rows[0]["rating"]), ("2026-10-12", 4.0))

    def test_capterra(self):
        (row,) = norm("capterra")
        self.assertEqual(row["date"], "2026-10-10")
        self.assertIn("Cons: Search is slow.", row["text"])
        self.assertEqual(row["rating"], 4.0)

    def test_reddit_has_no_stars_and_skips_communities(self):
        rows = norm("reddit")
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(r["rating"] is None for r in rows))
        brand = rv.brand_pattern("Acme Widgets")
        kept = [r for r in rows if rv.keep_review(r, "reddit", brand)]
        self.assertEqual(len(kept), 2)  # the off-topic mattress comment is dropped

    def test_ids_are_unique_when_review_urls_are_shared(self):
        items = raw("productreview")[:1] * 1 + [dict(raw("productreview")[0], id="pr-2", body="Different text")]
        rows = [rv.normalise("productreview", it, "x") for it in items]
        self.assertEqual(rows[0]["review_url"], rows[1]["review_url"])
        self.assertNotEqual(rows[0]["id"], rows[1]["id"])


class Privacy(unittest.TestCase):
    def test_author_hash_is_salted_and_stable(self):
        a = rv.author_hash("Test Person", "trustpilot", "salt-1")
        self.assertEqual(a, rv.author_hash("  test person ", "trustpilot", "salt-1"))
        self.assertNotEqual(a, rv.author_hash("Test Person", "trustpilot", "salt-2"))
        self.assertEqual(len(a), 16)
        self.assertEqual(rv.author_hash("", "trustpilot", "s"), "")

    def test_no_raw_names_survive_normalising_or_redaction(self):
        names = ["Test Person", "Map Person", "PR Person", "G2 Person", "Cap Person", "redditor1"]
        for platform in ("trustpilot", "google_maps", "productreview", "g2", "capterra", "reddit"):
            dumped = json.dumps(norm(platform)) + json.dumps(rv.redact(raw(platform)))
            for name in names:
                self.assertNotIn(name, dumped, platform)


class Dedupe(unittest.TestCase):
    def test_same_id_and_same_author_text_are_dropped(self):
        a = norm("trustpilot")[0]
        twin = dict(a, id="other-id")  # same author + text, different id
        other = dict(a, id="z", author_hash="different", text="Different text")
        self.assertEqual([r["id"] for r in rv.dedupe([a, a, twin, other])], [a["id"], "z"])

    def test_anonymous_identical_texts_are_kept(self):
        a = {"id": "1", "platform": "google_maps", "author_hash": "", "text": "Great"}
        b = dict(a, id="2")
        self.assertEqual(len(rv.dedupe([a, b])), 2)


class Costs(unittest.TestCase):
    spec = {"price_per_review_usd": 0.002, "start_fee_usd": 0.09}

    def test_estimate_maths(self):
        self.assertEqual(rv.est_cost(self.spec, 1560, 500), (500, round(0.09 + 500 * 0.002, 4)))
        self.assertEqual(rv.est_cost(self.spec, 63, 500), (63, round(0.09 + 63 * 0.002, 4)))
        self.assertEqual(rv.est_cost(self.spec, None, 30), (30, round(0.09 + 30 * 0.002, 4)))
        self.assertEqual(rv.est_cost(self.spec, 200, None), (200, 0.49))
        self.assertEqual(rv.est_cost(self.spec, None, None), (None, None))
        self.assertEqual(rv.est_cost(None, 100, 500), (None, None))

    def test_minimum_cap_is_respected(self):
        spec = {"price_per_review_usd": 0.001, "start_fee_usd": 0.005, "min_max": 10}
        self.assertEqual(rv.est_cost(spec, 4, 500), (10, 0.015))

    def test_charge_cap_never_below_the_actor_floor(self):
        self.assertEqual(rv.charge_cap({"price_per_review_usd": 0.00075, "charge_cap_floor_usd": 0.45}, 5), 0.45)
        self.assertEqual(rv.charge_cap({"price_per_review_usd": 0.002}, 500), round(500 * 0.002 * 1.5 + 0.02, 4))

    def test_template_fill_keeps_number_types(self):
        out = rv.actor_input({"input": {"urls": ["{url}"], "n": "{max}", "h": "{half}", "q": "x {query}"}},
                             "https://a", 30)
        self.assertEqual(out, {"urls": ["https://a"], "n": 30, "h": 15, "q": "x https://a"})

    def test_run_cost_prefers_apify_usage(self):
        self.assertEqual(rv.run_cost_usd({"usageTotalUsd": 0.0433}), 0.0433)


class Scoring(unittest.TestCase):
    def test_priority_formula(self):
        # 0.2 share x 2 severity x 1.5 recency x 1.25 visibility x (1 + 0.5 unreplied)
        self.assertEqual(rv.priority_score(0.2, 2.0, 1.5, 1.25, 0.5), round(0.2 * 2 * 1.5 * 1.25 * 1.5, 4))
        self.assertEqual(rv.priority_score(0.1, 1, 1, 1, 0), 0.1)

    def test_components_from_reviews(self):
        sc = rv.load_config("scoring.json")
        negs = [{"id": "a", "platform": "reddit", "date": "2026-09-01", "owner_reply": ""},
                {"id": "b", "platform": "trustpilot", "date": "2025-01-01", "owner_reply": "Sorry"},
                {"id": "c", "platform": "trustpilot", "date": "2026-08-01", "owner_reply": ""}]
        labels = {"a": {"severity": "safety_legal"}, "b": {"severity": "minor"}, "c": {"severity": "cost_money"}}
        comp = rv.fix_components(negs, labels, 10, "2026-10-12", sc)
        self.assertEqual(comp["volume_share"], 0.3)
        self.assertEqual(comp["avg_severity"], 2.0)          # (3 + 1 + 2) / 3
        self.assertEqual(comp["recency_weight"], round(5 / 3, 3))  # 2, 1, 2
        self.assertEqual(comp["visibility_weight"], round((1.5 + 1 + 1) / 3, 3))
        self.assertEqual(comp["unreplied_share"], 0.5)       # Reddit cannot be replied to; 1 of 2 Trustpilot unreplied

    def test_reddit_is_excluded_from_the_star_average(self):
        rows = [{"platform": "trustpilot", "rating": 5.0}, {"platform": "trustpilot", "rating": 3.0},
                {"platform": "reddit", "rating": 1.0}]
        self.assertEqual(rv.star_average(rows), (4.0, 2))

    def test_trend_direction_needs_ten_each_side(self):
        sc = rv.load_config("scoring.json")
        rising = [{"date": "2026-09-01"}] * 15 + [{"date": "2026-05-01"}] * 10
        self.assertEqual(rv.trend(rising, "2026-10-12", sc)["direction"], "rising")
        falling = [{"date": "2026-09-01"}] * 10 + [{"date": "2026-05-01"}] * 15
        self.assertEqual(rv.trend(falling, "2026-10-12", sc)["direction"], "falling")
        flat = [{"date": "2026-09-01"}] * 10 + [{"date": "2026-05-01"}] * 11
        self.assertEqual(rv.trend(flat, "2026-10-12", sc)["direction"], "flat")
        few = [{"date": "2026-09-01"}] * 12 + [{"date": "2026-05-01"}] * 9  # 9 on one side
        self.assertEqual(rv.trend(few, "2026-10-12", sc)["direction"], "too few to call")
        self.assertEqual(rv.trend([{"date": "2026-09-01"}] * 4, "2026-10-12", sc)["direction"], "too few to call")

    def test_monthly_drops_thin_months_and_says_why(self):
        rows = [{"id": str(i), "platform": "trustpilot", "rating": 5.0, "date": "2026-09-0%d" % (i % 9 + 1)}
                for i in range(6)] + [{"id": "x", "platform": "trustpilot", "rating": 1.0, "date": "2025-01-01"}]
        months = rv.monthly(rows, {}, 5)
        self.assertEqual([m["month"] for m in months], ["2026-09"])
        agg = {"meta": {"max_per_platform": 30}, "scoring": {"monthly_min_reviews": 5}, "monthly": months,
               "monthly_dropped": 1}
        text = rv.monthly_table(agg)
        self.assertIn("newest 30 reviews per platform", text)
        self.assertIn("not a trend", text)
        self.assertNotIn("2025-01", text)


class Framing(unittest.TestCase):
    aspect = {"id": "delivery", "label": "Delivery", "description": "Late or missed delivery days and couriers."}

    def test_snippet_starts_at_the_sentence_about_the_aspect(self):
        text = ("Lovely couch and great colours. We were thrilled at first. The delivery was three weeks late "
                "and nobody told us. Would not buy again.")
        snip, states = rv.issue_snippet(text, rv.aspect_terms(self.aspect))
        self.assertTrue(states)
        self.assertTrue(snip.startswith("… The delivery was three weeks late"))
        self.assertIn(snip.lstrip("… ").rstrip(" …"), text)  # still verbatim

    def test_quotes_rank_stating_snippets_and_skip_happy_reviews(self):
        negs = [{"id": "happy", "rating": 5.0, "date": "2026-10-01",
                 "text": "Best purchase ever. The colour is perfect."},
                {"id": "vague", "rating": 1.0, "date": "2026-10-02", "text": "Terrible company, avoid them."},
                {"id": "clear", "rating": 2.0, "date": "2026-01-01",
                 "text": "Nice fabric. Delivery was late twice and the courier never called."}]
        labels = {r["id"]: {"aspects": {"delivery": -2}} for r in negs}
        quotes = rv.pick_quotes(negs, labels, self.aspect)
        self.assertEqual([q["id"] for q in quotes], ["clear", "vague"])  # 5 star dropped, stating first
        self.assertTrue(quotes[0]["states_issue"])
        self.assertTrue(quotes[0]["snippet"].startswith("… Delivery was late"))

    def test_praise_never_ranks_a_net_complaint(self):
        tax = {"aspects": [{"id": "support", "label": "Support"}, {"id": "assembly", "label": "Assembly"}]}
        reviews = [{"id": f"s{i}", "platform": "trustpilot", "rating": 1.0, "date": "2026-09-01", "text": "x"}
                   for i in range(24)] + \
                  [{"id": f"p{i}", "platform": "trustpilot", "rating": 5.0, "date": "2026-09-01", "text": "x"}
                   for i in range(8)]
        labels = {f"s{i}": {"id": f"s{i}", "aspects": {"support": -2}, "severity": "minor"} for i in range(24)}
        labels.update({f"p{i}": {"id": f"p{i}", "aspects": {"support": 2, "assembly": 2}, "severity": "none"}
                       for i in range(8)})
        agg = rv.build_aggregate(reviews, labels, tax, rv.load_config("scoring.json"), "2026-10-12", {},
                                 {"brand": "Acme", "max_per_platform": 30})
        self.assertEqual(agg["praised"], ["assembly"])  # support has 8 positive but 24 negative
        lines = rv.headline_lines(agg, {})
        praised = next(line for line in lines if "Most praised" in line)
        self.assertIn("Assembly (8 positive / 0 negative mentions)", praised)
        net = next(line for line in lines if "Net sentiment" in line)
        self.assertIn("32 labelled reviews (32 name an aspect)", net)

    def test_small_counts_never_read_as_a_pattern(self):
        self.assertEqual(rv.few(1, "negative reviews", "one negative review"), "one negative review")
        self.assertEqual(rv.few(5, "negative reviews"), "5 negative reviews")
        agg = {"fixes": [{"id": "warranty", "negative": 2}]}
        notes = {"issues": {"warranty": {"root_cause": "Complaints cluster on one product line."}}}
        self.assertTrue(rv.pattern_word_warnings(agg, notes))
        self.assertEqual(rv.default_problem({"label": "Warranty", "negative": 1}), "Warranty: single mentions")


class Phrases(unittest.TestCase):
    def test_ngrams_drop_stopwords_and_brand(self):
        blocked = rv.brand_words("Acme Widgets", "acme.example")
        grams = rv.ngrams_of("The Acme delivery was really late. Customer service never replied to me", blocked)
        self.assertIn("customer service", grams)
        self.assertIn("service never replied", grams)
        for g in grams:
            words = g.split()
            self.assertNotIn(words[0], rv.STOPWORDS, g)
            self.assertNotIn(words[-1], rv.STOPWORDS, g)
            self.assertFalse({"acme", "widgets"} & set(words), g)
        self.assertNotIn("the acme", grams)

    def test_ngrams_never_cross_sentences(self):
        grams = rv.ngrams_of("Fast delivery. Great price", set())
        self.assertNotIn("delivery great", grams)

    def test_shared_phrases_score_near_zero(self):
        from collections import Counter
        z = rv.log_odds(Counter({"customer service": 10, "late delivery": 10}),
                        Counter({"customer service": 10, "refund never": 10}), 20, 20)
        self.assertAlmostEqual(z["customer service"], 0.0, places=6)
        self.assertGreater(z["late delivery"], 0)
        self.assertLess(z["refund never"], 0)


class Discovery(unittest.TestCase):
    def test_predictable_urls(self):
        plats = [p for p in rv.platform_index().values() if p["id"] in ("trustpilot", "productreview")]
        urls = [c["url"] for c in rv.predictable_candidates("Acme Widgets", "acme.example", plats)]
        self.assertIn("https://www.trustpilot.com/review/acme.example", urls)
        self.assertIn("https://www.productreview.com.au/listings/acme-widgets", urls)

    def test_country_and_industry_filter(self):
        ids = {p["id"] for p in rv.eligible_platforms("US", "", False)}
        self.assertNotIn("productreview", ids)  # AU only
        self.assertNotIn("g2", ids)             # SaaS only
        self.assertNotIn("glassdoor", ids)      # opt-in
        self.assertIn("g2", {p["id"] for p in rv.eligible_platforms("AU", "B2B SaaS", False)})

    def test_site_links_read_sameas_and_hrefs(self):
        page = ('<a href="https://www.trustpilot.com/review/acme.example">TP</a><script type="application/ld+json">'
                '{"@type":"Organization","sameAs":["https://www.facebook.com/acme"]}</script>')
        links = rv.site_links(page)
        self.assertIn("https://www.trustpilot.com/review/acme.example", links)
        self.assertIn("https://www.facebook.com/acme", links)

    def test_profile_stats_parsers(self):
        page = ('<script type="application/ld+json">{"@type":"Product","name":"Acme Widgets","aggregateRating":'
                '{"ratingValue":"4.2","reviewCount":"1,560"},"review":[{"@type":"Review","datePublished":'
                '"2026-10-01"},{"@type":"Review","datePublished":"2024-01-01"}]}</script>')
        s = rv.merge_stats(rv.jsonld_stats(page), rv.regex_stats(page))
        self.assertEqual((s["rating"], s["reviews"]), (4.2, 1560))
        self.assertEqual(rv.recency(s["dates"], "2026-10-12"), {"latest": "2026-10-01", "last_12m": 1})
        self.assertEqual(rv.recency(["2026-10-01", "2026-09-01"], "2026-10-12")["last_12m"], ">=2")
        self.assertEqual(rv.entity_match({"name": "Acme Widgets", "site": "au.acme.example"}, "", "Acme Widgets",
                                         "acme.example"), "match")
        self.assertEqual(rv.entity_match({"name": "Acme Widgets Ltd UK"}, "", "Acme Widgets", "acme.example"),
                         "partial")
        self.assertEqual(rv.flag("partial"), "⚠ partial match")


class Labels(unittest.TestCase):
    tax = json.loads((FIX / "run" / "data" / "taxonomy.json").read_text(encoding="utf-8"))

    def test_taxonomy_validation(self):
        self.assertEqual(rv.validate_taxonomy(self.tax), [])
        short = {"aspects": self.tax["aspects"][:3]}
        self.assertTrue(any("8 to 15" in e for e in rv.validate_taxonomy(short)))
        bad = {"aspects": [dict(a, owner="ceo") for a in self.tax["aspects"]]}
        self.assertTrue(any("owner" in e for e in rv.validate_taxonomy(bad)))

    def test_jev_request_and_parse(self):
        review = {"id": "r1", "platform": "trustpilot", "rating": 1.0, "title": "", "text": "Refund never came."}
        req = rv.jev_request(review, self.tax)
        self.assertEqual(len(req["questions"]), 2 * len(self.tax["aspects"]) + 2)
        answers = {f"mention__{a['id']}": {"type": "noul", "noul": 0.1} for a in self.tax["aspects"]}
        answers.update({"mention__refunds": {"noul": 0.93}, "polarity__refunds": {"score": 0.2},
                        "severity": {"choice": "cost_money"}, "churn": {"noul": 0.8}})
        lab = rv.parse_jev("r1", {"answers": answers}, self.tax, 0.5)
        self.assertEqual(lab["aspects"], {"refunds": -2})
        self.assertEqual((lab["severity"], lab["churn"]), ("cost_money", True))
        answers["polarity__refunds"] = {"score": 0.0}  # a zero score is strongly negative, not missing
        self.assertEqual(rv.parse_jev("r1", {"answers": answers}, self.tax, 0.5)["aspects"], {"refunds": -2})

    def test_import_validation(self):
        ids = {a["id"] for a in self.tax["aspects"]}
        self.assertEqual(rv.label_problem({"id": "x", "aspects": {"refunds": -2}}, ids, {"x"}), "")
        self.assertIn("bad aspect", rv.label_problem({"id": "x", "aspects": {"refunds": -3}}, ids, {"x"}))
        self.assertIn("unknown review", rv.label_problem({"id": "y", "aspects": {}}, ids, {"x"}))


class Path_(unittest.TestCase):
    def test_outside_a_brain(self):
        with tempfile.TemporaryDirectory() as tmp:
            got = rv.run_dir_for("Acme Widgets", "2026-10-12", Path(tmp))
            self.assertEqual(got, Path(tmp).resolve() / "outputs/geo/2026-10/acme-widgets/review-sentiment")


class OfflinePipeline(unittest.TestCase):
    """The committed fixture (62 synthetic reviews, 3 platforms, pre-made labels) through
    aggregate -> wordcloud -> report -> workbook."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.run_dir = Path(self.tmp.name) / "2026-10" / "review-sentiment"
        shutil.copytree(FIX / "run", self.run_dir)

    def tearDown(self):
        self.tmp.cleanup()

    def quiet(self, fn, **kw):
        with mock.patch("builtins.print"):
            return fn(Namespace(run_dir=str(self.run_dir), **kw))

    def test_aggregate_and_report(self):
        self.assertEqual(self.quiet(rv.cmd_aggregate, scoring=None, as_of="2026-10-12"), 0)
        agg = json.loads((self.run_dir / "data" / "aggregate.json").read_text(encoding="utf-8"))
        self.assertEqual(agg["meta"]["reviews"], 62)
        self.assertEqual(agg["meta"]["excluded_from_stars"], ["reddit"])
        self.assertEqual(agg["headline"]["rated"], 52)  # 62 minus 10 Reddit
        ranks = [f["priority"] for f in agg["fixes"]]
        self.assertEqual(ranks, sorted(ranks, reverse=True))
        self.assertEqual(self.quiet(rv.cmd_report), 0)
        for name in (rv.REPORT, rv.FIX_REPORT, rv.UNREPLIED):
            self.assertTrue((self.run_dir / name).is_file(), name)
        fix = (self.run_dir / rv.FIX_REPORT).read_text(encoding="utf-8")
        self.assertIn("hypothesis", fix)
        self.assertNotIn("—", fix + (self.run_dir / rv.REPORT).read_text(encoding="utf-8"))  # no em dashes
        rows = (self.run_dir / rv.UNREPLIED).read_text(encoding="utf-8").splitlines()
        self.assertEqual(len(rows) - 1, len(agg["unreplied_negatives"]))

    @unittest.skipUnless(HAS_WC, "needs wordcloud + matplotlib (uv run --with wordcloud --with matplotlib)")
    def test_wordcloud(self):
        self.assertEqual(self.quiet(rv.cmd_wordcloud, min_count=2, top=60, block=None), 0)
        for name in ("wordcloud-positive.png", "wordcloud-negative.png", "wordcloud.html"):
            self.assertGreater((self.run_dir / name).stat().st_size, 1000, name)
        page = (self.run_dir / "wordcloud.html").read_text(encoding="utf-8")
        self.assertIn("prefers-color-scheme:dark", page)
        self.assertNotIn("<script src", page)  # self-contained
        phrases = json.loads((self.run_dir / "data" / "phrases.json").read_text(encoding="utf-8"))
        for side in ("positive", "negative"):
            self.assertTrue(phrases[side])
            self.assertTrue(all(len(p["quotes"]) <= 3 for p in phrases[side]))
            self.assertFalse(any("acme" in p["phrase"] for p in phrases[side]))

    @unittest.skipUnless(HAS_XL, "needs openpyxl (uv run --with openpyxl)")
    def test_workbook(self):
        import openpyxl
        self.quiet(rv.cmd_aggregate, scoring=None, as_of="2026-10-12")
        self.assertEqual(self.quiet(rv.cmd_workbook), 0)
        book = self.run_dir.parent / rv.WORKBOOK
        wb = openpyxl.load_workbook(book)
        self.assertIn(rv.TAB, wb.sheetnames)
        self.assertLess(wb.sheetnames.index(rv.TAB), wb.sheetnames.index("Initiatives"))
        ws = wb[rv.TAB]
        ws.cell(row=7, column=15, value="Agree")  # the client's verdict
        wb.save(book)
        tasks_before = [r[3] for r in wb["Initiatives"].iter_rows(min_row=6, values_only=True) if r[3]]
        self.quiet(rv.cmd_workbook)  # re-run keeps the verdict and adds no duplicate tasks
        wb = openpyxl.load_workbook(book)
        self.assertEqual(wb[rv.TAB].cell(row=7, column=15).value, "Agree")
        tasks_after = [r[3] for r in wb["Initiatives"].iter_rows(min_row=6, values_only=True) if r[3]]
        self.assertEqual(sorted(tasks_before), sorted(tasks_after))
        self.assertTrue(any("unreplied" in t for t in tasks_after))
        checklist = {r[2]: r for r in wb["Checklist"].iter_rows(values_only=True) if r[2]}
        self.assertEqual(checklist[rv.SKILL_ID][8], "☑")


if __name__ == "__main__":
    unittest.main()
