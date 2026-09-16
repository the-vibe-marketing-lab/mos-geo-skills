#!/usr/bin/env python3
"""Build "Brand Audit Template.xlsx", the master GEO brand audit workbook.

The workbook is generated, never hand-edited, so every mos-geo-* skill is
added the same way: add an entry to skills.json, re-run this script, commit
both files.

    uv run --with openpyxl python _shared/brand-audit/build_template.py

Tabs
  Checklist            one row per mos-geo-* skill: link, run command, output,
                       status and a tick box
  Brand Truth Review   one row per factual claim from the brand-360 report;
                       the client marks each Accurate / Partly / Inaccurate
  Initiatives          ICE-scored fixes that come out of the audit
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
OUT = HERE / "Brand Audit Template.xlsx"

INK = "1F2937"
PAPER = "FFFFFF"
MUTED = "F3F4F6"
LINE = "D1D5DB"
DONE = "DCFCE7"
FLAG = "FEE2E2"
ROWS = 60  # blank, formatted rows under each table

HEAD = Font(bold=True, color=PAPER, size=11)
HEAD_FILL = PatternFill("solid", fgColor=INK)
NOTE_FILL = PatternFill("solid", fgColor=MUTED)
TITLE = Font(bold=True, size=16, color=INK)
LINK = Font(color="1D4ED8", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
EDGE = Border(*(Side(style="thin", color=LINE),) * 4)

TICK = '"☐,☑"'


def base(ws, title: str, notes: list[str], widths: dict[str, float]) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws["B2"] = title
    ws["B2"].font = TITLE
    row = 3
    for note in notes:
        ws.cell(row=row, column=2, value=note).alignment = Alignment(wrap_text=False, vertical="top")
        row += 1


def table(ws, top: int, headers: list[str]) -> None:
    for i, name in enumerate(headers, start=2):
        c = ws.cell(row=top, column=i, value=name)
        c.font, c.fill, c.alignment, c.border = HEAD, HEAD_FILL, WRAP, EDGE
    ws.row_dimensions[top].height = 32
    ws.freeze_panes = ws.cell(row=top + 1, column=2)
    ws.auto_filter.ref = f"B{top}:{ws.cell(row=top, column=len(headers) + 1).column_letter}{top + ROWS}"
    for r in range(top + 1, top + ROWS + 1):
        for i in range(2, len(headers) + 2):
            c = ws.cell(row=r, column=i)
            c.alignment, c.border = WRAP, EDGE


def dropdown(ws, options: str, cells: str, prompt: str) -> None:
    dv = DataValidation(type="list", formula1=options, allow_blank=True, showDropDown=False)
    dv.promptTitle, dv.prompt, dv.showInputMessage = "Pick one", prompt, True
    ws.add_data_validation(dv)
    dv.add(cells)


def checklist(wb: Workbook, registry: dict) -> None:
    ws = wb.active
    ws.title = "Checklist"
    base(ws, "GEO Brand Audit: Checklist", [
        "Client: [CLIENT]      Brand name (as customers say it): [BRAND]      Industry: [INDUSTRY]      Month: [YYYY-MM]",
        "How to use: run each skill below in the client's MarketingOS brain, paste the output location, set Status, tick Done.",
        "Rows are generated from the mos-geo-skills pack. New skills appear here when the template is rebuilt.",
    ], {"B": 20, "C": 24, "D": 60, "E": 20, "F": 44, "G": 22, "H": 16, "I": 8, "J": 13, "K": 40})
    top = 7
    headers = ["Category", "Skill", "What it does", "Run", "Output (where the result is saved)",
               "Feeds tab", "Status", "Done", "Date run", "Link to result / notes"]
    table(ws, top, headers)
    for n, s in enumerate(registry["skills"], start=top + 1):
        link = f"{registry['repo']}/tree/main/{s['skill']}"
        values = [s["category"], s["skill"], s["what_it_does"], s["run"], s["output"],
                  s["feeds_tab"], "Not started", "☐", None, None]
        for i, v in enumerate(values, start=2):
            ws.cell(row=n, column=i, value=v)
        skill = ws.cell(row=n, column=3)
        skill.hyperlink, skill.font = link, LINK
        ws.cell(row=n, column=5).font = Font(name="Consolas")
        ws.cell(row=n, column=6).font = Font(name="Consolas", size=9)
        ws.cell(row=n, column=10).number_format = "yyyy-mm-dd"
    last = top + ROWS
    dropdown(ws, '"Not started,In progress,Client review,Completed,Not needed"', f"H{top + 1}:H{last}",
             "Where this skill is up to.")
    dropdown(ws, TICK, f"I{top + 1}:I{last}", "Tick when the output is delivered.")
    ws.conditional_formatting.add(
        f"B{top + 1}:K{last}",
        FormulaRule(formula=[f'$I{top + 1}="☑"'], fill=PatternFill("solid", fgColor=DONE)))
    ws.cell(row=top + len(registry["skills"]) + 2, column=2,
            value="Needs: " + "; ".join(f"{s['skill']}: {s['needs']}" for s in registry["skills"])
            ).font = Font(italic=True, size=9, color="6B7280")


def truth_review(wb: Workbook) -> None:
    ws = wb.create_sheet("Brand Truth Review")
    base(ws, "Brand Truth Review: is this what AI should say about [BRAND]?", [
        "Replaces pasting the whole Brand Brain report. Each row is one factual claim from the brand-360 report "
        "(research sections 1 to 17 and what the AI engines said in Section 18).",
        "Client: mark each claim in 'Your verdict' and write the correct version where it is wrong. "
        "Skip anything you are not sure about; 'Not sure' is a valid answer.",
        "Agency: every Inaccurate or Partly row becomes an Initiative (fix the page, profile or schema the claim came from).",
    ], {"B": 6, "C": 20, "D": 55, "E": 30, "F": 22, "G": 16, "H": 45, "I": 12, "J": 30})
    top = 7
    headers = ["#", "Topic", "What the research / AI says", "Where it came from", "Said by (AI surfaces)",
               "Your verdict", "Correct version (client)", "Fix priority", "Agency notes"]
    table(ws, top, headers)
    last = top + ROWS
    for r in range(top + 1, last + 1):
        ws.cell(row=r, column=2, value=f"=IF(C{r}=\"\",\"\",ROW()-{top})")
    dropdown(ws, '"Accurate,Partly accurate,Inaccurate,Not sure,Out of date"', f"G{top + 1}:G{last}",
             "Is this claim true today?")
    dropdown(ws, '"High,Medium,Low"', f"I{top + 1}:I{last}", "How much a wrong answer here hurts.")
    for text, colour in (("Inaccurate", FLAG), ("Out of date", FLAG), ("Accurate", DONE)):
        ws.conditional_formatting.add(
            f"G{top + 1}:G{last}",
            FormulaRule(formula=[f'$G{top + 1}="{text}"'], fill=PatternFill("solid", fgColor=colour)))
    example = ["(example)", "Founder", "Run by [NAME] in [CITY].", "Section 1, source [2]",
               "ChatGPT (app), Claude (API)", None, None, None, "Delete this row"]
    for i, v in enumerate(example, start=2):
        c = ws.cell(row=top + 1, column=i, value=v)
        c.font, c.fill = Font(italic=True, color="6B7280"), NOTE_FILL


def initiatives(wb: Workbook) -> None:
    ws = wb.create_sheet("Initiatives")
    base(ws, "Initiatives: fixes from the audit, ranked by ICE", [
        "ICE Score = Impact x Confidence / Ease (Ease: 1 = easy, 10 = hard). Sort by ICE Score, highest first.",
    ], {"B": 9, "C": 20, "D": 40, "E": 50, "F": 22, "G": 20, "H": 12, "I": 18, "J": 11, "K": 11, "L": 11, "M": 30})
    top = 5
    headers = ["ICE Score", "Category", "Task", "Instructions", "From skill", "Status", "Needs dev?",
               "Owner", "Impact (10 = high)", "Confidence (10 = high)", "Ease (1 = easy)", "Deliverable / link"]
    table(ws, top, headers)
    last = top + ROWS
    for r in range(top + 1, last + 1):
        ws.cell(row=r, column=2, value=f'=IFERROR(ROUND(J{r}*K{r}/L{r},1),"")')
    dropdown(ws, '"Brand Optimisation,Content Strategy,Competitive Analysis,Technical,Digital PR"',
             f"C{top + 1}:C{last}", "Type of work.")
    dropdown(ws, '"Scheduled,In progress (agency),Implementing (client),QA,Completed"',
             f"G{top + 1}:G{last}", "Where this fix is up to.")
    dropdown(ws, '"Yes,No"', f"H{top + 1}:H{last}", "Does a developer need to do this?")
    dropdown(ws, '"[AGENCY],[CLIENT],[DEV TEAM]"', f"I{top + 1}:I{last}", "Who owns it.")
    score = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True)
    score.error, score.showErrorMessage = "Use a whole number from 1 to 10.", True
    ws.add_data_validation(score)
    score.add(f"J{top + 1}:L{last}")
    ws.conditional_formatting.add(
        f"B{top + 1}:M{last}",
        FormulaRule(formula=[f'$G{top + 1}="Completed"'], fill=PatternFill("solid", fgColor=DONE)))


def main() -> None:
    registry = json.loads((HERE / "skills.json").read_text(encoding="utf-8"))
    wb = Workbook()
    checklist(wb, registry)
    truth_review(wb)
    initiatives(wb)
    wb.properties.title = "Brand Audit Template"
    wb.save(OUT)
    print(f"Wrote {OUT} ({len(registry['skills'])} skills)")


if __name__ == "__main__":
    main()
