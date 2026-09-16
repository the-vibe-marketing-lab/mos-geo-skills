#!/usr/bin/env python3
"""Tick a skill's row in the month's brand-audit-master.xlsx.

Every mos-geo-* skill saves into campaigns/geo/YYYY-MM/<skill-folder>/ and
then calls this so the Checklist tab at the month root stays current:

    uv run --with openpyxl python _shared/brand-audit/tick_checklist.py \
        --skill mos-geo-llm-buttons --run-dir campaigns/geo/2026-09/llm-buttons \
        --status "Client review" --note "brief, code reference and demo in llm-buttons/"

If the month has no workbook yet, the blank template from this folder is
copied in first. mos-geo-brand-360 has its own richer `workbook` step and does
not use this script.
"""

from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "brand-audit-template.xlsx"
MASTER = "brand-audit-master.xlsx"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--skill", required=True, help="skill id as it appears in the Checklist, e.g. mos-geo-llm-buttons")
    ap.add_argument("--run-dir", required=True, help="the skill's folder inside the month folder")
    ap.add_argument("--status", default="Client review",
                    choices=["Not started", "In progress", "Client review", "Completed", "Not needed"])
    ap.add_argument("--note", default="", help="goes in 'Link to result / notes'")
    ap.add_argument("--date", default=time.strftime("%Y-%m-%d"))
    args = ap.parse_args()
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is needed: uv run --with openpyxl python tick_checklist.py ...")

    run_dir = Path(args.run_dir).resolve()
    book = run_dir.parent / MASTER
    if not book.is_file():
        if not TEMPLATE.is_file():
            sys.exit(f"no {MASTER} in {run_dir.parent} and no template at {TEMPLATE}")
        book.write_bytes(TEMPLATE.read_bytes())
    wb = openpyxl.load_workbook(book)
    ws = wb["Checklist"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=3).value == args.skill:
            ws.cell(row=r, column=8, value=args.status)
            ws.cell(row=r, column=9, value="☑" if args.status in ("Client review", "Completed") else "☐")
            ws.cell(row=r, column=10, value=args.date)
            ws.cell(row=r, column=11, value=args.note or f"files in {run_dir.name}/")
            break
    else:
        sys.exit(f"Checklist has no row for {args.skill}; add it to skills.json and rebuild the template")
    wb.save(book)
    print(f"Ticked {args.skill} in {book}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
