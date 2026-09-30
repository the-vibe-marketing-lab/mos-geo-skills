#!/usr/bin/env python3
"""mos-geo-homepage-faqs: homepage FAQs an LLM can lift one answer at a time.

Subcommands
  path      print this month's run folder (same rules as every mos-geo skill)
  doctor    preflight: Scrapling on PATH, uv, openpyxl, the sibling crawler, the workbook template
  crawl     sitemap(s) + homepage nav + same-host BFS -> data/pages/<slug>.md + data/link-inventory.csv
  gather    data/public-bundle.md (crawl, --include, interview, research Summary: the provenance source)
            + data/input-bundle.md (all that plus this month's ai-info / brand-360 / truths and the brain)
  links     BM25 shortlist of the 3 best internal pages per answer -> data/link-candidates.json
  lint      the deterministic gate (answer formula, provenance, links, voice) -> data/lint.md
  render    faqs.html, faqpage.jsonld, handover.md, research-note.md
  workbook  tick the Checklist, add the 'Homepage FAQs' tab and a publish row on Initiatives

faqs.md format (the file Claude writes at <run>/faqs.md)
---------------------------------------------------------
    # {Brand} FAQs                         <- optional title line

    ### What is {Brand}?                   <- one H3 per FAQ, the question, ending "?"
    {Brand} is a ... [descriptive anchor](https://brand.example/page/) ... .
    <!-- bucket: identity; placement: homepage; sources: https://brand.example/, https://brand.example/about/; evidence: first-party -->

  * The answer is every non-empty line between the H3 and the next H3 (joined with spaces),
    minus the metadata comment. It must be ONE sentence.
  * Links are inline markdown `[anchor](url)`; the anchor is words already in the answer.
    At most one per answer. `data/links.json` ([{"question" or "n", "target", "anchor"}]) is
    also accepted: lint checks the anchor is verbatim in the answer and render applies it. Inline
    links in faqs.md are the source of truth; an entry repeated in both counts once.
  * Metadata comment keys, `;`-separated `key: value`:
      bucket     identity | product | capability | proof | audience | resources | differentiator | founder
      placement  homepage | faq-page        (default: first 8 homepage, the rest faq-page)
      sources    comma-separated URLs the answer came from (at least one). This IS the evidence
                 ledger: there is no separate ledger file.
      evidence   first-party (brand's own pages) | third-party (independent sources) | mixed (both);
                 third-party and mixed need at least one source off the brand's domain
      theme      optional ledger theme name

Standard library only, except `workbook` (openpyxl, via `uv run --with openpyxl`). Scrapling
(a CLI installed with `uv tool install "scrapling[fetchers]"`) fetches page content; without it
the crawl falls back to the mos-geo-ai-info crawler (urllib).
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import html
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
SIBLING = SKILL_DIR.parent / "mos-geo-ai-info" / "scripts" / "aiinfo.py"
PACK_TEMPLATE = SKILL_DIR.parent / "_shared" / "brand-audit" / "brand-audit-template.xlsx"
SKILL_ID = "mos-geo-homepage-faqs"
SKILL_FOLDER = "homepage-faqs"
DATA = "data"
BUNDLE = "data/input-bundle.md"
PUBLIC_BUNDLE = "data/public-bundle.md"
INVENTORY = "data/link-inventory.csv"
CANDIDATES = "data/link-candidates.json"
LINKS_JSON = "data/links.json"
RESEARCH = "data/research.md"
FAQS_MD, FAQS_HTML, JSONLD, HANDOVER, NOTE = "faqs.md", "faqs.html", "faqpage.jsonld", "handover.md", "research-note.md"
WORKBOOK = "brand-audit-master.xlsx"
TAB = "Homepage FAQs"
INV_COLS = ["url", "status", "final_url", "canonical", "indexable", "title", "h1", "meta_description",
            "page_type", "word_count"]
# (min, max) per bucket; kept in step with references/question-bank.md "Target counts".
BUCKET_LIMITS = {"identity": (1, 1), "product": (0, 3), "capability": (4, 8), "proof": (1, 4), "audience": (2, 4),
                 "resources": (0, 2), "differentiator": (0, 2), "founder": (0, 1)}
BUCKETS = set(BUCKET_LIMITS)
PLACEMENTS = {"homepage", "faq-page"}
EVIDENCE = {"first-party", "third-party", "mixed"}
HOMEPAGE_TOP = 8
INSTALL = 'uv tool install "scrapling[fetchers]"  &&  scrapling install'


def sibling():
    if not SIBLING.is_file():
        sys.exit(f"needs mos-geo-ai-info next to this skill ({SIBLING}); install the whole mos-geo-skills pack")
    sys.path.insert(0, str(SIBLING.parent))
    import aiinfo  # noqa: E402
    return aiinfo


class Report:
    def __init__(self):
        self.fails, self.warns, self.notes = [], [], []


# ------------------------------------------------------------------ path --

def cmd_path(args) -> int:
    ai = sibling()
    ai.SKILL_FOLDER = SKILL_FOLDER
    print(ai.run_dir_for(args.brand, args.date or time.strftime("%Y-%m-%d"), Path(args.start or ".")))
    return 0


# ---------------------------------------------------------------- doctor --

def scrapling_bin() -> str | None:
    found = shutil.which("scrapling")
    if found:
        return found
    local = Path.home() / ".local" / "bin" / "scrapling"
    return str(local) if local.is_file() else None


def _has_openpyxl() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def cmd_doctor(args) -> int:
    ok = True
    sc = scrapling_bin()
    if sc:
        print(f"[ok]   scrapling CLI: {sc}")
    else:
        print("[warn] scrapling CLI not found; crawl falls back to the mos-geo-ai-info crawler (urllib).")
        print(f"       install: {INSTALL}")
    uv = shutil.which("uv")
    print(f"[ok]   uv: {uv}" if uv else "[warn] uv not found; workbook needs openpyxl (https://docs.astral.sh/uv/)")
    print("[ok]   openpyxl importable" if _has_openpyxl() else
          "[info] openpyxl not importable here; run workbook via `uv run --with openpyxl python faqs.py workbook ...`")
    if SIBLING.is_file():
        print(f"[ok]   sibling crawler: {SIBLING}")
    else:
        ok = False
        print(f"[fail] missing {SIBLING}: install the whole mos-geo-skills pack")
    print(f"[ok]   workbook template: {PACK_TEMPLATE}" if PACK_TEMPLATE.is_file() else
          f"[warn] no workbook template at {PACK_TEMPLATE}")
    print(f"python {sys.version.split()[0]}")
    return 0 if ok else 1


# ----------------------------------------------------------------- crawl --

EXCLUDE = re.compile(
    r"\.(jpe?g|png|gif|webp|avif|svg|ico|pdf|zip|mp4|mp3|mov|webm|css|js|json|xml|rss|atom|txt|woff2?)(\?|$)|"
    r"/wp-(admin|json|content|includes)/|/(feed|rss|atom)/?$|/(tag|tags|category|categories|author)/|/page/\d+|"
    r"[?&](page|paged|s|p|replytocom)=|/cdn-cgi/|"
    r"/(login|log-in|signin|sign-in|signup|sign-up|register|account|my-account|cart|checkout|logout|wp-login\.php)(/|$|\?)",
    re.I)
PAGE_TYPES = [
    ("blog-index", r"^(blog|news|insights|articles|guides|resources)$"),
    ("blog-post", r"^(blog|news|insights|articles|guides|resources)/."),
    ("about", r"about|who-we-are|our-story|team|founder"), ("contact", r"contact|locations?"),
    ("pricing", r"pricing|plans"), ("legal", r"privacy|terms|cookie|legal|disclaimer"),
    ("service", r"services?|solutions|what-we-do"), ("product", r"products?|shop|features?|platform"),
    ("case-study", r"case-stud|results|portfolio|our-work|clients?"),
    ("faq", r"faq|help"),
]
REDIRECTS = (301, 302, 303, 307, 308)


def clean_url(url: str) -> str:
    """Drop the fragment and lower-case scheme and host; everything else is kept exactly."""
    p = urllib.parse.urlsplit(url.strip())
    return urllib.parse.urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path or "/", p.query, ""))


def is_home(url: str) -> bool:
    p = urllib.parse.urlsplit(url)
    return p.path in ("", "/") and not p.query


def page_type(url: str) -> str:
    if is_home(url):
        return "homepage"
    path = urllib.parse.urlsplit(url).path.strip("/").lower()
    for name, pat in PAGE_TYPES:
        if re.search(pat, path):
            return name
    return "page"


def page_slug(url: str) -> str:
    """Readable path slug (any script) plus a short URL hash, so /a-b, /a/b and /a/ never collide."""
    path = urllib.parse.unquote(urllib.parse.urlsplit(url).path).lower()
    slug = re.sub(r"[^\w/]+", "-", path).strip("-/").replace("/", "__")
    digest = hashlib.sha1(clean_url(url).encode("utf-8")).hexdigest()[:6]
    return f"{(slug or 'home')[:90]}-{digest}"


def crawlable(url: str, host: str, same_site) -> bool:
    return url.startswith(("http://", "https://")) and same_site(url, host) and not EXCLUDE.search(url)


class _NoFollow(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):  # noqa: D401 - urllib hook
        return None


def _open_once(opener, url: str, headers: dict) -> tuple[int, dict, bytes, str]:
    req = urllib.request.Request(url, headers=headers)
    try:
        with opener.open(req, timeout=30) as r:
            body = r.read()
            if r.headers.get("Content-Encoding") == "gzip" or body[:2] == b"\x1f\x8b":
                try:
                    body = gzip.decompress(body)
                except OSError:
                    pass
            return r.status, {k.lower(): v for k, v in r.headers.items()}, body, ""
    except urllib.error.HTTPError as e:
        hdrs = {k.lower(): v for k, v in (e.headers or {}).items()}
        return e.code, hdrs, b"", hdrs.get("location", "")
    except Exception as e:  # noqa: BLE001 - network errors become status 0
        return 0, {"error": str(e)}, b"", ""


def probe(url: str, headers: dict, hops: int = 6) -> tuple[int, int, str, dict, bytes]:
    """(first status, final status, final url, final headers, final body), following redirects by hand
    so a 3xx is recorded as a 3xx."""
    opener = urllib.request.build_opener(_NoFollow)
    first, cur = None, url
    for _ in range(hops):
        st, hdrs, body, loc = _open_once(opener, cur, headers)
        first = st if first is None else first
        if st in REDIRECTS and loc:
            cur = clean_url(urllib.parse.urljoin(cur, loc))
            continue
        return first, st, cur, hdrs, body
    return first or 0, 0, cur, {}, b""


def scrapling_body(url: str, binary: str, js_only) -> tuple[bytes, str] | None:
    """get -> fetch -> stealthy-fetch until one returns a 200 page with real text."""
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "page.html"
        for mode in ("get", "fetch", "stealthy-fetch"):
            try:
                res = subprocess.run([binary, "extract", mode, url, str(out)], capture_output=True, text=True,
                                     timeout=120)
            except (subprocess.TimeoutExpired, OSError):
                continue
            m = re.search(r"Fetched \((\d+)\)", res.stdout + res.stderr)
            if not (m and m.group(1) == "200" and out.is_file()):
                continue
            body = out.read_bytes()
            if body and not js_only(body):
                return body, f"scrapling-{mode}"
    return None


def fetch_body(url: str, probe_body: bytes, ai, binary: str | None, fetcher: str) -> tuple[bytes, str]:
    if fetcher != "urllib" and binary:
        got = scrapling_body(url, binary, ai.js_only)
        if got:
            return got
    if probe_body:
        return probe_body, "urllib"
    status, body, _, via = ai.fetch_page(url, 0, "off")
    return (body if status == 200 else b""), via


def noindex(parser, headers: dict) -> bool:
    robots = " ".join([parser.meta.get("robots", ""), parser.meta.get("googlebot", ""),
                       headers.get("x-robots-tag", "")]).lower()
    return "noindex" in robots or "none" in robots.split(",")


def inventory_row(url: str, first: int, final_url: str, parser, hdrs: dict, words: int) -> dict:
    canonical = clean_url(parser.canonical) if parser and parser.canonical else ""
    ok = first == 200 and final_url == url and parser is not None
    indexable = ok and not noindex(parser, hdrs) and canonical in ("", url)
    h1 = [h for t, h in parser.headings if t == "h1"] if parser else []
    return {"url": url, "status": first, "final_url": final_url, "canonical": canonical,
            "indexable": "yes" if indexable else "no",
            "title": html.unescape(parser.title.strip()) if parser else "", "h1": h1[0] if h1 else "",
            "meta_description": parser.meta.get("description", "").strip() if parser else "",
            "page_type": page_type(url), "word_count": words}


def save_text(pages: Path, url: str, parser) -> int:
    txt = parser.text()
    head = [f"# {html.unescape(parser.title.strip())}", "", f"URL: {url}",
            f"Meta description: {parser.meta.get('description', '')}", ""]
    (pages / f"{page_slug(url)}.md").write_text("\n".join(head) + txt + "\n", encoding="utf-8")
    return len(txt.split())


def crawl_one(url: str, ctx: dict) -> tuple[dict, list[str]]:
    """One inventory row plus the same-host links found on the page."""
    ai = ctx["ai"]
    first, final_st, final_url, hdrs, pbody = probe(url, ai.HEADERS)
    if first != 200:
        row = inventory_row(url, first, final_url, None, hdrs, 0)
        nxt = [final_url] if first in REDIRECTS and final_st == 200 else []
        return row, nxt
    body, via = fetch_body(url, pbody, ai, ctx["bin"], ctx["fetcher"])
    ctx["via"][via] = ctx["via"].get(via, 0) + 1
    parser = ai.parse_page(url, body)
    canonical = clean_url(parser.canonical) if parser.canonical else ""
    duplicate = canonical not in ("", url)  # a canonicalised variant: its text belongs to the canonical page
    words = len(parser.text().split()) if duplicate else save_text(ctx["pages"], url, parser)
    row = inventory_row(url, first, final_url, parser, hdrs, words)
    return row, ([canonical] if duplicate else []) + [clean_url(u) for u, _ in parser.links]


def seed_queue(start: str, host: str, ai, limit: int) -> list[str]:
    robots_body = ai.fetch(urllib.parse.urljoin(start, "/robots.txt"))[2]
    robots = ai.text_of(robots_body) if isinstance(robots_body, bytes) else ""
    urls = [clean_url(u) for u in ai.sitemap_urls(start, robots, limit=2000)]
    urls = [u for u in urls if crawlable(u, host, ai.same_site)]
    urls.sort(key=lambda u: -ai.score_url(u))
    return urls[: limit * 4]


def resolve_start(url: str, ai) -> tuple[str, list[dict]]:
    start = clean_url(url if re.match(r"https?://", url) else "https://" + url)
    first, final_st, final_url, _, _ = probe(start, ai.HEADERS)
    if final_st != 200:
        sys.exit(f"homepage returned {final_st or first}: {start}")
    rows = [] if first == 200 else [inventory_row(start, first, final_url, None, {}, 0)]
    return final_url, rows


def enqueue(queue: list[str], seen: set, urls: list[str], host: str, ai) -> None:
    for u in urls:
        if u not in seen and crawlable(u, host, ai.same_site):
            seen.add(u)
            queue.append(u)


def cmd_crawl(args) -> int:
    ai = sibling()
    run = Path(args.run_dir)
    pages = run / DATA / "pages"
    pages.mkdir(parents=True, exist_ok=True)
    start, rows = resolve_start(args.url, ai)
    host = urllib.parse.urlsplit(start).netloc
    ctx = {"ai": ai, "bin": scrapling_bin(), "fetcher": args.fetcher, "pages": pages, "via": {}}
    if not ctx["bin"] and args.fetcher != "urllib":
        print(f"[warn] scrapling CLI not found, using the mos-geo-ai-info crawler. Install: {INSTALL}")
    seen, queue = {start}, [start]
    home_row, home_links = crawl_one(start, ctx)
    rows.append(home_row)
    enqueue(queue, seen, home_links, host, ai)
    enqueue(queue, seen, seed_queue(start, host, ai, args.max), host, ai)
    i = 1
    while i < len(queue) and len(rows) < args.max:
        row, links = crawl_one(queue[i], ctx)
        rows.append(row)
        enqueue(queue, seen, links, host, ai)
        i += 1
        time.sleep(args.delay)
    mark_hubs(rows)
    write_inventory(run, rows)
    idx = sum(r["indexable"] == "yes" for r in rows)
    redir = sum(str(r["status"]).startswith("3") for r in rows)
    print(f"Wrote {run / INVENTORY}: {len(rows)} URLs ({idx} indexable, {redir} redirects); "
          f"pages in {pages} (fetched via {ctx['via'] or 'none'})")
    return 0


def mark_hubs(rows: list[dict]) -> None:
    """A page whose path is a parent of another crawled page is a hub/category, not a post."""
    urls = [str(r["url"]) for r in rows]
    for r in rows:
        base = str(r["url"]).rstrip("/") + "/"
        if r["page_type"] in ("blog-post", "page") and any(u != r["url"] and u.startswith(base) for u in urls):
            r["page_type"] = "hub"


def write_inventory(run: Path, rows: list[dict]) -> None:
    out = run / INVENTORY
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=INV_COLS)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in INV_COLS})


def load_inventory(run: Path) -> dict[str, dict]:
    p = run / INVENTORY
    if not p.is_file():
        return {}
    with p.open(encoding="utf-8", newline="") as fh:
        return {clean_url(r["url"]): r for r in csv.DictReader(fh)}


# ---------------------------------------------------------------- gather --
BRAIN_GLOBS = ["CLAUDE.md", "BRAIN.md", "CONTEXT.md", "reference/core/*.md", "reference/proof/*.md",
               "reference/voice*.md", "reference/voice/*.md", "business/brand/*.md", "business/voice/*.md",
               "business/offer/*.md", "business/offers/*.md", "business/proof/*.md", "business/proof/*/*.md",
               "business/strategy/*.md", "business/audience/*.md", "voice.md"]
FILE_CAP = 6000  # words per file
MONTH_SOURCES = [("ai-info", ["*.md", "data/discrepancies.md", "data/facts.json"]),
                 ("ai-info-draft", ["page.md"]), ("brand-360-report", ["*.md"])]


def latest_run(month: Path, name: str) -> Path | None:
    """name, name-2, name-3 ... : the highest-numbered one that exists."""
    found = [p for p in month.glob(f"{name}*") if p.is_dir() and re.fullmatch(rf"{re.escape(name)}(-\d+)?", p.name)]
    return max(found, key=lambda p: int(p.name.rsplit("-", 1)[1]) if p.name != name else 1) if found else None


def month_files(run: Path) -> list[tuple[str, Path]]:
    out = []
    for name, globs in MONTH_SOURCES:
        d = latest_run(run.resolve().parent, name)
        if d:
            for g in globs:
                out += [(name, p) for p in sorted(d.glob(g)) if p.is_file()]
    return out


def truth_corrections(run: Path) -> str:
    book = run.resolve().parent / WORKBOOK
    if not book.is_file():
        return ""
    try:
        import openpyxl
    except ImportError:
        print("  [info] openpyxl not importable: Brand Truth Review corrections skipped (run gather via uv)")
        return ""
    wb = openpyxl.load_workbook(book, data_only=True)
    if "Brand Truth Review" not in wb.sheetnames:
        return ""
    ws, lines = wb["Brand Truth Review"], []
    for r in range(8, ws.max_row + 1):
        claim, verdict = ws.cell(row=r, column=4).value, ws.cell(row=r, column=7).value
        if claim and verdict:
            fix = ws.cell(row=r, column=8).value
            lines.append(f"- [{verdict}] {ws.cell(row=r, column=3).value}: {claim}"
                         + (f" | correct version: {fix}" if fix else ""))
    return "\n".join(lines)


def gather_files(run: Path, brain: str | None, include: list[str]) -> list[tuple[str, Path]]:
    files: list[tuple[str, Path]] = []
    if brain:
        for g in BRAIN_GLOBS:
            files += [("brain", p) for p in sorted(Path(brain).glob(g)) if p.is_file()]
    for inc in include:
        p = Path(inc)
        found = sorted(p.rglob("*.md")) if p.is_dir() else ([p] if p.is_file() else [])
        if not found:
            print(f"  [warn] --include {inc}: nothing found")
        files += [("include", f) for f in found]
    files += month_files(run)
    files += [("crawl", p) for p in sorted((run / DATA / "pages").glob("*.md"))]
    if (run / DATA / "interview.md").is_file():
        files.append(("interview", run / DATA / "interview.md"))
    return files


PUBLIC_KINDS = {"crawl", "include", "interview"}


def research_summary(run: Path) -> str:
    """Only the public '## Summary' section of data/research.md; the rest is private working notes."""
    p = run / RESEARCH
    if not p.is_file():
        return ""
    m = re.search(r"^##\s+Summary\s*$(.*?)(?=^##\s|\Z)", p.read_text(encoding="utf-8"), re.M | re.S | re.I)
    return strip_brokers(m.group(1).strip()) if m else ""


def cmd_gather(args) -> int:
    run = Path(args.run_dir)
    seen, parts, manifest, public = set(), [], [], []
    for kind, p in gather_files(run, args.brain, args.include or []):
        if p.resolve() in seen:
            continue
        seen.add(p.resolve())
        raw = p.read_text(encoding="utf-8", errors="replace")
        words = raw.split()
        text = " ".join(words[:FILE_CAP]) if len(words) > FILE_CAP else raw
        note = f", truncated at {FILE_CAP}" if len(words) > FILE_CAP else ""
        parts.append(f"\n\n=== {kind.upper()}: {p} ({len(words)} words{note}) ===\n\n{text}")
        manifest.append({"kind": kind, "path": str(p), "words": len(words), "truncated": bool(note)})
        if kind in PUBLIC_KINDS:
            public.append(parts[-1])
    summary = research_summary(run)
    if summary:
        public.append(f"\n\n=== RESEARCH SUMMARY: {run / RESEARCH} ===\n\n{summary}")
    truths = truth_corrections(run)
    if truths:
        parts.append(f"\n\n=== TRUTHS: client corrections from Brand Truth Review ===\n\n{truths}")
        manifest.append({"kind": "truths", "path": str(run.resolve().parent / WORKBOOK), "words": len(truths.split())})
    if not parts:
        sys.exit("nothing gathered: run crawl first, or pass --brain / --include")
    out = run / BUNDLE
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("# Input bundle\n\nEvery fact, number and named tool in faqs.md must come from this file."
                   + "".join(parts) + "\n", encoding="utf-8")
    (run / PUBLIC_BUNDLE).write_text("# Public bundle\n\nProvenance source: the brand's published site, approved "
                                     "--include files, the research Summary and the interview. Brain files are for "
                                     "voice only." + "".join(public) + "\n", encoding="utf-8")
    (run / DATA / "inputs.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    by = Counter(m["kind"] for m in manifest)
    print(f"Wrote {out}: {len(manifest)} files, {sum(m['words'] for m in manifest)} words "
          f"({', '.join(f'{k} {v}' for k, v in by.items())})")
    if "interview" not in by:
        print("  [info] no data/interview.md yet: save the gap-interview answers there, then gather again")
    return 0


# ------------------------------------------------------------- faqs.md io --
LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
META = re.compile(r"<!--(.*?)-->", re.S)


def parse_meta(raw: str) -> dict:
    meta = {}
    for part in raw.split(";"):
        key, sep, val = part.partition(":")
        if sep:
            meta[key.strip().lower()] = val.strip()
    srcs = re.split(r"[,\s]+", meta.get("sources", ""))
    meta["sources"] = [s for s in srcs if s]
    return meta


def parse_faqs(md: str) -> tuple[str, list[dict]]:
    """(title, [{n, question, answer (markdown), meta}])"""
    title, faqs, cur = "", [], None
    for line in md.splitlines():
        h = re.match(r"^###\s+(.+?)\s*$", line)
        if h:
            cur = {"n": len(faqs) + 1, "question": h.group(1).strip(), "lines": []}
            faqs.append(cur)
        elif re.match(r"^#\s+", line) and not faqs:
            title = line.lstrip("#").strip()
        elif cur is not None and not re.match(r"^#{1,2}\s", line):
            cur["lines"].append(line)
    for f in faqs:
        block = "\n".join(f.pop("lines"))
        m = META.search(block)
        f["meta"] = parse_meta(m.group(1)) if m else {"sources": []}
        f["has_meta"] = bool(m)
        f["answer"] = " ".join(META.sub(" ", block).split())
    for i, f in enumerate(faqs):
        f["meta"].setdefault("placement", "homepage" if i < HOMEPAGE_TOP else "faq-page")
    return title, faqs


def brand_from_title(title: str) -> str:
    return re.sub(r"\s+(FAQs?|Frequently Asked Questions)$", "", title, flags=re.I).strip()


def plain(md: str) -> str:
    return LINK.sub(lambda m: m.group(1), md)


def load_links_json(run: Path) -> list[dict]:
    p = run / LINKS_JSON
    if not p.is_file():
        return []
    data = json.loads(p.read_text(encoding="utf-8"))
    return data if isinstance(data, list) else data.get("links", [])


def links_for(faq: dict, extra: list[dict]) -> list[dict]:
    """Inline markdown links plus any data/links.json picks for this FAQ."""
    out = [{"anchor": m.group(1), "url": m.group(2), "via": "inline"} for m in LINK.finditer(faq["answer"])]
    for e in extra:
        if e.get("n") == faq["n"] or (e.get("question") or "").strip() == faq["question"]:
            out.append({"anchor": e.get("anchor", ""), "url": e.get("target") or e.get("url", ""), "via": "links.json"})
    seen, uniq = set(), []
    for link in out:  # an inline link repeated in data/links.json counts once
        key = (clean_url(link["url"]) if link["url"] else "", link["anchor"])
        if key not in seen:
            seen.add(key)
            uniq.append(link)
    return uniq


def linked_answer(faq: dict, extra: list[dict]) -> str:
    """The answer markdown with data/links.json picks applied (first verbatim occurrence)."""
    md = faq["answer"]
    for link in links_for(faq, extra):
        if link["via"] == "links.json" and link["anchor"] and link["anchor"] in plain(md) and not LINK.search(md):
            md = md.replace(link["anchor"], f"[{link['anchor']}]({link['url']})", 1)
    return md


# ------------------------------------------------------------------ lint --
STOP = set("""a an and are as at be been but by can for from has have how i if in into is it its of on or so
such than that the their them there these they this to was were what when where which who why will with
your you does do did not no yes about also more most other our we us any each own may might should could
across alongside including using used use make makes helps help""".split())
COMMON_CAPS = {"AI", "I", "FAQ", "FAQs", "URL", "URLs", "HTML", "API", "APIs", "PDF", "JSON", "CEO", "B2B",
               "B2C", "SaaS", "DIY", "OK"}
BAD_OPENERS = {"yes", "no", "it", "it's", "its", "we", "our", "you", "your", "they", "this", "absolutely", "sure"}
ABBREV = re.compile(r"\b(e\.g|i\.e|etc|vs|Inc|Ltd|Pty|Co|Dr|Mr|Mrs|Ms|St|No|approx)\.", re.I)
PROMISE = re.compile(r"\bguarantee[sd]?\b|\bguaranteed\b|\bwill (rank|double|triple|get|boost)\b|"
                     r"\bgets? (you |brands? |your brand )?cited\b|#1\b|\bnumber one\b", re.I)
SUPERLATIVE = re.compile(r"\b(best|leading|world[- ]class|premier|unrivalled|unrivaled|unmatched|most trusted|"
                         r"award[- ]winning|cheapest|fastest|biggest|greatest|ultimate)\b", re.I)
VOLATILE = re.compile(r"[$€£]\s?\d|\d[\d,.]*\s*[kKmM]?\+?\s*(members|subscribers|followers|customers|clients|students|"
                      r"users|people|downloads)\b|\b(19|20)\d{2}\b|\b(?-i:January|February|March|April|May|June|July|"
                      r"August|September|October|November|December)\b", re.I)
US_SPELL = re.compile(r"\b(optimiz\w*|organiz\w*|speciali[z]\w*|prioritiz\w*|summariz\w*|recogniz\w*|analyz\w*|"
                      r"customiz\w*|standardiz\w*|color\w*|behavior\w*|center(s|ed)?|favorite\w*)\b", re.I)
AU_SPELL = re.compile(r"\b(optimis\w*|organis\w*|specialis\w*|prioritis\w*|summaris\w*|recognis\w*|analys(e|es|ed|ing)\b|"
                      r"customis\w*|colour\w*|behaviour\w*|centre\w*|favourite\w*)\b", re.I)
NEGATED = re.compile(r"\b(not|never|no|rather than|without)\b", re.I)
BROKERS = re.compile(r"https?://(?:[\w-]+\.)*(?:zoominfo|signalhire|wiza|rocketreach|lusha|contactout)\.[a-z.]+[^\s)\"'<>]*|"
                     r"https?://(?:[\w-]+\.)*apollo\.io\b[^\s)\"'<>]*", re.I)
GENERIC_ANCHOR = re.compile(r"^(click here|here|learn more|read more|find out more|see more|more info|more|this page|"
                            r"this link|this|link|website|homepage|home|details|more details)$", re.I)
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿]")
NUM = re.compile(r"(?<![\w.])(\$?)(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)(\s?[kKmM](?![a-zA-Z]))?(%?)")
PLACEHOLDER = re.compile(r"\[VERIFY:[^\]]*\]|\bTODO\b|\bTBC\b")
TOKEN = re.compile(r"[^\W_][\w'’&+.-]*")


def num_values(text: str) -> set[float]:
    """Every number in text as a float, with 113.4k == 113400 and 1,000 == 1000."""
    vals = set()
    for m in NUM.finditer(text):
        try:
            v = float(m.group(2).replace(",", ""))
        except ValueError:
            continue
        vals.add(round(v, 4))
        suf = (m.group(3) or "").strip().lower()
        if suf in ("k", "m"):
            vals.add(round(v * (1000 if suf == "k" else 1_000_000), 4))
    return vals


def words_of(text: str) -> list[str]:
    return re.findall(r"[^\W_]+(?:'[^\W_]+)?", text.lower())


def content_set(text: str, drop: set[str]) -> set[str]:
    return {w for w in words_of(text) if len(w) > 3 and w not in STOP and w not in drop}


def names_words(names: list[str]) -> set[str]:
    return {w for n in names for w in words_of(n)}


class Bundle:
    def __init__(self, text: str):
        self.text = text
        self.lower = text.lower()
        self.nums = num_values(text)
        self._cache: dict[str, bool] = {}

    def has(self, phrase: str) -> bool:
        key = phrase.lower()
        if key not in self._cache:
            hits = re.finditer(re.escape(key), self.lower)
            self._cache[key] = any(self._starts_cleanly(m.start()) and self._ends_cleanly(m.end()) for m in hits)
        return self._cache[key]

    def _starts_cleanly(self, start: int) -> bool:
        if start == 0:
            return True
        prev = self.text[start - 1]
        return not (prev.isalnum() or prev == "_") or (prev.islower() and self.text[start].isupper())

    def _ends_cleanly(self, end: int) -> bool:
        # Scraped text often glues a heading to the next line ("RecordingsRaw and ..."):
        # a capital letter straight after the match counts as a word boundary.
        if end >= len(self.text):
            return True
        nxt = self.text[end]
        return not (nxt.isalnum() or nxt == "_") or (nxt.isupper() and self.text[end - 1].islower())


def strip_names(text: str, names: list[str]) -> str:
    for n in sorted(names, key=len, reverse=True):
        if n:
            text = re.sub(re.escape(n) + r"(['’]s)?", " , ", text)
    return text


def capital_part(tok: str) -> str:
    """'Claude-powered' -> 'Claude', 'AI-powered' -> '' (common), 'eBay' -> 'eBay', 'widget' -> ''."""
    parts = [x for x in tok.split("-") if x and (x[0].isupper() or any(c.isupper() for c in x[1:]))]
    parts = [x for x in parts if x not in COMMON_CAPS] if "-" in tok else parts
    return "-".join(parts)


def cap_phrases(text: str) -> list[str]:
    """Runs of capitalised tokens (proper nouns, acronyms, camel case), common edge tokens trimmed."""
    phrases, cur, prev_end = [], [], -1
    for m in TOKEN.finditer(text):
        tok = capital_part(re.sub(r"(['’]s|[.'’-])$", "", m.group(0)))
        capital = bool(tok)
        joined = cur and text[prev_end:m.start()] == " "
        if capital and (joined or not cur):
            cur.append(tok)
        elif capital:
            phrases.append(cur)
            cur = [tok]
        else:
            if cur:
                phrases.append(cur)
            cur = []
        prev_end = m.end()
    if cur:
        phrases.append(cur)
    out = []
    for p in phrases:
        while p and p[0] in COMMON_CAPS:
            p = p[1:]
        while p and p[-1] in COMMON_CAPS:
            p = p[:-1]
        if p:
            out.append(" ".join(p))
    return out


def check_provenance(faq: dict, text: str, names: list[str], bundle: Bundle, r: Report,
                     private: Bundle | None = None) -> int:
    """Numbers and capitalised names in the answer must appear in the public bundle."""
    tag = f"FAQ {faq['n']}"
    for m in NUM.finditer(text):
        if not num_values(m.group(0)) & bundle.nums:
            r.fails.append(f"{tag}: number not in the public bundle (invented?): {m.group(0).strip()!r}")
            if private and num_values(m.group(0)) & private.nums:
                r.warns.append(f"{tag}: {m.group(0).strip()!r} is only in private inputs (brain/month files)")
    rest = strip_names(text, names)
    starts_with_name = rest.lstrip().startswith(",")
    if not starts_with_name:
        rest = re.sub(r"^\s*\S+", " , ", rest)  # sentence-initial word is capitalised anyway
    missing = [p for p in dict.fromkeys(cap_phrases(rest)) if not bundle.has(p)]
    for p in missing:
        r.fails.append(f"{tag}: name {p!r} is not in the public bundle; cut it or add the source that says it")
        if private and private.has(p):
            r.warns.append(f"{tag}: {p!r} is only in private inputs (brain/month files); publish a source or cut it")
    return len(missing)


def one_sentence(text: str) -> bool:
    t = ABBREV.sub(lambda m: m.group(1), text.strip())
    return t.endswith(".") and not re.search(r"[.!?][\"”’)]?\s+\S", t)


def check_answer_shape(faq: dict, text: str, names: list[str], r: Report) -> None:
    tag = f"FAQ {faq['n']}"
    n = len(text.split())
    if not one_sentence(text):
        r.fails.append(f"{tag}: answer must be exactly one sentence ending in a full stop: {text[:80]!r}")
    if not 20 <= n <= 55:
        r.fails.append(f"{tag}: answer is {n} words; write 20 to 55")
    first = (text.split() or [""])[0].strip(",.").lower()
    if first in BAD_OPENERS:
        r.fails.append(f"{tag}: answer opens with {first.title()!r}; open with the brand or product name as subject")
    elif not any(text.startswith(nm) for nm in names if nm):
        r.fails.append(f"{tag}: answer must start with the brand or a product name ({', '.join(names)})")


def check_language(faq: dict, text: str, opts: dict, r: Report) -> None:
    tag = f"FAQ {faq['n']}"
    for rx, what in ((PROMISE, "promise / outcome language"), (SUPERLATIVE, "superlative")):
        m = rx.search(text)
        if m and rx is PROMISE and NEGATED.search(text[max(0, m.start() - 25):m.start()]):
            m = None  # "not a guarantee of", "rather than guarantees": the recommended non-guarantee line
        if m:
            r.fails.append(f"{tag}: {what} {m.group(0)!r}; describe capability, never promise results")
    if "—" in text or "—" in faq["question"]:
        r.fails.append(f"{tag}: em dash; use a comma, colon or full stop")
    vm = VOLATILE.search(text)
    if vm and not opts["allow_volatile"]:
        r.fails.append(f"{tag}: volatile fact {vm.group(0)!r} (price, count or date); leave it out or pass --allow-volatile")
    for term in opts["forbid"]:
        if term and re.search(re.escape(term), text + " " + faq["question"], re.I):
            r.fails.append(f"{tag}: never-say term present: {term!r}")
    if EMOJI.search(text):
        r.warns.append(f"{tag}: emoji in the answer")
    check_spelling(faq, text, opts["locale"], r)


def check_spelling(faq: dict, text: str, locale: str, r: Report) -> None:
    loc = locale.lower()
    wrong = US_SPELL if loc in ("en-au", "en-gb", "en-nz", "en-ie") else AU_SPELL if loc == "en-us" else None
    if not wrong:
        return
    hits = sorted({m.group(0) for m in wrong.finditer(text) if not m.group(0)[0].isupper() or m.start() == 0})
    if hits:
        r.fails.append(f"FAQ {faq['n']}: spelling does not match {locale}: {hits}")


def host_of(url: str) -> str:
    h = urllib.parse.urlsplit(url).netloc.lower()
    return h[4:] if h.startswith("www.") else h


def brand_hosts(inv: dict) -> set[str]:
    return {host_of(u) for u in inv}


def is_brand(url: str, hosts: set[str]) -> bool:
    h = host_of(url)
    return any(h == b or h.endswith("." + b) for b in hosts)


def check_evidence(tag: str, meta: dict, hosts: set[str], r: Report) -> None:
    srcs = meta.get("sources", [])
    for s in srcs:
        if BROKERS.search(s):
            r.fails.append(f"{tag}: source is a people-data broker, never cite it: {s}")
    if meta.get("evidence") in ("third-party", "mixed") and hosts and not [s for s in srcs if not is_brand(s, hosts)]:
        r.fails.append(f"{tag}: evidence {meta['evidence']} but every source is on the brand's own domain; "
                       "cite the third-party URL or mark it first-party")


def check_meta(faq: dict, final: bool, r: Report, hosts: set[str] = frozenset()) -> None:
    tag, meta = f"FAQ {faq['n']}", faq["meta"]
    if not faq["has_meta"]:
        r.fails.append(f"{tag}: no metadata comment (<!-- bucket: ...; sources: ...; evidence: ... -->)")
        return
    if meta.get("bucket") not in BUCKETS:
        r.fails.append(f"{tag}: bucket {meta.get('bucket')!r} is not one of {sorted(BUCKETS)}")
    if meta.get("placement") not in PLACEMENTS:
        r.fails.append(f"{tag}: placement {meta.get('placement')!r} is not homepage or faq-page")
    if not [s for s in meta["sources"] if s.startswith(("http://", "https://"))]:
        r.fails.append(f"{tag}: no source URL; every FAQ cites the page(s) it came from")
    if meta.get("evidence") not in EVIDENCE:
        r.fails.append(f"{tag}: evidence {meta.get('evidence')!r} is not first-party, third-party or mixed")
    check_evidence(tag, meta, hosts, r)
    for ph in PLACEHOLDER.findall(faq["answer"] + faq["question"]):
        (r.fails if final else r.notes).append(f"{tag}: placeholder left: {ph}")


def check_questions(faqs: list[dict], names: list[str], r: Report, cap: int = 20) -> None:
    if not 12 <= len(faqs) <= cap:
        r.fails.append(f"{len(faqs)} FAQs; write 12 to {cap}")
    drop = names_words(names)
    sets = [content_set(f["question"], drop) for f in faqs]
    for i, f in enumerate(faqs):
        if not f["question"].endswith("?"):
            r.fails.append(f"FAQ {f['n']}: question must end with '?': {f['question']!r}")
        for j in range(i):
            a, b = sets[i], sets[j]
            same = f["question"].lower() == faqs[j]["question"].lower()
            if same or (a and b and len(a & b) / len(a | b) >= 0.8):
                r.fails.append(f"FAQ {f['n']}: duplicate or near-duplicate of FAQ {faqs[j]['n']}: {f['question']!r}")


def check_coverage(faqs: list[dict], r: Report) -> None:
    buckets = Counter(f["meta"].get("bucket") for f in faqs)
    for name, (lo, hi) in BUCKET_LIMITS.items():
        if not lo <= buckets[name] <= hi:
            r.fails.append(f"bucket {name}: {buckets[name]} FAQ(s); allowed {lo} to {hi} (references/question-bank.md)")
    home = sum(f["meta"].get("placement") == "homepage" for f in faqs)
    if home > HOMEPAGE_TOP:
        r.warns.append(f"{home} FAQs placed on the homepage; keep the top {HOMEPAGE_TOP} there, the rest on an FAQ page")
    r.notes.append("buckets: " + ", ".join(f"{k} {v}" for k, v in sorted(buckets.items(), key=lambda kv: str(kv[0]))))


def check_overlap(faqs: list[dict], texts: list[str], names: list[str], r: Report) -> None:
    drop = names_words(names)
    sets = [content_set(t, drop) for t in texts]
    for i in range(len(faqs)):
        for j in range(i):
            a, b = sets[i], sets[j]
            if a and b and len(a & b) / min(len(a), len(b)) > 0.6:
                r.warns.append(f"FAQ {faqs[i]['n']} repeats FAQ {faqs[j]['n']} (over 60% shared terms); merge or differentiate")


# ------------------------------------------------------------ lint: links --

def check_target(tag: str, url: str, inv: dict, r: Report) -> None:
    row = inv.get(clean_url(url))
    if row is None:
        r.fails.append(f"{tag}: link target not in link-inventory.csv (crawl it or pick another): {url}")
        return
    if str(row["status"]) != "200" or clean_url(row.get("final_url") or url) != clean_url(url):
        r.fails.append(f"{tag}: link target redirects or errors (status {row['status']} -> {row.get('final_url')}): {url}")
    elif row.get("indexable") != "yes":
        r.fails.append(f"{tag}: link target is not indexable (noindex or canonicalised elsewhere): {url}")
    if row.get("canonical") and clean_url(row["canonical"]) != clean_url(url):
        r.fails.append(f"{tag}: link target is not self-canonical (canonical {row['canonical']}): {url}")
    if row.get("page_type") == "homepage" or is_home(url):
        r.fails.append(f"{tag}: links to the homepage, where the FAQs live: {url}")


def check_anchor(tag: str, anchor: str, text: str, brand: str, r: Report) -> None:
    if anchor not in text:
        r.fails.append(f"{tag}: anchor {anchor!r} is not verbatim in the answer; never reword an answer to fit a link")
    n = len(anchor.split())
    if not 2 <= n <= 6:
        r.fails.append(f"{tag}: anchor {anchor!r} is {n} word(s); use 2 to 6 descriptive words")
    if GENERIC_ANCHOR.match(anchor.strip()):
        r.fails.append(f"{tag}: generic anchor {anchor!r}; use words that say what the page is")
    if brand and re.sub(r"['’]s$", "", anchor.strip()).lower() == brand.lower():
        r.fails.append(f"{tag}: anchor is the bare brand name; use words that describe the target page")


def check_links(faqs: list[dict], texts: list[str], ctx: dict, r: Report) -> None:
    inv, extra, brand = ctx["inventory"], ctx["links"], ctx["brand"]
    used, linked = Counter(), 0
    for faq, text in zip(faqs, texts):
        tag, links = f"FAQ {faq['n']}", links_for(faq, extra)
        if len({l_["url"] for l_ in links}) > 1 or len([l_ for l_ in links if l_["via"] == "inline"]) > 1:
            r.fails.append(f"{tag}: {len(links)} links; at most one per answer")
        for link in links:
            if inv:
                check_target(tag, link["url"], inv, r)
            else:
                r.fails.append(f"{tag}: has a link but data/link-inventory.csv is missing; run crawl")
            check_anchor(tag, link["anchor"], text, brand, r)
            used[clean_url(link["url"])] += 1
        linked += bool(links)
    for url, n in used.items():
        if n > 2:
            r.fails.append(f"link target used {n} times (max 2): {url}")
    linkable = any(v.get("indexable") == "yes" and v.get("page_type") != "homepage" for v in inv.values())
    if faqs and (linkable or ctx["has_candidates"]) and linked / len(faqs) < 0.5:
        r.warns.append(f"only {linked} of {len(faqs)} answers carry an internal link; aim for at least half")
    r.notes.append(f"links: {linked} of {len(faqs)} answers linked, {len(used)} distinct targets")


def lint_faqs(md: str, bundle_text: str, opts: dict) -> Report:
    """opts: brand, products, locale, forbid, allow_volatile, final, inventory, links, has_candidates."""
    r = Report()
    title, faqs = parse_faqs(md)
    brand = opts.get("brand") or brand_from_title(title)
    if not faqs:
        r.fails.append("no FAQs found: one '### Question?' heading per FAQ")
        return r
    opts = {**opts, "brand": brand}
    names = [n for n in [brand, *opts.get("products", [])] if n]
    bundle = Bundle(bundle_text)
    private = Bundle(opts["private"]) if opts.get("private") else None
    texts = [plain(linked_answer(f, opts["links"])) for f in faqs]
    check_questions(faqs, names, r, opts.get("count") or 20)
    orphans = 0
    for faq, text in zip(faqs, texts):
        check_meta(faq, opts["final"], r, brand_hosts(opts["inventory"]))
        check_answer_shape(faq, text, names, r)
        check_language(faq, text, opts, r)
        orphans += check_provenance(faq, text, names, bundle, r, private)
    check_coverage(faqs, r)
    check_overlap(faqs, texts, names, r)
    check_links(faqs, texts, opts, r)
    r.notes.append(f"provenance: {orphans} unsupported name(s) across {len(faqs)} answers")
    return r


def split_list(values: list[str] | None) -> list[str]:
    return [x.strip() for v in values or [] for x in v.split(",") if x.strip()]


def cmd_lint(args) -> int:
    run = Path(args.run_dir)
    page = Path(args.faqs) if args.faqs else run / FAQS_MD
    bundle_p = Path(args.bundle) if args.bundle else run / PUBLIC_BUNDLE
    for p, hint in ((page, "write faqs.md first"), (bundle_p, "run gather first")):
        if not p.is_file():
            sys.exit(f"missing {p}: {hint}")
    opts = {"brand": args.brand, "products": split_list(args.products), "locale": args.locale,
            "forbid": split_list(args.forbid), "allow_volatile": args.allow_volatile, "final": args.final,
            "inventory": load_inventory(run), "links": load_links_json(run),
            "has_candidates": has_candidates(run), "count": args.count,
            "private": (run / BUNDLE).read_text(encoding="utf-8", errors="replace") if (run / BUNDLE).is_file() else ""}
    r = lint_faqs(page.read_text(encoding="utf-8"), bundle_p.read_text(encoding="utf-8", errors="replace"), opts)
    r.fails += [f"{name}: never-say term {term!r} in a rendered file" for name, term in rendered_leaks(run, opts["forbid"])]
    verdict = "FAIL" if r.fails else "PASS"
    lines = [f"# Lint: {page}", "", f"Checked {time.strftime('%Y-%m-%d %H:%M')}", ""]
    lines += [f"- FAIL {f}" for f in r.fails] + [f"- WARN {w}" for w in r.warns] + [f"- NOTE {n}" for n in r.notes]
    lines += ["", f"**Result: {verdict}** ({len(r.fails)} fail, {len(r.warns)} warn)"]
    print("\n".join(lines[4:]))
    out = run / DATA / "lint.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return 1 if r.fails else 0


RENDERED = [FAQS_HTML, "faq-page.html", JSONLD, HANDOVER, NOTE]


def rendered_leaks(run: Path, forbid: list[str]) -> list[tuple[str, str]]:
    """(file, term) for every never-say term or data-broker URL found in a public rendered file."""
    out = []
    for name in RENDERED:
        p = run / name
        if not p.is_file():
            continue
        text = html.unescape(p.read_text(encoding="utf-8", errors="replace"))
        out += [(name, t) for t in forbid if t and re.search(re.escape(t), text, re.I)]
        out += [(name, m.group(0)) for m in BROKERS.finditer(text)]
    return out


def strip_brokers(text: str) -> str:
    text = re.sub(r'<a [^>]*href="[^"]*"[^>]*>(.*?)</a>',
                  lambda m: m.group(1) if BROKERS.search(m.group(0)) else m.group(0), text)
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)",
                  lambda m: m.group(1) if BROKERS.search(m.group(2)) else m.group(0), text)
    return BROKERS.sub("[link removed]", text)


def has_candidates(run: Path) -> bool:
    p = run / CANDIDATES
    if not p.is_file():
        return False
    return any(e.get("candidates") for e in json.loads(p.read_text(encoding="utf-8")))


# ----------------------------------------------------------------- links --

def tokens(text: str) -> list[str]:
    return [w for w in words_of(text) if len(w) > 1 and w not in STOP]


def page_doc(run: Path, row: dict) -> list[str]:
    body = ""
    p = run / DATA / "pages" / f"{page_slug(row['url'])}.md"
    if p.is_file():
        body = " ".join(p.read_text(encoding="utf-8", errors="replace").split()[:3000])
    head = " ".join([row.get("title", ""), row.get("h1", ""), row.get("meta_description", "")])
    return tokens(head) * 2 + tokens(body)


class BM25:
    def __init__(self, docs: list[list[str]], k1: float = 1.5, b: float = 0.75):
        self.docs, self.k1, self.b = [Counter(d) for d in docs], k1, b
        self.lens = [len(d) for d in docs]
        self.avg = (sum(self.lens) / len(docs)) if docs else 0
        df = Counter(t for d in self.docs for t in d)
        n = len(docs)
        self.idf = {t: math.log(1 + (n - f + 0.5) / (f + 0.5)) for t, f in df.items()}

    def score(self, query: list[str], i: int) -> float:
        d, s = self.docs[i], 0.0
        norm = self.k1 * (1 - self.b + self.b * self.lens[i] / (self.avg or 1))
        for t in set(query):
            tf = d.get(t, 0)
            if tf:
                s += self.idf[t] * tf * (self.k1 + 1) / (tf + norm)
        return s


def link_targets(inv: dict) -> list[dict]:
    return [row for url, row in inv.items() if row.get("indexable") == "yes" and str(row.get("status")) == "200"
            and row.get("page_type") not in ("homepage", "legal") and not is_home(url)]


def shortlist(faqs: list[dict], rows: list[dict], docs: list[list[str]], drop: set[str], k: int = 3) -> list[dict]:
    bm = BM25(docs)
    out = []
    for f in faqs:
        q = [t for t in tokens(plain(f["answer"]) + " " + f["question"]) if t not in drop]
        scored = sorted(((bm.score(q, i), i) for i in range(len(rows))), key=lambda x: -x[0])
        cands = [{"url": rows[i]["url"], "title": rows[i].get("title", ""), "score": round(s, 3)}
                 for s, i in scored[:k] if s > 0]
        out.append({"n": f["n"], "question": f["question"], "candidates": cands})
    return out


def cmd_links(args) -> int:
    run = Path(args.run_dir)
    inv = load_inventory(run)
    if not inv:
        sys.exit(f"missing {run / INVENTORY}: run crawl first")
    title, faqs = parse_faqs((run / FAQS_MD).read_text(encoding="utf-8"))
    brand = args.brand or brand_from_title(title)
    rows = link_targets(inv)
    drop = names_words([brand, *split_list(args.products)])
    result = shortlist(faqs, rows, [page_doc(run, r) for r in rows], drop)
    out = run / CANDIDATES
    out.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    with_c = sum(bool(e["candidates"]) for e in result)
    print(f"Wrote {out}: {len(result)} answers, {with_c} with candidates, from {len(rows)} linkable pages")
    return 0


# ---------------------------------------------------------------- render --

def esc(s: str) -> str:
    return html.escape(s, quote=True)


def md_inline_html(md: str, quote: bool = True) -> str:
    """Markdown links to <a href>. quote=False (JSON-LD) escapes only <, > and &, never quotes."""
    out, pos = [], 0
    for m in LINK.finditer(md):
        out.append(html.escape(md[pos:m.start()], quote=quote))
        href = html.escape(m.group(2), quote=False).replace('"', "%22")
        out.append(f'<a href="{href}">{html.escape(m.group(1), quote=quote)}</a>')
        pos = m.end()
    out.append(html.escape(md[pos:], quote=quote))
    return "".join(out)


def faqs_html(items: list[tuple[str, str]], heading: str) -> str:
    rows = [f'<details class="faq-item">\n  <summary>{esc(q)}</summary>\n  <p>{md_inline_html(a)}</p>\n</details>'
            for q, a in items]
    style = ("<style>.faq details{border-bottom:1px solid #ddd;padding:.75rem 0}"
             ".faq summary{cursor:pointer;font-weight:600}.faq p{margin:.5rem 0 0}</style>")
    return (f'<section class="faq" id="faq" aria-labelledby="faq-heading">\n{style}\n'
            f'<h2 id="faq-heading">{esc(heading)}</h2>\n' + "\n".join(rows) + "\n</section>\n")


def faq_jsonld(items: list[tuple[str, str]], url: str) -> dict:
    url = clean_url(url)  # https://x.com -> https://x.com/, matching the canonical
    return {"@context": "https://schema.org", "@type": "FAQPage", "@id": f"{url}#faq", "url": url,
            "mainEntity": [{"@type": "Question", "name": q,
                            "acceptedAnswer": {"@type": "Answer", "text": md_inline_html(a, quote=False)}}
                           for q, a in items]}


HANDOVER_TMPL = """# Homepage FAQs: handover for {brand}

Generated {date} by mos-geo-homepage-faqs. Every answer passed `faqs.py lint`: one sentence,
brand-first, every number and named tool traced to the brand's own inputs, every link checked
against the crawl.

## Files

- `faqs.html`: the homepage FAQ block ({n_home} questions) as a `<details>` accordion. No scripts or
  external assets; the small inline `<style>` can be deleted if the theme styles it.
- `faqpage.jsonld`: `FAQPage` schema for the same {n_home} questions. Paste inside
  `<script type="application/ld+json">` on the homepage only.
{faq_page_line}- `faqs.md`: the source copy. `data/lint.md`: the gate result. `research-note.md`: sources and claims policy.

## Where it goes on the homepage

- Below the main content, above the footer: the section a reader reaches after the offer is clear.
- The text must be in the HTML the server sends (not loaded by JavaScript after the page renders),
  so crawlers and AI engines that do not run scripts still read it.
- Keep the answers visible on click (a `<details>` accordion is fine); do not hide them behind a
  script-only tab.

## Homepage vs FAQ page

The top {top} questions (placement `homepage`) go on the homepage. The rest (placement `faq-page`)
belong on a dedicated FAQ page or the AI Info Page, where a longer list does not push the offer down.

## Internal links

Each linked answer points to the page that backs it up, with an anchor copied from the answer.
Keep the links when pasting. If a target page moves, update the link or remove it: never leave a
link to a redirect.

## Schema: what to expect

`FAQPage` markup helps machines read the question/answer pairs. It does **not** earn a rich
result: Google removed FAQ rich results from Search entirely on 7 May 2026 (they had been
limited to government and health sites since August 2023). Ship it for machine readability, not for a SERP feature, and never
mark up questions that are not visible on the page.

## Before publishing

- Client approves every row on the "{tab}" tab of the month's `{workbook}`.
- Check each link opens a live page (not a redirect).
"""


def placement_split(faqs: list[dict], extra: list[dict]) -> tuple[list, list]:
    home = [(f["question"], linked_answer(f, extra)) for f in faqs if f["meta"].get("placement") == "homepage"]
    rest = [(f["question"], linked_answer(f, extra)) for f in faqs if f["meta"].get("placement") != "homepage"]
    return home, rest


def research_note(run: Path, brand: str, faqs: list[dict]) -> str:
    body = research_summary(run) or (
        "_The off-site research is kept as private working notes in `data/`; no public `## Summary` "
        "section was written in data/research.md._")
    sources = sorted({s for f in faqs for s in f["meta"].get("sources", []) if not BROKERS.search(s)})
    evidence = Counter(f["meta"].get("evidence", "unknown") for f in faqs)
    lines = [f"# Research note: {brand} homepage FAQs", "", "## Research and disambiguation", "", body, "",
             "## Evidence", "", f"{len(faqs)} FAQs: " + ", ".join(f"{k} {v}" for k, v in sorted(evidence.items())),
             "", "Sources cited:", ""] + [f"- {s}" for s in sources]
    lines += ["", "## Claims policy", "",
              "Answers describe capabilities (teaches, helps, shows, focuses on), never outcomes. No prices, "
              "counts or dates, no superlatives, no guarantees. Every number and named tool is in the "
              "public bundle (`data/public-bundle.md`)."]
    return "\n".join(lines) + "\n"


def write_public(path: Path, text: str) -> None:
    path.write_text(strip_brokers(text), encoding="utf-8")


def cmd_render(args) -> int:
    run = Path(args.run_dir)
    title, faqs = parse_faqs((run / FAQS_MD).read_text(encoding="utf-8"))
    brand = args.brand or brand_from_title(title) or "the brand"
    extra = load_links_json(run)
    home, rest = placement_split(faqs, extra)
    heading = args.heading or f"{brand}: frequently asked questions"
    write_public(run / FAQS_HTML, faqs_html(home, heading))
    write_public(run / JSONLD, json.dumps(faq_jsonld(home, args.url), indent=2, ensure_ascii=False) + "\n")
    line = ""
    if rest:
        write_public(run / "faq-page.html", faqs_html(rest, f"More about {brand}"))
        line = (f"- `faq-page.html`: the other {len(rest)} questions for an FAQ page (give that page its own "
                "`FAQPage` schema if you publish them there).\n")
    elif (run / "faq-page.html").is_file():
        (run / "faq-page.html").unlink()
    write_public(run / HANDOVER, HANDOVER_TMPL.format(
        brand=brand, date=time.strftime("%d %B %Y").lstrip("0"), n_home=len(home), faq_page_line=line,
        top=HOMEPAGE_TOP, tab=TAB, workbook=WORKBOOK))
    write_public(run / NOTE, research_note(run, brand, faqs))
    print(f"Wrote {FAQS_HTML} ({len(home)} homepage FAQs), {JSONLD}, {HANDOVER}, {NOTE}"
          + (f", faq-page.html ({len(rest)})" if rest else "") + f" in {run}")
    leaks = rendered_leaks(run, split_list(args.forbid))
    for name, term in leaks:
        print(f"FAIL {name}: never-say term {term!r} in a public file; fix the source and render again")
    return 1 if leaks else 0


# -------------------------------------------------------------- workbook --
TAB_HEADERS = ["#", "Placement", "Question", "Answer", "Bucket", "Source URL(s)", "Evidence", "Link target",
               "Anchor", "Client verdict", "Edited version", "Notes"]
TAB_ORDER = ["Checklist", "Brand Truth Review", "Brand 360 Report", "AI Visibility", "AI Info Page", TAB,
             "Initiatives"]


def open_book(run: Path):
    import openpyxl
    book = run / WORKBOOK if (run / WORKBOOK).is_file() else run.parent / WORKBOOK
    if not book.is_file():
        if not PACK_TEMPLATE.is_file():
            sys.exit(f"no {WORKBOOK} in {run.parent} and no pack template at {PACK_TEMPLATE}")
        book.write_bytes(PACK_TEMPLATE.read_bytes())
    return book, openpyxl.load_workbook(book)


def tick_checklist(wb, run: Path, day: str) -> None:
    ws = wb["Checklist"]
    row = next((r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == SKILL_ID), None)
    if row is None:
        row = next(r for r in range(8, ws.max_row + 2) if not ws.cell(row=r, column=3).value)
        ws.cell(row=row, column=2, value="Brand Optimisation")
        ws.cell(row=row, column=3, value=SKILL_ID)
        print(f"  [warn] Checklist had no {SKILL_ID} row; added one (rebuild the template to make it permanent)")
    if ws.cell(row=row, column=8).value in (None, "", "Not started"):
        ws.cell(row=row, column=8, value="Client review")  # never overwrite a status someone has moved on
    for col, v in {9: "☑", 10: day, 11: f"See the '{TAB}' tab; files in {run.name}/"}.items():
        ws.cell(row=row, column=col, value=v)


def kept_verdicts(wb) -> dict:
    """Client columns from the old tab, keyed by question text and, as a fallback, by row #."""
    if TAB not in wb.sheetnames:
        return {}
    old, kept = wb[TAB], {}
    for r in range(1, old.max_row + 1):
        q, n = old.cell(row=r, column=4).value, old.cell(row=r, column=2).value
        if q and q != "Question":
            vals = [old.cell(row=r, column=c).value for c in (11, 12, 13)]
            kept[str(q).strip()] = vals
            if isinstance(n, int):
                kept.setdefault(("n", n), vals)
    del wb[TAB]
    return kept


def faq_row(f: dict, extra: list[dict]) -> list:
    links = links_for(f, extra)
    link = links[0] if links else {"url": None, "anchor": None}
    return [f["n"], "Homepage" if f["meta"].get("placement") == "homepage" else "FAQ page", f["question"],
            plain(linked_answer(f, extra)), f["meta"].get("bucket"), "\n".join(f["meta"].get("sources", [])),
            f["meta"].get("evidence"), link["url"], link["anchor"]]


def write_tab(wb, faqs: list[dict], extra: list[dict], brand: str, run: Path, kept: dict) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    ws = wb.create_sheet(TAB)
    ws.sheet_view.showGridLines = False
    widths = [2, 5, 11, 40, 60, 14, 40, 13, 40, 24, 14, 45, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws["B2"] = f"Homepage FAQs: approve every answer before {brand} publishes them"
    ws["B2"].font = Font(bold=True, size=16, color="1F2937")
    ws["B3"] = f"In {run.name}/: {FAQS_HTML} (paste-ready), {JSONLD} (schema), {HANDOVER} (where it goes)."
    ws["B4"] = "Client: pick a verdict per row. Edit = write the new version; Remove = it is cut before publishing."
    top = 6
    wrap, edge = Alignment(wrap_text=True, vertical="top"), Border(*(Side(style="thin", color="D1D5DB"),) * 4)
    for i, h in enumerate(TAB_HEADERS, start=2):
        c = ws.cell(row=top, column=i, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2937")
        c.alignment, c.border = wrap, edge
    ws.freeze_panes = ws.cell(row=top + 1, column=2)
    for n, f in enumerate(faqs, start=top + 1):
        vals = faq_row(f, extra) + kept.get(f["question"], kept.get(("n", f["n"]), [None, None, None]))
        for i, v in enumerate(vals, start=2):
            c = ws.cell(row=n, column=i, value=v)
            c.alignment, c.border = wrap, edge
    last = top + max(len(faqs), 1)
    dv = DataValidation(type="list", formula1='"Approve,Edit,Remove"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K{top + 1}:K{last}")
    ws.auto_filter.ref = f"B{top}:M{last}"


def add_initiative(wb, brand: str, run: Path) -> None:
    ws = wb["Initiatives"]
    task = f"Publish the homepage FAQs for {brand}"
    if any(ws.cell(row=r, column=4).value == task for r in range(1, ws.max_row + 1)):
        return
    r = next(r for r in range(6, ws.max_row + 2) if not ws.cell(row=r, column=4).value)
    vals = {3: "Brand Optimisation", 4: task,
            5: f"Client approves the '{TAB}' tab, then follow {run.name}/{HANDOVER}: paste {FAQS_HTML} into the "
               f"homepage (server-rendered), add {JSONLD}, keep the internal links.",
            6: SKILL_ID, 7: "Scheduled", 8: "Yes", 10: 5, 11: 6, 12: 2, 13: f"{run.name}/{FAQS_HTML}"}
    for col, v in vals.items():
        ws.cell(row=r, column=col, value=v)


def cmd_workbook(args) -> int:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.exit("openpyxl is needed: uv run --with openpyxl python faqs.py workbook ...")
    run = Path(args.run_dir).resolve()
    if not (run / FAQS_MD).is_file():
        sys.exit(f"missing {run / FAQS_MD}")
    title, faqs = parse_faqs((run / FAQS_MD).read_text(encoding="utf-8"))
    brand = args.brand or brand_from_title(title) or "the brand"
    book, wb = open_book(run)
    tick_checklist(wb, run, time.strftime("%Y-%m-%d"))
    kept = kept_verdicts(wb)
    write_tab(wb, faqs, load_links_json(run), brand, run, kept)
    if "Initiatives" in wb.sheetnames:
        add_initiative(wb, brand, run)
    wb._sheets = [wb[n] for n in TAB_ORDER if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in TAB_ORDER]
    wb.save(book)
    print(f"Wrote {book}: Checklist ticked, {len(faqs)} FAQs on '{TAB}' ({sum(any(v) for k, v in kept.items() if isinstance(k, str))} earlier verdict row(s) kept), "
          "publish task on Initiatives")
    return 0


# ------------------------------------------------------------------ main --

def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("path", help="print the run folder")
    p.add_argument("--brand", required=True)
    p.add_argument("--date", help="YYYY-MM-DD (default today)")
    p.add_argument("--start", help="where to look for a MarketingOS brain (default .)")
    p.set_defaults(fn=cmd_path)
    sub.add_parser("doctor", help="preflight: Scrapling, uv, openpyxl, sibling crawler").set_defaults(fn=cmd_doctor)
    p = sub.add_parser("crawl", help="crawl the brand's own site into pages + link inventory")
    p.add_argument("--url", required=True)
    p.add_argument("--run-dir", required=True)
    p.add_argument("--max", type=int, default=50, help="max URLs (default 50)")
    p.add_argument("--delay", type=float, default=0.3)
    p.add_argument("--fetcher", choices=["auto", "urllib"], default="auto",
                   help="auto: Scrapling CLI (get > fetch > stealthy-fetch) when installed, else urllib")
    p.set_defaults(fn=cmd_crawl)
    p = sub.add_parser("gather", help="build data/input-bundle.md")
    p.add_argument("--run-dir", required=True)
    p.add_argument("--brain", help="MarketingOS brain root (identity and voice files)")
    p.add_argument("--include", action="append",
                   help="extra public file or folder of .md (repeatable). Drafts count as public only if "
                        "they are published or client-approved")
    p.set_defaults(fn=cmd_gather)
    p = sub.add_parser("links", help="BM25 shortlist of internal link targets per answer")
    p.add_argument("--run-dir", required=True)
    p.add_argument("--brand")
    p.add_argument("--products", action="append", help='"A,B" product names (repeatable)')
    p.set_defaults(fn=cmd_links)
    p = sub.add_parser("lint", help="the deterministic gate")
    p.add_argument("--run-dir", required=True)
    p.add_argument("--brand", required=True)
    p.add_argument("--products", action="append", help='"A,B" product names allowed as answer subjects')
    p.add_argument("--locale", default="en-AU")
    p.add_argument("--forbid", action="append", help="never-say term(s), comma-separated (repeatable)")
    p.add_argument("--allow-volatile", action="store_true", help="allow prices, counts and dates")
    p.add_argument("--final", action="store_true", help="also fail on [VERIFY]/TODO placeholders")
    p.add_argument("--count", type=int, default=20, help="max FAQs (default 20; min is always 12)")
    p.add_argument("--faqs", help="lint this file instead of <run>/faqs.md")
    p.add_argument("--bundle", help="provenance bundle (default <run>/data/public-bundle.md)")
    p.set_defaults(fn=cmd_lint)
    p = sub.add_parser("render", help="faqs.html + faqpage.jsonld + handover.md + research-note.md")
    p.add_argument("--run-dir", required=True)
    p.add_argument("--url", required=True, help="the homepage URL the FAQs will live on")
    p.add_argument("--brand")
    p.add_argument("--heading", help="section heading (default '{Brand}: frequently asked questions')")
    p.add_argument("--forbid", action="append", help="never-say term(s); any found in a rendered file fails")
    p.set_defaults(fn=cmd_render)
    p = sub.add_parser("workbook", help="tick Checklist, add the Homepage FAQs tab, add an Initiative")
    p.add_argument("--run-dir", required=True)
    p.add_argument("--brand")
    p.set_defaults(fn=cmd_workbook)
    return ap


def main() -> int:
    args = build_parser().parse_args()
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(main())
