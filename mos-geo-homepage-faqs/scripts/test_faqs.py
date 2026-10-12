"""Offline tests for faqs.py. No network.

Run from the repo root:
    uv run --with openpyxl python -m unittest mos-geo-homepage-faqs/scripts/test_faqs.py
"""

import csv
import json
import re
import sys
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import faqs as f  # noqa: E402

BRAND, PRODUCT = "Acme Widgets", "WidgetOS"
HOME = "https://acme.example/"
OS_PAGE = "https://acme.example/widgetos/"
CUSTOM = "https://acme.example/services/custom-widgets/"
ABOUT = "https://acme.example/about/"
OLD = "https://acme.example/old-page"
BUNDLE = """# Input bundle
Acme Widgets is a Geelong widget maker. WidgetOS is its production planning software and connects to Shopify
and Xero. Acme Widgets designs custom widgets, runs a factory tour, and publishes maintenance guides.
Acme Widgets serves small manufacturers, hardware retailers and engineering teams. Founded by Jane Smith.
"""

INVENTORY = [
    {"url": HOME, "status": 200, "final_url": HOME, "canonical": HOME, "indexable": "yes", "title": "Acme Widgets",
     "h1": "Widgets made in Geelong", "meta_description": "Custom widgets", "page_type": "homepage", "word_count": 400},
    {"url": OS_PAGE, "status": 200, "final_url": OS_PAGE, "canonical": "", "indexable": "yes",
     "title": "WidgetOS production planning software", "h1": "WidgetOS", "meta_description":
     "Plan production scheduling and stock in one place", "page_type": "product", "word_count": 800},
    {"url": CUSTOM, "status": 200, "final_url": CUSTOM, "canonical": CUSTOM, "indexable": "yes",
     "title": "Custom widget design services", "h1": "Custom widgets", "meta_description":
     "Bespoke widget design and prototyping", "page_type": "service", "word_count": 600},
    {"url": ABOUT, "status": 200, "final_url": ABOUT, "canonical": ABOUT, "indexable": "yes",
     "title": "About Acme Widgets", "h1": "Our Geelong factory", "meta_description": "Factory tour and team",
     "page_type": "about", "word_count": 500},
    {"url": OLD, "status": 301, "final_url": OS_PAGE, "canonical": "", "indexable": "no", "title": "",
     "h1": "", "meta_description": "", "page_type": "page", "word_count": 0},
]

# (question, answer markdown, bucket)
GOOD = [
    ("What is Acme Widgets?", "Acme Widgets is a Geelong widget maker that designs, builds and supports "
     "industrial widgets for small manufacturers, with a focus on durable parts, clear documentation and practical "
     "maintenance advice.", "identity"),
    ("What is WidgetOS?", f"WidgetOS is the Acme Widgets [production planning software]({OS_PAGE}) that helps "
     "teams schedule jobs, track stock and keep orders moving without juggling separate spreadsheets.", "product"),
    ("Does WidgetOS connect to Shopify and Xero?", "WidgetOS connects to Shopify and Xero so that online orders, "
     "invoices and stock levels stay in step with the factory schedule instead of being copied across by hand.",
     "capability"),
    ("Does Acme Widgets design custom widgets?", f"Acme Widgets offers [custom widget design services]({CUSTOM}) "
     "covering sketches, prototypes, material choices and small production runs for teams that cannot find a "
     "suitable part off the shelf.", "capability"),
    ("Can Acme Widgets help me plan production?", f"Acme Widgets helps manufacturers plan production through "
     f"WidgetOS, which lays out [jobs, machines and deadlines]({OS_PAGE}) on one board so planners can spot "
     "clashes before they reach the floor.", "capability"),
    ("Does Acme Widgets publish maintenance guides?", "Acme Widgets publishes maintenance guides that explain "
     "cleaning, lubrication, inspection intervals and common faults, so owners can keep their equipment running "
     "between scheduled service visits.", "capability"),
    ("Can I visit the Acme Widgets factory?", f"Acme Widgets runs a [factory tour in Geelong]({ABOUT}) where "
     "visitors can watch widgets being machined, assembled and tested, and ask the engineers questions about "
     "their own projects.", "proof"),
    ("Who founded Acme Widgets?", f"Acme Widgets was founded by Jane Smith, who set up the [Geelong widget "
     f"workshop]({ABOUT}) to give local manufacturers dependable parts and direct access to the people who make "
     "them.", "founder"),
    ("Is Acme Widgets useful for hardware retailers?", "Acme Widgets supplies hardware retailers with stocked "
     "widget ranges, shelf-ready packaging and reorder support, which lets store owners carry dependable parts "
     "without holding excess inventory.", "audience"),
    ("Is Acme Widgets useful for engineering teams?", f"Acme Widgets works with engineering teams on "
     f"[bespoke widget prototypes]({CUSTOM}), sharing drawings, tolerances and test notes so that designers can "
     "move from concept to a working part quickly.", "audience"),
    ("What resources does Acme Widgets provide?", "Acme Widgets provides drawings, installation notes, "
     "maintenance checklists and spare part lists alongside every order, giving customers the paperwork they need "
     "to fit and service each widget.", "resources"),
    ("How is Acme Widgets different from importers?", "Acme Widgets machines and assembles its widgets locally in "
     "Geelong, which means customers deal directly with the makers, receive parts sooner and can request changes "
     "without overseas delays.", "differentiator"),
]


def faqs_md(items, brand=BRAND):
    out = [f"# {brand} FAQs", ""]
    for i, (q, a, bucket) in enumerate(items):
        place = "homepage" if i < 8 else "faq-page"
        out += [f"### {q}", a, f"<!-- bucket: {bucket}; placement: {place}; sources: {HOME}, {ABOUT}; "
                "evidence: first-party -->", ""]
    return "\n".join(out)


def opts(**over):
    base = {"brand": BRAND, "products": [PRODUCT], "locale": "en-AU", "forbid": [], "allow_volatile": False,
            "final": True, "inventory": {f.clean_url(r["url"]): {k: str(v) for k, v in r.items()} for r in INVENTORY},
            "links": [], "has_candidates": False}
    base.update(over)
    return base


def with_answer(n, answer, question=None):
    items = [list(x) for x in GOOD]
    items[n][1] = answer
    if question:
        items[n][0] = question
    return faqs_md([tuple(x) for x in items])


def fails(md, **over):
    return f.lint_faqs(md, BUNDLE, opts(**over)).fails


class LintTest(unittest.TestCase):
    def assertFailsWith(self, fl, needle):
        self.assertTrue(any(needle in x for x in fl), f"expected a fail containing {needle!r}, got {fl}")

    def test_good_fixture_passes(self):
        r = f.lint_faqs(faqs_md(GOOD), BUNDLE, opts())
        self.assertEqual(r.fails, [], r.fails)

    def test_two_sentences_pass(self):
        md = with_answer(0, "Acme Widgets is a Geelong widget maker for small manufacturers. It designs, builds and "
                            "supports industrial widgets with durable parts and clear documentation for owners.")
        self.assertEqual(fails(md), [])

    def test_yes_opener_passes(self):
        md = with_answer(5, "Yes, Acme Widgets publishes maintenance guides that explain cleaning, lubrication, "
                            "inspection intervals and common faults so owners can keep equipment running.")
        self.assertEqual(fails(md), [])

    def test_number_not_in_bundle_fails(self):
        md = with_answer(5, "Acme Widgets publishes 37 maintenance guides that explain cleaning, lubrication, "
                            "inspection intervals and common faults so owners can keep equipment running.")
        self.assertFailsWith(fails(md), "number not in the public bundle")

    def test_unsupported_tool_name_fails(self):
        md = with_answer(2, "WidgetOS connects to Shopify, Xero and Codex so that online orders, invoices and stock "
                            "levels stay in step with the factory schedule instead of being copied by hand.")
        self.assertFailsWith(fails(md), "'Codex'")

    def test_names_do_not_join_across_sentences(self):
        self.assertEqual(f.cap_phrases('WidgetOS. Shopify connects to Xero.'), ['WidgetOS', 'Shopify', 'Xero'])
        md = with_answer(0, 'Acme Widgets makes widgets. Payments are non-refundable. Cancelling is free.')
        self.assertEqual(fails(md), [])
        md = with_answer(0, 'Acme Widgets makes widgets. Codex connects to the factory.')
        self.assertFailsWith(fails(md), "'Codex'")

    def test_link_to_redirect_fails(self):
        md = faqs_md(GOOD).replace(f"[production planning software]({OS_PAGE})",
                                   f"[production planning software]({OLD})")
        self.assertFailsWith(fails(md), "redirects")

    def test_anchor_not_in_answer_fails(self):
        links = [{"n": 6, "target": CUSTOM, "anchor": "robotic welding cells"}]
        self.assertFailsWith(fails(faqs_md(GOOD), links=links), "not verbatim")

    def test_generic_anchor_fails(self):
        md = with_answer(10, f"Acme Widgets provides drawings, installation notes and maintenance checklists with "
                             f"every order, and customers can [learn more]({CUSTOM}) about fitting and servicing "
                             "each widget.")
        self.assertFailsWith(fails(md), "generic anchor")

    def test_homepage_link_fails(self):
        md = with_answer(10, f"Acme Widgets provides drawings, installation notes, [maintenance checklists]({HOME}) "
                             "and spare part lists alongside every order, giving customers the paperwork they need "
                             "to fit each widget.")
        self.assertFailsWith(fails(md), "homepage")

    def test_link_in_links_json_is_applied_and_checked(self):
        links = [{"question": GOOD[5][0], "target": CUSTOM, "anchor": "maintenance guides"}]
        md = faqs_md(GOOD).replace(f"[custom widget design services]({CUSTOM})", "custom widget design services")
        self.assertEqual(fails(md, links=links), [])
        title, items = f.parse_faqs(md)
        self.assertIn(f"[maintenance guides]({CUSTOM})", f.linked_answer(items[5], links))

    def test_volatile_and_promise_fail(self):
        md = with_answer(5, "Acme Widgets guarantees the best maintenance guides in 2026, explaining cleaning, "
                            "lubrication, inspection intervals and common faults so owners keep equipment running.")
        fl = fails(md)
        self.assertFailsWith(fl, "promise")
        self.assertFailsWith(fl, "superlative")
        self.assertFailsWith(fl, "volatile")

    def test_lowercase_may_and_hyphenated_ai_are_fine(self):
        md = with_answer(5, "Acme Widgets publishes AI-assisted maintenance guides that owners may use to plan "
                            "cleaning, lubrication, inspection intervals and common faults between service visits.")
        self.assertEqual(fails(md), [])

    def test_us_spelling_fails_for_en_au(self):
        md = with_answer(5, "Acme Widgets publishes maintenance guides that help owners organize cleaning, "
                            "lubrication, inspection intervals and common faults so their equipment keeps running.")
        self.assertFailsWith(fails(md), "spelling")


class ReviewFixesTest(unittest.TestCase):
    def assertFailsWith(self, fl, needle):
        self.assertTrue(any(needle in x for x in fl), f"expected a fail containing {needle!r}, got {fl}")

    def test_inline_plus_links_json_counts_once(self):
        links = [{"n": 2, "target": OS_PAGE, "anchor": "production planning software"}]
        r = f.lint_faqs(faqs_md(GOOD), BUNDLE, opts(links=links))
        self.assertEqual(r.fails, [])
        self.assertIn("6 of 12 answers linked", " ".join(r.notes))

    def test_bucket_quotas_do_not_drop_questions(self):
        items = [list(x) for x in GOOD]
        items[11][2] = "identity"
        self.assertEqual(fails(faqs_md([tuple(x) for x in items])), [])

    def test_third_party_evidence_needs_off_domain_source(self):
        md = faqs_md(GOOD).replace("evidence: first-party", "evidence: third-party", 1)
        self.assertFailsWith(fails(md), "every source is on the brand's own domain")
        ok = md.replace(f"sources: {HOME}, {ABOUT}; evidence: third-party",
                        f"sources: {HOME}, https://news.example.org/acme; evidence: third-party")
        self.assertEqual(fails(ok), [])

    def test_negated_guarantee_passes(self):
        md = with_answer(5, "Acme Widgets publishes maintenance guides as practical advice rather than a guarantee "
                            "of uptime, covering cleaning, lubrication, inspection intervals and common faults.")
        self.assertEqual(fails(md), [])

    def test_fact_only_in_private_bundle_fails_and_warns(self):
        md = with_answer(2, "WidgetOS connects to Shopify, Xero and MYOB so that online orders, invoices and stock "
                            "levels stay in step with the factory schedule instead of being copied by hand.")
        r = f.lint_faqs(md, BUNDLE, opts(private=BUNDLE + "\nBrain: WidgetOS also talks to MYOB."))
        self.assertTrue(any("'MYOB' is not in the public bundle" in x for x in r.fails), r.fails)
        self.assertTrue(any("only in private inputs" in w for w in r.warns), r.warns)

    def test_gather_public_bundle_excludes_brain_and_private_research(self):
        with tempfile.TemporaryDirectory() as tmp:
            run, brain = Path(tmp) / "2026-10" / "homepage-faqs", Path(tmp) / "brain"
            (run / "data" / "pages").mkdir(parents=True)
            (brain / "business" / "brand").mkdir(parents=True)
            (run / "data" / "pages" / "home.md").write_text("Acme Widgets makes widgets.", encoding="utf-8")
            (brain / "business" / "brand" / "voice.md").write_text("Secret client Globex.", encoding="utf-8")
            (run / "data" / "research.md").write_text("## Summary\nPress coverage is sparse.\n\n## Notes\n"
                                                       "Private: Initech https://www.zoominfo.com/p/x\n", encoding="utf-8")
            f.cmd_gather(Namespace(run_dir=str(run), brain=str(brain), include=None))
            public = (run / f.PUBLIC_BUNDLE).read_text(encoding="utf-8")
            self.assertIn("Press coverage is sparse", public)
            self.assertNotIn("Globex", public)
            self.assertNotIn("Initech", public)
            self.assertIn("Globex", (run / f.BUNDLE).read_text(encoding="utf-8"))

    def test_unicode_words_names_and_slugs(self):
        self.assertEqual(f.words_of("Café Zürich naïve"), ["café", "zürich", "naïve"])
        self.assertIn("Zürich", f.cap_phrases("widgets made in Zürich for Ærø teams"))
        self.assertIn("Ærø", f.cap_phrases("widgets made in Zürich for Ærø teams"))
        slugs = {f.page_slug(u) for u in ("https://a.example/a-b", "https://a.example/a/b", "https://a.example/a/b/")}
        self.assertEqual(len(slugs), 3)
        self.assertTrue(f.page_slug("https://a.example/%C3%BCber-uns/").startswith("über-uns-"))
        self.assertTrue(f.Bundle("Made in Zürich.").has("Zürich"))

    def test_hub_pages(self):
        rows = [{"url": "https://a.example/guides/geo/", "page_type": "blog-post"},
                {"url": "https://a.example/guides/geo/post/", "page_type": "blog-post"}]
        f.mark_hubs(rows)
        self.assertEqual([r["page_type"] for r in rows], ["hub", "blog-post"])


def steve_sample():
    text = (HERE.parent / "references" / "source-sample-steve-toth-gpt.md").read_text(encoding="utf-8")
    pairs = re.findall(r"^\*\*(.+?\?)\*\*\n(.+)$", text, re.M)
    buckets = ["identity"] + ["capability"] * (len(pairs) - 2) + ["audience"]
    return [(q, a, b) for (q, a), b in zip(pairs, buckets)]


class SteveSampleTest(unittest.TestCase):
    def test_raw_gpt_answers_flag_unsupported_names(self):
        items = steve_sample()
        self.assertEqual(len(items), 20)
        bundle = ("The Vibe Marketing Lab is a free community teaching AI SEO and GEO. MarketingOS is its system. "
                  "Members work with ChatGPT and Claude.")
        md = faqs_md(items, brand="The Vibe Marketing Lab")
        r = f.lint_faqs(md, bundle, opts(brand="The Vibe Marketing Lab", products=["MarketingOS"], locale="en-US",
                                         inventory={}))
        flagged = " ".join(r.fails)
        for name in ("'Codex'", "'Perplexity'", "'Google'", "'Claude Code'"):
            self.assertIn(name, flagged)
        for ok in ("'ChatGPT'", "'MarketingOS'", "'GEO'"):
            self.assertNotIn(f"name {ok}", flagged)


class LinksTest(unittest.TestCase):
    def test_bm25_ranks_matching_page_first(self):
        inv = {f.clean_url(r["url"]): {k: str(v) for k, v in r.items()} for r in INVENTORY}
        rows = f.link_targets(inv)
        self.assertNotIn(HOME, [r["url"] for r in rows])
        self.assertNotIn(OLD, [r["url"] for r in rows])
        docs = [f.tokens(" ".join([r["title"], r["h1"], r["meta_description"]])) for r in rows]
        _, items = f.parse_faqs(faqs_md(GOOD))
        out = f.shortlist(items, rows, docs, f.names_words([BRAND, PRODUCT]))
        self.assertEqual(out[3]["candidates"][0]["url"], CUSTOM)  # custom widget design
        self.assertEqual(out[6]["candidates"][0]["url"], ABOUT)   # factory tour
        self.assertLessEqual(len(out[0]["candidates"]), 3)

    def test_crawl_exclusions(self):
        for bad in ("https://acme.example/tag/news/", "https://acme.example/blog/page/2/", "https://acme.example/feed/",
                    "https://acme.example/login", "https://acme.example/logo.png", "https://acme.example/category/x/"):
            self.assertIsNotNone(f.EXCLUDE.search(bad), bad)
        self.assertIsNone(f.EXCLUDE.search("https://acme.example/services/custom-widgets/"))
        self.assertEqual(f.page_type(HOME), "homepage")
        self.assertEqual(f.page_type("https://acme.example/blog/how-to/"), "blog-post")


class RenderWorkbookTest(unittest.TestCase):
    def test_legacy_layout_migration_preserves_copy_and_files(self):
        before = (self.run / f.FAQS_MD).read_bytes()
        artifacts = {'faqs.html': b'old html', 'faqpage.jsonld': b'old schema',
                     'handover.md': b'old handover', 'comparison.md': b'old comparison',
                     'README.md': b'old readme', 'extra-notes.md': b'user notes',
                     'source.csv': b'question,answer', 'capture.png': b'image bytes'}
        for name, content in artifacts.items():
            (self.run / name).write_bytes(content)
        (self.run / 'page-faqs').mkdir()
        (self.run / 'page-faqs/item.html').write_bytes(b'page block')
        self.assertEqual(f.organize_run(self.run), len(artifacts) + 1)
        self.assertEqual((self.run / f.FAQS_MD).read_bytes(), before)
        self.assertEqual([p.name for p in self.run.iterdir() if p.is_file()], ['faqs.md'])
        for name, content in artifacts.items():
            self.assertEqual(f.artifact_target(self.run, self.run / name).read_bytes(), content)
        self.assertEqual((self.run / f.PAGE_FAQS / 'item.html').read_bytes(), b'page block')
        self.assertEqual(f.organize_run(self.run), 0)

    def test_layout_conflict_does_not_overwrite_or_partially_move(self):
        (self.run / 'comparison.md').write_text('original', encoding='utf-8')
        (self.run / 'README.md').write_text('readme', encoding='utf-8')
        (self.run / 'review').mkdir()
        (self.run / 'review/comparison.md').write_text('client edits', encoding='utf-8')
        with self.assertRaises(ValueError):
            f.organize_run(self.run)
        self.assertEqual((self.run / 'review/comparison.md').read_text(encoding='utf-8'), 'client edits')
        self.assertTrue((self.run / 'README.md').is_file())
        self.assertTrue((self.run / 'comparison.md').is_file())

    def test_root_stays_clean_after_render_workbook_and_rerun(self):
        self.make_book()
        for _ in range(2):
            self.assertEqual(self.lint(), 0)
            self.assertEqual(f.cmd_render(Namespace(run_dir=str(self.run), url=HOME, brand=None, heading=None, forbid=None)), 0)
            self.assertEqual(f.cmd_workbook(Namespace(run_dir=str(self.run), brand=None)), 0)
            self.assertEqual([p.name for p in self.run.iterdir() if p.is_file()], ['faqs.md'])
            self.assertTrue((self.run / f.FAQS_HTML).is_file())
            self.assertTrue((self.run / f.PLACEMENT).is_file())
            text = (self.run / f.HANDOVER).read_text(encoding='utf-8')
            self.assertIn('publishing/faqs.html', text)
            self.assertIn('review/placement-plan.md', text)

    def test_local_workbook_moves_into_review_and_preserves_client_state(self):
        import openpyxl
        self.make_book().rename(self.run / f.WORKBOOK)
        args = Namespace(run_dir=str(self.run), brand=None)
        self.assertEqual(f.cmd_workbook(args), 0)
        book = self.run / 'review' / f.WORKBOOK
        self.assertTrue(book.is_file())
        wb = openpyxl.load_workbook(book)
        wb[f.TAB].cell(row=7, column=11, value='Approve')
        ini = wb['Initiatives']
        ini.cell(row=6, column=7, value='In progress')
        ini.cell(row=6, column=13, value='homepage-faqs/faqs.html')
        wb.save(book)
        self.assertEqual(f.cmd_workbook(args), 0)
        wb = openpyxl.load_workbook(book)
        self.assertEqual(wb[f.TAB].cell(row=7, column=11).value, 'Approve')
        self.assertEqual(wb['Initiatives'].cell(row=6, column=7).value, 'In progress')
        self.assertEqual(wb['Initiatives'].cell(row=6, column=13).value, 'homepage-faqs/publishing/faqs.html')
        self.assertIn('review/handover.md', wb['Initiatives'].cell(row=6, column=5).value)
        self.assertEqual([p.name for p in self.run.iterdir() if p.is_file()], ['faqs.md'])

    def test_multiblock_answers_match_visible_html_and_schema(self):
        answer = ('We make widgets at Acme Widgets. Choose the format that fits your work.\n\n'
                  '- Stock widgets.\n- Custom widgets.\n\n'
                  '| Type | Use |\n| --- | --- |\n| Stock | Maintenance |\n\n'
                  '<script>alert("x")</script>')
        _, parsed = f.parse_faqs(faqs_md([(GOOD[0][0], answer, 'identity')]))
        self.assertEqual(parsed[0]['answer'], answer)
        block = f.answer_html(answer)
        self.assertIn('<ul><li>Stock widgets.</li>', block)
        self.assertIn('<table><thead>', block)
        self.assertNotIn('<script>', block)
        self.assertIn('&lt;script&gt;', block)
        self.assertIn(block, f.faqs_html([(GOOD[0][0], answer)], 'FAQs'))
        self.assertEqual(f.faq_jsonld([(GOOD[0][0], answer)], HOME)['mainEntity'][0]['acceptedAnswer']['text'], block)

    def test_page_routing_and_unanswered_dealbreaker_retention(self):
        md = faqs_md(GOOD).replace('bucket: product; placement: homepage;',
                                   f'bucket: product; placement: page; target: {OS_PAGE};')
        md = md.replace('bucket: resources; placement: faq-page;',
                        'bucket: resources; placement: content-gap; target: https://acme.example/resources/;')
        (self.run / f.FAQS_MD).write_text(md, encoding='utf-8')
        q = 'What is Acme Widgets refund policy?'
        (self.run / 'data/question-plan.json').write_text(json.dumps([
            {'question': q, 'status': 'needs-client-answer', 'placement': 'homepage',
             'reason': 'Client must confirm the refund policy.'}]), encoding='utf-8')
        self.assertEqual(self.lint(), 0)
        self.assertEqual(f.cmd_render(Namespace(run_dir=str(self.run), url=HOME, brand=None, heading=None, forbid=None)), 0)
        schema = json.loads((self.run / f.PAGE_FAQS / (f.page_slug(OS_PAGE) + '.jsonld')).read_text(encoding='utf-8'))
        self.assertEqual(schema['url'], OS_PAGE)
        self.assertEqual(schema['mainEntity'][0]['name'], GOOD[1][0])
        self.assertIn(q, (self.run / f.PLACEMENT).read_text(encoding='utf-8'))
        home = (self.run / f.JSONLD).read_text(encoding='utf-8')
        self.assertNotIn(GOOD[1][0], home)
        for path in list((self.run / 'publishing').glob('*.html')) + list((self.run / f.PAGE_FAQS).glob('*')):
            text = path.read_text(encoding='utf-8')
            self.assertNotIn(q, text)
            self.assertNotIn(GOOD[10][0], text)
        self.assertEqual(len(list((self.run / f.PAGE_FAQS).glob('*.html'))), 1)

    def test_unresolved_or_omitted_ready_question_blocks_export(self):
        path = self.run / 'data/question-plan.json'
        for row in ({'question': GOOD[0][0], 'status': 'needs-client-answer'},
                    {'question': 'Can I get a refund?', 'status': 'ready', 'placement': 'homepage'},
                    {'question': 'Different wording?', 'status': 'merged', 'merged_into': 'Missing question?'}):
            path.write_text(json.dumps([row]), encoding='utf-8')
            self.assertEqual(self.lint(), 1)

    def test_question_and_voice_files_are_private_inputs(self):
        questions, voice = self.run / 'questions.md', self.run / 'voice.md'
        questions.write_text('PrivateDealbreaker and proposed 37 seats.', encoding='utf-8')
        voice.write_text('PrivateVoice: We speak casually.', encoding='utf-8')
        f.cmd_gather(Namespace(run_dir=str(self.run), brain=None, include=None,
                               questions=[str(questions)], voice=[str(voice)]))
        private = (self.run / f.BUNDLE).read_text(encoding='utf-8')
        public = (self.run / f.PUBLIC_BUNDLE).read_text(encoding='utf-8')
        for term in ('PrivateDealbreaker', 'PrivateVoice'):
            self.assertIn(term, private)
            self.assertNotIn(term, public)

    def test_first_person_short_answer_and_client_placeholder(self):
        self.assertEqual(fails(with_answer(0, 'We make widgets at Acme Widgets.')), [])
        self.assertEqual(fails(with_answer(0, 'We make widgets at Acme Widgets — built locally.')), [])
        self.assertTrue(fails('# Acme Widgets FAQs\n'))
        self.assertTrue(fails(with_answer(0, 'Acme Widgets policy: [CLIENT TO ANSWER: confirm refunds].')))

    def test_page_target_must_exist_and_be_indexable(self):
        md = faqs_md(GOOD).replace('bucket: product; placement: homepage;',
                                   'bucket: product; placement: page; target: https://acme.example/missing/;')
        self.assertTrue(any('target' in error for error in fails(md)))
        md = md.replace('https://acme.example/missing/', 'https://elsewhere.example/')
        self.assertTrue(fails(md))

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.month = Path(self.tmp.name) / "2026-10"
        self.run = self.month / "homepage-faqs"
        (self.run / "data").mkdir(parents=True)
        (self.run / "faqs.md").write_text(faqs_md(GOOD), encoding="utf-8")
        with (self.run / f.INVENTORY).open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=f.INV_COLS)
            w.writeheader()
            w.writerows(INVENTORY)

    def tearDown(self):
        self.tmp.cleanup()

    def lint(self, final=True, forbid=None):
        (self.run / f.PUBLIC_BUNDLE).write_text(BUNDLE, encoding="utf-8")
        return f.cmd_lint(Namespace(run_dir=str(self.run), faqs=None, bundle=None, brand=BRAND,
                                    products=[PRODUCT], locale="en-AU", forbid=forbid, allow_volatile=False,
                                    final=final, count=20))

    def test_render(self):
        self.assertEqual(self.lint(), 0)
        f.cmd_render(Namespace(run_dir=str(self.run), url="https://acme.example", brand=None, heading=None,
                               forbid=None))
        page = (self.run / f.FAQS_HTML).read_text(encoding="utf-8")
        self.assertEqual(page.count("<details"), 8)
        self.assertIn(f'<a href="{OS_PAGE}">production planning software</a>', page)
        self.assertNotIn("<script", page)
        ld = json.loads((self.run / f.JSONLD).read_text(encoding="utf-8"))
        self.assertEqual(ld["@type"], "FAQPage")
        self.assertIn(f'<a href="{OS_PAGE}">', ld["mainEntity"][1]["acceptedAnswer"]["text"])
        self.assertNotIn("&quot;", (self.run / f.JSONLD).read_text(encoding="utf-8"))
        self.assertEqual(ld["url"], HOME)
        self.assertEqual(ld["@id"], HOME + "#faq")
        self.assertIn("August 2023", (self.run / f.HANDOVER).read_text(encoding="utf-8"))
        self.assertTrue((self.run / f.FAQ_PAGE).is_file())
        self.assertTrue((self.run / f.NOTE).is_file())
        self.assertEqual([p.name for p in self.run.iterdir() if p.is_file()], ['faqs.md'])

    def test_render_fails_on_forbidden_term_and_strips_brokers(self):
        (self.run / "data" / "research.md").write_text(
            "## Summary\nFounder profile seen at https://www.zoominfo.com/p/jane and in Globex press.\n\n"
            "## Working notes\nEmployer: Initech.\n", encoding="utf-8")
        self.assertEqual(self.lint(), 0)
        args = Namespace(run_dir=str(self.run), url=HOME, brand=None, heading=None, forbid=["Globex"])
        self.assertEqual(f.cmd_render(args), 1)
        note = (self.run / f.NOTE).read_text(encoding="utf-8")
        self.assertNotIn("zoominfo", note)
        self.assertNotIn("Initech", note)
        args.forbid = ["Initech"]
        self.assertEqual(f.cmd_render(args), 0)
        self.assertEqual(f.rendered_leaks(self.run, ["Globex"]), [(f.NOTE, "Globex")])

    def test_export_rejects_missing_nonfinal_failed_and_stale_lint(self):
        args = Namespace(run_dir=str(self.run), url=HOME, brand=None, heading=None, forbid=None)
        self.assertEqual(f.cmd_render(args), 1)
        self.assertFalse((self.run / f.FAQS_HTML).exists())
        self.assertEqual(self.lint(final=False), 0)
        self.assertEqual(f.cmd_render(args), 1)
        self.assertEqual(self.lint(), 0)
        self.assertEqual(f.cmd_render(args), 0)
        (self.run / f.FAQS_MD).write_text(with_answer(0, "Yes. A second sentence."), encoding="utf-8")
        self.assertEqual(f.cmd_render(args), 1)
        self.assertEqual(self.lint(), 1)
        self.assertFalse((self.run / f.DATA / "lint-receipt.json").exists())

    def test_export_rejects_changed_evidence_links_and_research(self):
        args = Namespace(run_dir=str(self.run), url=HOME, brand=None, heading=None, forbid=None)
        for name in [f.PUBLIC_BUNDLE, f.INVENTORY, f.LINKS_JSON, f.RESEARCH,
                     'data/question-plan.json', 'data/voice-brief.md']:
            with self.subTest(name=name):
                self.assertEqual(self.lint(), 0)
                path = self.run / name
                before = path.read_bytes() if path.exists() else None
                path.write_bytes((before or b"") + b"\nchanged\n")
                self.assertEqual(f.cmd_render(args), 1)
                if before is None:
                    path.unlink()
                else:
                    path.write_bytes(before)

    def test_export_preserves_never_say_policy(self):
        self.assertEqual(self.lint(forbid=["Secret client"]), 0)
        args = Namespace(run_dir=str(self.run), url=HOME, brand=None, heading=None, forbid=None)
        self.assertEqual(f.cmd_render(args), 1)
        args.forbid = ["Secret client"]
        self.assertEqual(f.cmd_render(args), 0)

    def make_book(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Checklist"
        ws.cell(row=8, column=3, value=f.SKILL_ID)
        wb.create_sheet("AI Info Page")
        ini = wb.create_sheet("Initiatives")
        ini.cell(row=5, column=4, value="Task")
        wb._sheets = [wb["Checklist"], wb["AI Info Page"], wb["Initiatives"]]
        book = self.month / f.WORKBOOK
        wb.save(book)
        return book

    def test_workbook_tab_and_verdicts_survive_rerun(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed (run via uv run --with openpyxl)")
        book = self.make_book()
        args = Namespace(run_dir=str(self.run), brand=None)
        self.assertEqual(f.cmd_workbook(args), 0)
        wb = openpyxl.load_workbook(book)
        self.assertEqual(wb.sheetnames, ["Checklist", "AI Info Page", f.TAB, "Initiatives"])
        self.assertEqual(wb["Checklist"].cell(row=8, column=9).value, "☑")
        ws = wb[f.TAB]
        self.assertEqual([ws.cell(row=6, column=c).value for c in range(2, 14)], f.TAB_HEADERS)
        self.assertEqual(ws.cell(row=7, column=4).value, GOOD[0][0])
        self.assertEqual(ws.cell(row=8, column=9).value, OS_PAGE)
        ws.cell(row=7, column=11, value="Edit")
        ws.cell(row=7, column=12, value="A better answer")
        wb.save(book)
        self.assertEqual(f.cmd_workbook(args), 0)
        wb = openpyxl.load_workbook(book)
        ws = wb[f.TAB]
        self.assertEqual(ws.cell(row=7, column=11).value, "Edit")
        self.assertEqual(ws.cell(row=7, column=12).value, "A better answer")
        tasks = [c.value for c in wb["Initiatives"]["D"] if c.value and str(c.value).startswith("Publish the homepage")]
        self.assertEqual(tasks, [f"Publish the homepage FAQs for {BRAND}"])


class GluedTextTest(unittest.TestCase):
    def test_heading_glued_to_next_line_still_matches(self):
        b = f.Bundle("Build In Public RecordingsRaw and unfiltered recordings.")
        self.assertTrue(b.has("Build In Public Recordings"))
        self.assertFalse(f.Bundle("Recordingsraw text").has("Recordings"))
        self.assertTrue(f.Bundle("PremiumThe Prompt LibrarySteal my prompts").has("The Prompt Library"))
        self.assertFalse(f.Bundle("xthe prompt library").has("The Prompt Library"))


if __name__ == "__main__":
    unittest.main()
