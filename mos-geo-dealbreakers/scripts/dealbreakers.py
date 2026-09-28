#!/usr/bin/env python3
"""mos-geo-dealbreakers: build a Dealbreaker Detector report from a brand's public pages.

Subcommands
  path      print this month's run folder (MarketingOS-aware, same rules as ai-info)
  crawl     fetch the brand's own site: sitemap, nav/footer links, buyer-page ranking,
            optional extra URLs and a Wayback homepage snapshot
  add       fetch more pages (reviews, community posts) or save text fetched another way
  build     lint data/findings.json against references/findings-schema.md, then render the
            report, the FAQ draft and the findings CSV
  workbook  tick the Checklist, add a 'Dealbreakers' tab and (own mode) Initiatives rows

Standard library only, except `workbook` (openpyxl, via uv). Scrapling is an optional
fallback for pages that block plain requests or only render with JavaScript:
    uv run --with "scrapling[fetchers]" python dealbreakers.py crawl ...
"""

from __future__ import annotations

import argparse
import csv
import gzip
import html
import json
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
PACK_TEMPLATE = Path(__file__).resolve().parents[2] / "_shared" / "brand-audit" / "brand-audit-template.xlsx"
SKILL_ID = "mos-geo-dealbreakers"
SKILL_FOLDER = "dealbreakers"
DATA_DIR = "data"
WORKBOOK = "brand-audit-master.xlsx"
REPORT = "dealbreakers-report.md"
FAQ_DRAFT = "faq-draft.md"
FINDINGS_JSON = "data/findings.json"
FINDINGS_CSV = "data/findings.csv"

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36")
HEADERS = {"User-Agent": UA, "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
           "Accept-Language": "en-AU,en;q=0.9"}

# ------------------------------------------------------------------ path --
# Kept in step with aiinfo.py / brand360.py so every mos-geo skill lands in the same
# month folder.


def _main_checkout(git_file: Path) -> Path | None:
    m = re.match(r"gitdir:\s*(.+)", git_file.read_text(encoding="utf-8").strip())
    if not m:
        return None
    raw = m.group(1).strip()
    drive = re.match(r"^([A-Za-z]):[\\/](.*)$", raw)
    if drive and not sys.platform.startswith("win"):
        raw = f"/mnt/{drive.group(1).lower()}/{drive.group(2)}"
    gitdir = Path(raw.replace("\\", "/"))
    if not gitdir.is_absolute():
        gitdir = (git_file.parent / gitdir).resolve()
    for d in [gitdir, *gitdir.parents]:
        if d.name == ".git":
            return d.parent
    return None


def brain_config(root: Path) -> Path | None:
    own = root / ".mos" / "config.yaml"
    if own.is_file():
        return own
    git = root / ".git"
    if git.is_file():
        main = _main_checkout(git)
        if main and (main / ".mos" / "config.yaml").is_file():
            return main / ".mos" / "config.yaml"
    return None


def find_brain(start: Path) -> Path | None:
    for d in [start, *start.parents]:
        if brain_config(d):
            return d
        if (d / ".git").exists():
            return None
    return None


def brain_mode(brain: Path) -> str:
    config = brain_config(brain)
    if not config:
        return "in-house"
    m = re.search(r'["\']?mode["\']?\s*:\s*["\']?([a-z-]+)', config.read_text(encoding="utf-8"))
    return m.group(1) if m else "in-house"


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def run_dir_for(brand: str, day: str, start: Path) -> Path:
    """campaigns/geo/YYYY-MM/dealbreakers/ in a one-brand brain, with a brand folder in an
    agency brain, outputs/geo/YYYY-MM/<brand>/dealbreakers/ outside a brain. A repeat run in
    the same month gets dealbreakers-2, -3 ... so nothing is overwritten."""
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", day):
        sys.exit(f"--date must be YYYY-MM-DD, got {day}")
    slug = slugify(brand)
    brain = find_brain(start.resolve())
    if brain and brain_mode(brain) != "agency":
        month = brain / "campaigns" / "geo" / day[:7]
    elif brain:
        month = brain / "campaigns" / "geo" / day[:7] / slug
    else:
        month = start.resolve() / "outputs" / "geo" / day[:7] / slug
    target, n = month / SKILL_FOLDER, 2
    while target.exists() and any(target.iterdir()):
        target = month / f"{SKILL_FOLDER}-{n}"
        n += 1
    return target


def cmd_path(args) -> int:
    print(run_dir_for(args.brand, args.date or time.strftime("%Y-%m-%d"), Path(args.start or ".")))
    return 0


# ----------------------------------------------------------------- fetch --

def fetch(url: str, timeout: int = 30, method: str = "GET") -> tuple[int, dict, bytes, str]:
    """(status, headers, body, final_url). Never raises on HTTP errors."""
    req = urllib.request.Request(url, headers=HEADERS, method=method)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            body = r.read() if method == "GET" else b""
            if r.headers.get("Content-Encoding") == "gzip" or body[:2] == b"\x1f\x8b":
                try:
                    body = gzip.decompress(body)
                except OSError:
                    pass
            return r.status, {k.lower(): v for k, v in r.headers.items()}, body, r.geturl()
    except urllib.error.HTTPError as e:
        return e.code, {k.lower(): v for k, v in (e.headers or {}).items()}, b"", url
    except Exception as e:  # noqa: BLE001 - network errors become a status the caller reports
        return 0, {"error": str(e)}, b"", url


def fetch_scrapling(url: str, render: bool = False) -> tuple[int, dict, bytes, str] | None:
    """Optional fallback. None when Scrapling is not installed or the fetch fails."""
    try:
        from scrapling.fetchers import DynamicFetcher, Fetcher
    except ImportError:
        return None
    try:
        r = DynamicFetcher.fetch(url, network_idle=True) if render else Fetcher.get(url, stealthy_headers=True)
        body = r.body if isinstance(r.body, bytes) else str(r.body).encode("utf-8")
        return r.status, {k.lower(): v for k, v in dict(r.headers or {}).items()}, body, str(r.url or url)
    except Exception:  # noqa: BLE001 - a failed fallback leaves the original result in place
        return None


SCRAPLING_HINT = ('uv run --with "scrapling[fetchers]" python dealbreakers.py ...  '
                  "(and `uv run --with \"scrapling[fetchers]\" scrapling install` once for JavaScript pages)")


def fetch_page(url: str, delay: float, mode: str = "auto") -> tuple[int, bytes, str, str]:
    """(status, body, final_url, via) for crawling. Backs off once on 403/429/503, then
    tries Scrapling (mode auto/always) when the page is blocked or looks JavaScript-only."""
    via = "urllib"
    if mode == "always":
        got = fetch_scrapling(url)
        if got and got[0] == 200:
            return got[0], got[2], got[3], "scrapling"
    st, _, b, fu = fetch(url)
    if st in (403, 429, 503):
        time.sleep(max(5.0, delay * 8))
        st, _, b, fu = fetch(url)
    if mode != "off" and st in (0, 403, 429, 503):
        got = fetch_scrapling(url)
        if got and got[0] == 200:
            st, _, b, fu = got
            via = "scrapling"
    if mode != "off" and st == 200 and b and js_only(b):
        got = fetch_scrapling(url, render=True)
        if got and got[0] == 200 and not js_only(got[2]):
            st, _, b, fu = got
            via = "scrapling-browser"
    time.sleep(delay)
    return st, b, fu, via


def js_only(body: bytes) -> bool:
    """A page whose server HTML carries almost no text but plenty of script."""
    return len(parse_page("https://x/", body).text().split()) < 80 and text_of(body).count("<script") >= 5


def text_of(body: bytes) -> str:
    return body.decode("utf-8", errors="replace")


class PageParser(HTMLParser):
    """Pulls what a dealbreaker report needs out of one HTML page, stdlib only."""

    SKIP = {"script", "style", "noscript", "svg", "template", "iframe"}
    BLOCK = {"p", "div", "section", "article", "li", "br", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
             "header", "footer", "nav", "aside", "main", "table", "ul", "ol", "dd", "dt", "blockquote"}

    def __init__(self, base: str):
        super().__init__(convert_charrefs=True)
        self.base = base
        self.title = ""
        self.meta: dict[str, str] = {}
        self.links: list[tuple[str, str]] = []
        self.headings: list[tuple[str, str]] = []
        self._stack: list[str] = []
        self._skip = 0
        self._in_title = False
        self._heading: str | None = None
        self._heading_buf: list[str] = []
        self._link: str | None = None
        self._link_buf: list[str] = []
        self.parts: list[str] = []

    def handle_starttag(self, tag, attrs):
        a = {k: (v or "") for k, v in attrs}
        if tag in self.SKIP:
            self._skip += 1
            return
        if tag == "title":
            self._in_title = True
        elif tag == "meta":
            key = (a.get("name") or a.get("property") or "").lower()
            if key:
                self.meta[key] = a.get("content", "")
        elif tag == "a" and a.get("href"):
            self._link = urllib.parse.urljoin(self.base, a["href"].strip())
            self._link_buf = []
        elif re.fullmatch(r"h[1-4]", tag):
            self._heading, self._heading_buf = tag, []
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skip = max(0, self._skip - 1)
            return
        if tag == "title":
            self._in_title = False
        elif tag == "a" and self._link:
            self.links.append((self._link, " ".join("".join(self._link_buf).split())))
            self._link = None
        elif self._heading == tag:
            txt = " ".join("".join(self._heading_buf).split())
            if txt:
                self.headings.append((tag, txt))
                self.parts.append(f"\n{'#' * int(tag[1])} {txt}\n")
            self._heading = None
            return
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._skip:
            return
        if self._in_title:
            self.title += data
        if self._link is not None:
            self._link_buf.append(data)
        if self._heading:
            self._heading_buf.append(data)
            return
        self.parts.append(data)

    def text(self) -> str:
        raw = "".join(self.parts)
        lines = [" ".join(l.split()) for l in raw.splitlines()]
        out, blank = [], False
        for l in lines:
            if not l:
                if not blank:
                    out.append("")
                blank = True
                continue
            out.append(l)
            blank = False
        return "\n".join(out).strip()


def parse_page(url: str, body: bytes) -> PageParser:
    p = PageParser(url)
    try:
        p.feed(text_of(body))
    except Exception:  # noqa: BLE001 - malformed HTML still yields what was parsed
        pass
    return p


def same_site(url: str, host: str) -> bool:
    h = urllib.parse.urlparse(url).netloc.lower()
    return h == host or h == "www." + host or "www." + h == host


SKIP_URL = re.compile(r"\.(jpg|jpeg|png|gif|webp|svg|pdf|zip|mp4|css|js|ico|xml)(\?|$)|/wp-(admin|json|content)/|"
                      r"/feed/?$|/tag/|/author/|/page/\d|[?&](s|p|replytocom)=|/cart|/checkout|/my-account|"
                      r"/wp-login|^mailto:|^tel:", re.I)


def sitemap_urls(root: str, robots: str, limit: int = 4000) -> list[str]:
    """Page URLs from the sitemaps, recursing sitemap indexes."""
    queue = re.findall(r"(?im)^\s*sitemap:\s*(\S+)", robots) or []
    queue += [urllib.parse.urljoin(root, p) for p in ("/sitemap_index.xml", "/sitemap.xml", "/wp-sitemap.xml")]
    seen, pages = set(), []
    while queue and len(seen) < 40 and len(pages) < limit:
        sm = queue.pop(0)
        if sm in seen:
            continue
        seen.add(sm)
        status, _, body, _ = fetch(sm)
        if status != 200 or not body:
            continue
        xml = text_of(body)
        locs = [html.unescape(x.strip()) for x in re.findall(r"<loc>\s*(.*?)\s*</loc>", xml, re.S)]
        if "<sitemapindex" in xml:
            locs.sort(key=lambda u: (bool(re.search(r"image|author|tag|categor|attachment", u)), u))
            queue.extend(l for l in locs if not re.search(r"image|attachment", l))
        else:
            pages.extend(locs)
        time.sleep(0.3)
    return list(dict.fromkeys(pages))[:limit]


# ------------------------------------------------------------- buy-page ranking --

BUY_KEYWORDS = ["pricing", "plans", "price", "faq", "about", "features", "product", "how-it-works",
                "case-stud", "results", "testimonial", "review", "customer", "guarantee", "refund",
                "return", "shipping", "delivery", "terms", "privacy", "security", "trust", "compliance",
                "integration", "docs", "help", "support", "contact", "community", "join", "course",
                "classroom", "compare", "vs", "alternative", "enterprise", "team", "agency", "partner",
                "who-its-for"]
BLOG_PATTERN = re.compile(r"/blog/|/news/|/articles?/|/insights/|/\d{4}/\d{2}/", re.I)


def buy_score(url: str) -> int:
    """Higher is a more likely buying page (pricing, FAQ, proof, policy, ...). Blog posts
    score low; the homepage always wins (handled by the caller, which always fetches it
    first)."""
    path = urllib.parse.urlparse(url).path.lower().rstrip("/") or "/"
    if path == "/":
        return 1000
    if BLOG_PATTERN.search(path):
        return -50
    hits = sum(1 for kw in BUY_KEYWORDS if kw in path)
    depth = len([x for x in path.split("/") if x])
    return hits * 10 - depth


# ---------------------------------------------------------------- wayback --

def wayback_snapshot_url(home: str, months: int) -> str | None:
    """The web.archive.org snapshot URL closest to `months` ago, or None (never raises;
    the caller treats a miss as a warning, not a failure)."""
    target = time.time() - months * 30 * 86400
    ts = time.strftime("%Y%m%d", time.gmtime(target))
    api = f"https://archive.org/wayback/available?url={urllib.parse.quote(home, safe='')}&timestamp={ts}"
    st, _, body, _ = fetch(api)
    if st != 200 or not body:
        return None
    try:
        data = json.loads(text_of(body))
    except json.JSONDecodeError:
        return None
    snap = (data.get("archived_snapshots") or {}).get("closest") or {}
    if snap.get("available") and snap.get("url"):
        return snap["url"]
    return None


# ------------------------------------------------------------------ crawl --

def save_page_text(pages_dir: Path, name: str, url: str, body: bytes) -> tuple[str, int, str]:
    """Write one page's visible text to data/pages/<name>.txt. Returns (title, words, rel path)."""
    p = parse_page(url, body)
    txt = p.text()
    head = [f"URL: {url}", ""]
    (pages_dir / name).write_text("\n".join(head) + txt + "\n", encoding="utf-8")
    return html.unescape(p.title.strip()), len(txt.split()), f"{DATA_DIR}/pages/{name}"


def page_name(kind: str, url: str, used: set[str]) -> str:
    base = f"{kind}-{slugify(urllib.parse.urlparse(url).path) or 'home'}"[:76]
    name, n = base + ".txt", 2
    while name in used:
        name = f"{base}-{n}.txt"
        n += 1
    used.add(name)
    return name


def cmd_crawl(args) -> int:
    start = args.url if re.match(r"https?://", args.url) else "https://" + args.url
    status, body, final, via = fetch_page(start, 0, args.scrapling)
    if status != 200:
        sys.exit(f"homepage returned {status}: {start}")
    parsed = urllib.parse.urlparse(final)
    root = f"{parsed.scheme}://{parsed.netloc}/"
    host = parsed.netloc.lower().removeprefix("www.")
    out = Path(args.out)
    pages_dir = out / DATA_DIR / "pages"
    if pages_dir.is_dir():  # a re-crawl must not leave stale pages behind
        for old in pages_dir.glob("*.txt"):
            old.unlink()
    pages_dir.mkdir(parents=True, exist_ok=True)

    _, _, rb, _ = fetch(urllib.parse.urljoin(root, "/robots.txt"))
    robots = text_of(rb)
    home = parse_page(final, body)
    nav_links = {u.split("#")[0] for u, _ in home.links if same_site(u, host)}
    smap = sitemap_urls(root, robots)
    (out / DATA_DIR).mkdir(parents=True, exist_ok=True)
    (out / DATA_DIR / "sitemap-urls.txt").write_text("\n".join(smap) + "\n", encoding="utf-8")

    candidates = {u for u in (nav_links | {u for u in smap if same_site(u, host)})
                  if not SKIP_URL.search(u)}
    candidates.discard(final)
    candidates.discard(root)
    ranked = sorted(candidates, key=lambda u: (-buy_score(u), len(u)))
    site_urls = [final] + [u for u in ranked if u.rstrip("/") != final.rstrip("/")]
    site_urls = site_urls[: max(1, args.max)]

    extra_urls = [u if re.match(r"https?://", u) else urllib.parse.urljoin(root, u) for u in (args.extra or [])]

    pages: list[dict] = []
    seen: set[str] = set()
    used_names: set[str] = set()

    def record(requested_url: str, kind: str, st: int, b: bytes, fu: str, via_: str) -> None:
        seen.add(fu.rstrip("/"))
        if st != 200 or not b:
            pages.append({"url": requested_url, "final_url": fu, "status": st, "via": via_, "kind": kind,
                          "title": "", "words": 0, "file": None})
            print(f"  [FAILED {st}] {requested_url}")
            return
        name = page_name(kind, fu, used_names)
        title, words, rel = save_page_text(pages_dir, name, fu, b)
        pages.append({"url": requested_url, "final_url": fu, "status": st, "via": via_, "kind": kind,
                      "title": title, "words": words, "file": rel})
        print(f"  [{st}] {fu}  ({words} words)" + ("" if via_ == "urllib" else f"  via {via_}"))

    record(start, "site", status, body, final, via)
    for url in site_urls[1:]:
        if url.rstrip("/") in seen:
            continue
        st, b, fu, via_ = fetch_page(url, args.delay, args.scrapling)
        if fu.rstrip("/") in seen:
            continue
        record(url, "site", st, b, fu, via_)

    for url in extra_urls:
        if url.rstrip("/") in seen:
            continue
        st, b, fu, via_ = fetch_page(url, args.delay, args.scrapling)
        record(url, "extra", st, b, fu, via_)

    if args.wayback_months and args.wayback_months > 0:
        try:
            snap = wayback_snapshot_url(root, args.wayback_months)
        except Exception as e:  # noqa: BLE001 - never fatal
            snap = None
            print(f"  [warn] wayback lookup failed: {e}")
        if snap:
            st, b, fu, via_ = fetch_page(snap, args.delay, "off")
            record(snap, "wayback", st, b, fu, via_)
        else:
            print("  [warn] no wayback snapshot found; continuing without one")

    (out / DATA_DIR / "crawl.json").write_text(
        json.dumps({"crawled_at": time.strftime("%Y-%m-%dT%H:%M:%S"), "start": start, "root": root,
                    "sitemap_urls": len(smap), "pages": pages}, indent=2, ensure_ascii=False),
        encoding="utf-8")

    counts: dict[str, int] = {}
    for p in pages:
        if p.get("file"):
            counts[p["kind"]] = counts.get(p["kind"], 0) + 1
    ok = sum(counts.values())
    print(f"\nCrawled {ok}/{len(pages)} pages from {len(smap)} sitemap URLs -> {out / DATA_DIR / 'crawl.json'}")
    print("  " + ", ".join(f"{k}: {counts.get(k, 0)}" for k in ("site", "extra", "wayback")))
    return 0


# -------------------------------------------------------------------- add --

def load_crawl(run_dir: Path) -> dict:
    path = run_dir / DATA_DIR / "crawl.json"
    if not path.is_file():
        return {"pages": []}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {"pages": []}


def cmd_add(args) -> int:
    if not args.urls and not (args.url and args.text_file):
        sys.exit("add needs URL(s) to fetch, or --url with --text-file")
    if args.text_file and not args.url:
        sys.exit("--text-file requires --url")
    run_dir = Path(args.run_dir)
    pages_dir = run_dir / DATA_DIR / "pages"
    pages_dir.mkdir(parents=True, exist_ok=True)
    crawl = load_crawl(run_dir)
    entries: list[dict] = crawl.setdefault("pages", [])
    used_names = {Path(e["file"]).name for e in entries if e.get("file")}

    def upsert(entry: dict) -> None:
        key = normalize_url(entry["url"])
        crawl["pages"] = [e for e in entries if normalize_url(e.get("url", "")) != key]
        crawl["pages"].append(entry)

    saved = 0
    if args.text_file:
        text = Path(args.text_file).read_text(encoding="utf-8")
        name = page_name("added", args.url, used_names)
        (pages_dir / name).write_text(text, encoding="utf-8")
        title = next((l.strip() for l in text.splitlines() if l.strip()), "")[:120]
        upsert({"url": args.url, "final_url": args.url, "status": 200, "via": "manual", "kind": "added",
               "title": title, "words": len(text.split()), "file": f"{DATA_DIR}/pages/{name}"})
        saved += 1
        print(f"  [manual] {args.url}  ({len(text.split())} words)")

    for url in args.urls or []:
        st, b, fu, via_ = fetch_page(url, args.delay, args.scrapling)
        if st != 200 or not b:
            print(f"  [FAILED {st}] {url}")
            upsert({"url": url, "final_url": fu, "status": st, "via": via_, "kind": "added",
                   "title": "", "words": 0, "file": None})
            continue
        name = page_name("added", fu, used_names)
        title, words, rel = save_page_text(pages_dir, name, fu, b)
        upsert({"url": url, "final_url": fu, "status": st, "via": via_, "kind": "added",
               "title": title, "words": words, "file": rel})
        saved += 1
        print(f"  [{st}] {fu}  ({words} words)")

    crawl_path = run_dir / DATA_DIR / "crawl.json"
    crawl_path.parent.mkdir(parents=True, exist_ok=True)
    crawl_path.write_text(json.dumps(crawl, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nSaved {saved} page(s) -> {crawl_path}")
    return 0


# ------------------------------------------------------------- normalisation --

def normalize_url(u: str) -> str:
    """URL identity for matching: scheme, www and a trailing slash never matter."""
    u = (u or "").strip()
    u = re.sub(r"^https?://", "", u, flags=re.I)
    u = re.sub(r"^www\.", "", u, flags=re.I)
    return u.rstrip("/").lower()


SMART_QUOTES = str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"',
                              "–": "-", "—": "-", "…": "...", " ": " "})


def normalize_text(s: str) -> str:
    """NFKC, lowercase, smart quotes/dashes to ascii, strip everything but alphanumerics
    and spaces, collapse whitespace. Used for the quote substring check. Dashes become a
    space (so "BEST-no contest" doesn't fuse into "bestno"); everything else that isn't
    alphanumeric or whitespace is simply dropped (so "we're" becomes "were", matching the
    other side of the comparison the same way)."""
    s = unicodedata.normalize("NFKC", s or "")
    s = s.translate(SMART_QUOTES)
    s = s.lower().replace("-", " ")
    s = "".join(ch for ch in s if ch.isalnum() or ch.isspace())
    return re.sub(r"\s+", " ", s).strip()


def build_url_index(run_dir: Path, pages: list[dict]) -> dict[str, str]:
    """normalized url/final_url -> normalized saved page text, for every page that was
    actually fetched (status 200, has a file)."""
    idx: dict[str, str] = {}
    cache: dict[str, str | None] = {}
    for e in pages:
        if e.get("status") != 200 or not e.get("file"):
            continue
        rel = e["file"]
        if rel not in cache:
            fpath = run_dir / rel
            cache[rel] = normalize_text(fpath.read_text(encoding="utf-8")) if fpath.is_file() else None
        text = cache[rel]
        if text is None:
            continue
        for u in (e.get("url"), e.get("final_url")):
            if u:
                idx[normalize_url(u)] = text
    return idx


# ------------------------------------------------------------------ build --

CATEGORY_EMOJI = {
    "Proof & Results": "🧪",
    "Product & Setup Requirements": "🧠",
    "Maturity & Completeness": "📚",
    "Pricing & Offer Clarity": "💰",
    "Risk & Guarantees": "🛡️",
    "Data, Security & Compliance": "🔒",
    "Team & Enterprise Adoption": "🏢",
    "Measurement & Attribution": "📈",
    "Integrations & Ecosystem": "🔌",
    "Audience Fit": "👤",
    "Support & Service": "🤝",
    "Delivery & Fulfilment": "📦",
    "Ingredients, Safety & Claims": "🧾",
    "Location & Access": "📍",
    "Category Credibility": "🎯",
}
KNOWN_CATEGORIES = set(CATEGORY_EMOJI)
MODES = {"own", "competitor"}
SEVERITIES = ["critical", "major", "minor"]
SEV_ORDER = {s: i for i, s in enumerate(SEVERITIES)}
SEV_DOT = {"critical": "🔴", "major": "🟠", "minor": "🟡"}
EVIDENCE_TYPES = {"stated", "contradicted", "absent", "inferred"}
FIX_TYPES = {"faq", "pricing-page", "proof-asset", "policy-doc", "positioning", "product", "none"}
FAQ_FIX_TYPES = {"faq", "pricing-page", "policy-doc"}
ID_RE = re.compile(r"^[a-z0-9]+(-[a-z0-9]+)*$")


def load_findings(run_dir: Path) -> dict:
    path = run_dir / DATA_DIR / "findings.json"
    if not path.is_file():
        sys.exit(f"missing {path}: write it first (shape in references/findings-schema.md)")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        sys.exit(f"{path} is not valid JSON: {e}")


def check_sources(where: str, sources, url_index: dict[str, str], errors: list[str]) -> None:
    for s in sources or []:
        url = (s.get("url") or "").strip() if isinstance(s, dict) else ""
        quote = (s.get("quote") or "") if isinstance(s, dict) else ""
        if not url:
            errors.append(f"{where}: a source has no url")
            continue
        key = normalize_url(url)
        if key not in url_index:
            errors.append(f"{where}: {url} was never saved (run: dealbreakers.py add --run-dir ... {url})")
            continue
        if not quote.strip():
            errors.append(f"{where}: source {url} has no quote")
            continue
        if normalize_text(quote) not in url_index[key]:
            errors.append(f"{where}: quote not found on {url}: {quote!r}")


def lint(findings: dict, url_index: dict[str, str]) -> tuple[list[str], list[str]]:
    """(errors, warnings) against references/findings-schema.md's build rules."""
    errors: list[str] = []
    warnings: list[str] = []

    if not (findings.get("brand") or "").strip():
        errors.append("brand is empty")
    if not (findings.get("website") or "").strip():
        errors.append("website is empty")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", findings.get("date") or ""):
        errors.append("date must be YYYY-MM-DD")
    mode = findings.get("mode")
    if mode not in MODES:
        errors.append(f"mode must be one of {sorted(MODES)}")

    if not (findings.get("offer_summary") or "").strip():
        errors.append("offer_summary is empty")
    offer_sources = findings.get("offer_sources") or []
    if not offer_sources:
        errors.append("offer_sources: at least one source required")
    check_sources("offer_sources", offer_sources, url_index, errors)

    pd = findings.get("positioning_drift")
    if pd:
        if not (pd.get("summary") or "").strip():
            errors.append("positioning_drift.summary is empty")
        pd_sources = pd.get("sources") or []
        if not pd_sources:
            errors.append("positioning_drift: at least one source required")
        check_sources("positioning_drift", pd_sources, url_index, errors)

    personas = findings.get("personas") or []
    if not (2 <= len(personas) <= 4):
        errors.append(f"personas: need 2 to 4, got {len(personas)}")
    persona_ids: set[str] = set()
    seen_persona_ids: set[str] = set()
    for p in personas:
        pid = p.get("id") or ""
        if not ID_RE.match(pid):
            errors.append(f"persona id {pid!r} is not lowercase-hyphenated")
        if pid in seen_persona_ids:
            errors.append(f"persona id {pid!r} is duplicated")
        seen_persona_ids.add(pid)
        persona_ids.add(pid)
        if not (p.get("label") or "").strip():
            errors.append(f"persona {pid!r}: label is empty")
        if not (p.get("needs") or "").strip():
            errors.append(f"persona {pid!r}: needs is empty")
        if not (p.get("source") or "").strip():
            errors.append(f"persona {pid!r}: source is empty")

    dbs = findings.get("dealbreakers") or []
    if len(dbs) < 1:
        warnings.append("no dealbreakers at all")
    seen_ids: set[str] = set()
    persona_used: set[str] = set()
    cat_counts: dict[str, dict[str, int]] = {}
    crit_total = 0
    for d in dbs:
        did = d.get("id") or ""
        where = f"dealbreakers[{did or '?'}]"
        if not ID_RE.match(did):
            errors.append(f"{where}: id {did!r} is not lowercase-hyphenated")
        if did in seen_ids:
            errors.append(f"{where}: id is duplicated")
        seen_ids.add(did)

        cat = d.get("category") or ""
        if not cat.strip():
            errors.append(f"{where}: category is empty")
        else:
            if cat not in KNOWN_CATEGORIES:
                warnings.append(f"{where}: new category name {cat!r}")
            counts = cat_counts.setdefault(cat, {"critical": 0, "major": 0, "minor": 0})
            if d.get("severity") in counts:
                counts[d["severity"]] += 1

        if not (d.get("claim") or "").strip():
            errors.append(f"{where}: claim is empty")
        if not (d.get("buyer_question") or "").strip():
            errors.append(f"{where}: buyer_question is empty")

        sev = d.get("severity")
        if sev not in SEVERITIES:
            errors.append(f"{where}: severity must be one of {SEVERITIES}")
        elif sev == "critical":
            crit_total += 1

        ev = d.get("evidence")
        sources = d.get("sources") or []
        if ev not in EVIDENCE_TYPES:
            errors.append(f"{where}: evidence must be one of {sorted(EVIDENCE_TYPES)}")
        elif ev == "stated":
            if len(sources) < 1:
                errors.append(f"{where}: stated needs at least 1 source")
        elif ev == "contradicted":
            urls = {s.get("url") for s in sources if isinstance(s, dict)}
            if len(sources) < 2 or len(urls) < 2:
                errors.append(f"{where}: contradicted needs 2+ sources with different URLs")
        elif ev == "absent":
            checked = d.get("checked") or []
            if len(checked) < 2:
                errors.append(f"{where}: absent needs 2+ entries in checked")
            for c in checked:
                if not str(c).startswith("search:") and normalize_url(str(c)) not in url_index:
                    errors.append(f"{where}: checked URL {c} was never saved "
                                  f"(run: dealbreakers.py add --run-dir ... {c})")
        elif ev == "inferred":
            if not (d.get("reason") or "").strip():
                errors.append(f"{where}: inferred needs reason")
            if sev == "critical":
                errors.append(f"{where}: inferred cannot be critical")
        check_sources(where, sources, url_index, errors)

        d_personas = d.get("personas") or []
        if not d_personas:
            errors.append(f"{where}: needs at least one persona")
        for pid in d_personas:
            if pid not in persona_ids:
                errors.append(f"{where}: persona {pid!r} is not defined in personas")
            persona_used.add(pid)

        fix = d.get("fix") or {}
        ftype = fix.get("type")
        if ftype not in FIX_TYPES:
            errors.append(f"{where}: fix.type must be one of {sorted(FIX_TYPES)}")
        elif ftype != "none" and not (fix.get("action") or "").strip():
            errors.append(f"{where}: fix.action is required unless fix.type is 'none'")
        answer = (fix.get("answer") or "").strip()
        answer_sources = fix.get("answer_sources") or []
        if answer and not answer_sources:
            errors.append(f"{where}: fix.answer without answer_sources")
        if answer_sources:
            check_sources(f"{where}: fix.answer_sources", answer_sources, url_index, errors)

    for pid in persona_ids:
        if pid not in persona_used:
            errors.append(f"persona {pid!r} is not used by any finding")

    handled = findings.get("handled") or []
    if not handled:
        errors.append("handled: at least one entry required")
    for h in handled:
        obj = h.get("objection") or ""
        if not obj.strip():
            errors.append("handled: objection is empty")
        if not (h.get("how") or "").strip():
            errors.append(f"handled {obj!r}: how is empty")
        hs = h.get("sources") or []
        if not hs:
            errors.append(f"handled {obj!r}: at least one source required")
        check_sources(f"handled {obj!r}", hs, url_index, errors)

    priorities = findings.get("priorities") or []
    if not (3 <= len(priorities) <= 7):
        errors.append(f"priorities: need 3 to 7 ids, got {len(priorities)}")
    dbs_by_id = {d.get("id"): d for d in dbs}
    for pid in priorities:
        if pid not in dbs_by_id:
            errors.append(f"priorities: {pid!r} does not exist")
        elif dbs_by_id[pid].get("severity") == "minor":
            errors.append(f"priorities: {pid!r} is minor; only critical or major allowed")

    if not (findings.get("not_prioritised") or "").strip():
        errors.append("not_prioritised is empty")

    # Warnings (never block).
    total = len(dbs)
    if total and crit_total / total > 0.25:
        warnings.append(f"more than 25% critical ({crit_total}/{total})")
    real_cats = [c for c in cat_counts if c]
    if len(real_cats) < 6:
        warnings.append(f"fewer than 6 categories ({len(real_cats)})")
    if len(real_cats) > 10:
        warnings.append(f"more than 10 categories ({len(real_cats)})")
    if total < 12:
        warnings.append(f"fewer than 12 findings ({total})")
    for cat in real_cats:
        n = sum(cat_counts[cat].values())
        if n == 1:
            warnings.append(f"category {cat!r} has only one finding")

    return errors, warnings


# --------------------------------------------------------------- rendering --

def short_link(url: str) -> str:
    p = urllib.parse.urlparse(url)
    host = p.netloc.lower().removeprefix("www.")
    path = p.path.rstrip("/")
    return f"{host}{path}" if path else host


def md_link(url: str) -> str:
    return f"[{short_link(url)}]({url})"


def sources_suffix(sources: list[dict]) -> str:
    """One ` ([host+path](url))` per source, concatenated."""
    return "".join(f" ({md_link(s['url'])})" for s in sources or [] if s.get("url"))


def evidence_suffix(d: dict) -> str:
    ev = d.get("evidence")
    if ev in ("stated", "contradicted"):
        return sources_suffix(d.get("sources") or [])
    if ev == "absent":
        n = len(d.get("checked") or [])
        return f" _(not found on the {n} public pages checked)_"
    if ev == "inferred":
        return f" _(inferred: {(d.get('reason') or '').strip()})_"
    return ""


def category_order(by_cat: dict[str, list[dict]]) -> list[str]:
    def counts(cat: str) -> tuple[int, int]:
        crit = sum(1 for d in by_cat[cat] if d.get("severity") == "critical")
        maj = sum(1 for d in by_cat[cat] if d.get("severity") == "major")
        return crit, maj

    known = [c for c in by_cat if c in KNOWN_CATEGORIES and c != "Category Credibility"]
    unknown = [c for c in by_cat if c not in KNOWN_CATEGORIES]
    known.sort(key=lambda c: (-counts(c)[0], -counts(c)[1]))
    unknown.sort(key=lambda c: (-counts(c)[0], -counts(c)[1]))
    order = known + unknown
    if "Category Credibility" in by_cat:
        order.append("Category Credibility")
    return order


def crawl_counts(crawl: dict) -> dict[str, int]:
    counts: dict[str, int] = {}
    for p in crawl.get("pages", []):
        if p.get("status") == 200 and p.get("file"):
            counts[p.get("kind", "site")] = counts.get(p.get("kind", "site"), 0) + 1
    return counts


def render_report(findings: dict, crawl: dict) -> str:
    brand = findings["brand"]
    counts = crawl_counts(crawl)
    n_total = sum(counts.values())
    breakdown = " / ".join(f"{k} {counts.get(k, 0)}" for k in ("site", "extra", "added", "wayback"))
    out = [f"# Dealbreaker Detector: {brand}", "",
           f"> {findings.get('date', '')} · {findings.get('mode', '')} · {n_total} pages read "
           f"({breakdown}) · {findings.get('website', '')}", ""]

    out += ["## What a buyer sees today", "",
           (findings.get("offer_summary", "") or "").strip() + sources_suffix(findings.get("offer_sources") or [])]
    out.append("")
    pd = findings.get("positioning_drift")
    if pd:
        out.append((pd.get("summary") or "").strip() + sources_suffix(pd.get("sources") or []))
        out.append("")

    out += ["## Who's buying", ""]
    for p in findings.get("personas") or []:
        out.append(f"- **{p['label']}**: {p['needs']} ([{short_link(p['source'])}]({p['source']}))")
    out.append("")

    dbs = findings.get("dealbreakers") or []
    dbs_by_id = {d.get("id"): d for d in dbs}
    out += ["## Fix these first", ""]
    for i, pid in enumerate(findings.get("priorities") or [], start=1):
        d = dbs_by_id.get(pid)
        if not d:
            continue
        dot = SEV_DOT.get(d.get("severity"), "⚪")
        fix_action = (d.get("fix") or {}).get("action", "")
        out.append(f"{i}. **{d.get('claim', '')}** — {dot} {d.get('severity', '')} — Fix: {fix_action}")
    out.append("")
    out.append(f"*{(findings.get('not_prioritised') or '').strip()}*")
    out.append("")

    by_cat: dict[str, list[dict]] = {}
    for d in dbs:
        by_cat.setdefault(d.get("category") or "", []).append(d)
    for cat in category_order(by_cat):
        emoji = CATEGORY_EMOJI.get(cat, "•")
        out += [f"## {emoji} {cat}", ""]
        for d in sorted(by_cat[cat], key=lambda d: SEV_ORDER.get(d.get("severity"), 3)):
            dot = SEV_DOT.get(d.get("severity"), "⚪")
            if cat == "Category Credibility":
                line = f"- {dot} **\"{d.get('buyer_question', '')}\"** {d.get('claim', '')}"
            else:
                detail = f" {d['detail'].strip()}" if d.get("detail") else ""
                line = f"- {dot} **{d.get('claim', '')}**{detail}"
            out.append(line + evidence_suffix(d))
        out.append("")

    out += ["## 🟢 Already handled", ""]
    for h in findings.get("handled") or []:
        out.append(f"- **{h.get('objection', '')}** — {h.get('how', '')}{sources_suffix(h.get('sources') or [])}")
    out.append("")

    out += ["## Method and limits", "",
           f"Pages read: {counts.get('site', 0)} from the site, {counts.get('extra', 0)} extra, "
           f"{counts.get('added', 0)} added, {counts.get('wayback', 0)} wayback.", "",
           "- **Stated** — the brand's own public words create this objection.",
           "- **Contradicted** — two public sources disagree.",
           "- **Absent** — a buyer needs this and the public pages checked don't answer it.",
           "- **Inferred** — a reasonable buyer concern given how the offer works, with no page to point to.",
           "",
           f"Public information only, as of {findings.get('date', '')}. 'Not found' means a buyer could not "
           "verify it from these pages, not that it doesn't exist.", "",
           "Dealbreaker Detector is a method coined by Steve Toth (Notebook Agency). This skill is an "
           "independent implementation, not affiliated with or endorsed by them.", ""]

    return "\n".join(out).rstrip() + "\n"


def render_faq(findings: dict) -> str:
    brand = findings["brand"]
    priorities = findings.get("priorities") or []
    prio_rank = {pid: i for i, pid in enumerate(priorities)}
    dbs = [d for d in findings.get("dealbreakers") or [] if (d.get("fix") or {}).get("type") in FAQ_FIX_TYPES]
    dbs.sort(key=lambda d: (prio_rank.get(d.get("id"), len(priorities)), SEV_ORDER.get(d.get("severity"), 3)))
    out = [f"# FAQ draft: {brand}", "",
           "Answers marked [CLIENT TO ANSWER] must come from the client before this goes live.", ""]
    for d in dbs:
        out.append(f"### {d.get('buyer_question', '')}")
        out.append("")
        fix = d.get("fix") or {}
        answer = (fix.get("answer") or "").strip()
        if answer:
            out.append(answer + sources_suffix(fix.get("answer_sources") or []))
        else:
            out.append(f"[CLIENT TO ANSWER: {fix.get('action', '')}]")
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def write_findings_csv(run_dir: Path, findings: dict) -> int:
    priorities = findings.get("priorities") or []
    prio_rank = {pid: i + 1 for i, pid in enumerate(priorities)}
    dbs = findings.get("dealbreakers") or []
    path = run_dir / FINDINGS_CSV
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["id", "priority", "severity", "category", "claim", "buyer_question", "evidence",
                    "sources", "personas", "fix_type", "fix_action"])
        for d in dbs:
            fix = d.get("fix") or {}
            w.writerow([d.get("id"), prio_rank.get(d.get("id"), ""), d.get("severity"), d.get("category"),
                        d.get("claim"), d.get("buyer_question"), d.get("evidence"),
                        " | ".join(s.get("url", "") for s in d.get("sources") or []),
                        "; ".join(d.get("personas") or []), fix.get("type"), fix.get("action")])
    return len(dbs)


def cmd_build(args) -> int:
    run_dir = Path(args.run_dir)
    findings = load_findings(run_dir)
    crawl = load_crawl(run_dir)
    url_index = build_url_index(run_dir, crawl.get("pages", []))
    errors, warnings = lint(findings, url_index)
    for w in warnings:
        print(f"  [WARN] {w}")
    for e in errors:
        print(f"  [FAIL] {e}")
    if errors:
        print(f"\n{len(errors)} problem(s), {len(warnings)} warning(s) in data/findings.json. "
              "Nothing written; fix them and run build again.")
        return 1
    report = render_report(findings, crawl)
    faq = render_faq(findings)
    (run_dir / REPORT).write_text(report, encoding="utf-8")
    (run_dir / FAQ_DRAFT).write_text(faq, encoding="utf-8")
    n = write_findings_csv(run_dir, findings)
    print(f"\nBuilt {REPORT}, {FAQ_DRAFT}, {n} row(s) in {FINDINGS_CSV}. {len(warnings)} warning(s).")
    return 0


# -------------------------------------------------------------- workbook --

def cmd_workbook(args) -> int:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        sys.exit("openpyxl is needed: uv run --with openpyxl python dealbreakers.py workbook ...")
    run_dir = Path(args.run_dir).resolve()
    findings = load_findings(run_dir)
    if not (run_dir / REPORT).is_file():
        sys.exit("run build first: the workbook is filled from its outputs")
    book = run_dir.parent / WORKBOOK
    if not book.is_file():
        if not PACK_TEMPLATE.is_file():
            sys.exit(f"no {WORKBOOK} in {run_dir.parent} and no pack template at {PACK_TEMPLATE}")
        book.write_bytes(PACK_TEMPLATE.read_bytes())
    wb = openpyxl.load_workbook(book)
    day = time.strftime("%Y-%m-%d")

    ws = wb["Checklist"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=3).value == SKILL_ID:
            ws.cell(row=r, column=8, value="Client review")
            ws.cell(row=r, column=9, value="☑")
            ws.cell(row=r, column=10, value=day)
            ws.cell(row=r, column=11, value=f"report + FAQ draft in {run_dir.name}/")
            break
    else:
        sys.exit(f"Checklist has no {SKILL_ID} row; rebuild the template from the pack (build_template.py)")

    head = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F2937")
    wrap = Alignment(wrap_text=True, vertical="top")
    edge = Border(*(Side(style="thin", color="D1D5DB"),) * 4)
    brand = findings.get("brand", "")
    mode = findings.get("mode", "own")
    title = "Dealbreakers" if mode == "own" else f"Dealbreakers - {brand}"[:31]

    kept: dict[str, list] = {}
    for name in list(wb.sheetnames):
        if name == "Dealbreakers" or name.startswith("Dealbreakers - "):
            old = wb[name]
            for r in range(1, old.max_row + 1):
                did = old.cell(row=r, column=2).value
                if did:
                    kept[str(did)] = [old.cell(row=r, column=c).value for c in (12, 13)]
            del wb[name]

    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 14, "C": 9, "D": 10, "E": 20, "F": 45, "G": 45, "H": 16, "I": 40,
             "J": 14, "K": 45, "L": 16, "M": 30}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws["B2"] = f"Dealbreakers: what stops a qualified buyer from choosing {brand}"
    ws["B2"].font = Font(bold=True, size=16, color="1F2937")
    ws["B3"] = (f"Full write-up in {run_dir.name}/{REPORT}; FAQ draft in {run_dir.name}/{FAQ_DRAFT}. "
               "Client: mark each row with a verdict.")
    top = 6
    headers = ["ID", "Priority", "Severity", "Category", "Dealbreaker", "Buyer's question", "Evidence",
              "Source", "Fix type", "Suggested fix", "Client verdict", "Client notes"]
    for i, h in enumerate(headers, start=2):
        c = ws.cell(row=top, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = head, head_fill, wrap, edge
    ws.freeze_panes = ws.cell(row=top + 1, column=2)

    priorities = findings.get("priorities") or []
    prio_rank = {pid: i + 1 for i, pid in enumerate(priorities)}
    dbs = findings.get("dealbreakers") or []
    n = 0
    for row, d in enumerate(dbs, start=top + 1):
        n += 1
        fix = d.get("fix") or {}
        vals = [d.get("id"), prio_rank.get(d.get("id"), ""), d.get("severity"), d.get("category"),
               d.get("claim"), d.get("buyer_question"), d.get("evidence"),
               " | ".join(s.get("url", "") for s in d.get("sources") or []),
               fix.get("type"), fix.get("action"), *kept.get(str(d.get("id")), [None, None])]
        for i, v in enumerate(vals, start=2):
            c = ws.cell(row=row, column=i, value=v)
            c.alignment, c.border = wrap, edge
    last = top + max(n, 1)
    dv = DataValidation(type="list", formula1='"Agree,Partly agree,Disagree,Already fixed"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"L{top + 1}:L{last}")
    ws.auto_filter.ref = f"B{top}:M{last}"

    if mode == "own":
        ws2 = wb["Initiatives"]
        existing_tasks = {ws2.cell(row=r, column=4).value for r in range(1, ws2.max_row + 1)}
        dbs_by_id = {d.get("id"): d for d in dbs}
        ease = {"faq": 2, "pricing-page": 3, "policy-doc": 4, "positioning": 5, "proof-asset": 6,
               "product": 8, "none": 1}
        confidence = {"stated": 8, "contradicted": 8, "absent": 6, "inferred": 4}
        content_types = {"faq", "pricing-page", "proof-asset", "positioning"}
        for pid in priorities:
            d = dbs_by_id.get(pid)
            if not d:
                continue
            claim = (d.get("claim") or "").strip()
            task = f"Answer dealbreaker: {claim}"
            if len(task) > 120:
                task = task[:117].rstrip() + "..."
            if task in existing_tasks:
                continue
            fix = d.get("fix") or {}
            ftype = fix.get("type", "none")
            category = "Content Strategy" if ftype in content_types else "Brand Optimisation"
            impact = 9 if d.get("severity") == "critical" else 6
            deliverable = f"{run_dir.name}/{FAQ_DRAFT}" if ftype == "faq" else f"{run_dir.name}/{REPORT}"
            r = next(r for r in range(6, ws2.max_row + 2) if not ws2.cell(row=r, column=4).value)
            for col, v in {3: category, 4: task, 5: fix.get("action", ""), 6: SKILL_ID, 7: "Scheduled",
                          8: "Yes" if ftype == "pricing-page" else "No", 10: impact,
                          11: confidence.get(d.get("evidence"), 4), 12: ease.get(ftype, 4),
                          13: deliverable}.items():
                ws2.cell(row=r, column=col, value=v)
            existing_tasks.add(task)

    fixed = ["Checklist", "Brand Truth Review", "Brand 360 Report", "AI Visibility", "AI Info Page"]
    dealbreaker_sheets = [s for s in wb.sheetnames if s == "Dealbreakers" or s.startswith("Dealbreakers - ")]
    ordered = [s for s in fixed if s in wb.sheetnames] + dealbreaker_sheets \
        + ([s for s in ["Initiatives"] if s in wb.sheetnames])
    wb._sheets = [wb[n] for n in ordered] + [s for s in wb._sheets if s.title not in ordered]
    wb.save(book)
    print(f"Wrote {book}: Checklist ticked, {n} finding(s) on '{title}'"
         + (", Initiatives updated" if mode == "own" else ", no Initiatives (competitor mode)"))
    return 0


# ------------------------------------------------------------------ main --

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("path", help="print the run folder for this brand (MarketingOS-aware)")
    p.add_argument("--brand", required=True)
    p.add_argument("--date", help="YYYY-MM-DD (default today)")
    p.add_argument("--start", help="where to look for the brain (default cwd)")
    p.set_defaults(fn=cmd_path)

    p = sub.add_parser("crawl", help="fetch the brand's public pages into data/pages/ and data/crawl.json")
    p.add_argument("--url", required=True)
    p.add_argument("--out", required=True, help="the run folder")
    p.add_argument("--max", type=int, default=40, help="site pages to fetch (homepage always first)")
    p.add_argument("--delay", type=float, default=0.6, help="seconds between requests")
    p.add_argument("--extra", action="append", help="extra URL to always fetch, any host (repeatable)")
    p.add_argument("--wayback-months", type=int, default=12, dest="wayback_months",
                  help="months back for a homepage Wayback snapshot (0 disables)")
    p.add_argument("--scrapling", choices=["auto", "always", "off"], default="auto")
    p.set_defaults(fn=cmd_crawl)

    p = sub.add_parser("add", help="fetch more pages, or save text fetched another way")
    p.add_argument("--run-dir", required=True)
    p.add_argument("urls", nargs="*", help="URL(s) to fetch and save (kind 'added')")
    p.add_argument("--url", help="the page's URL, used with --text-file")
    p.add_argument("--text-file", help="local file with the page's already-fetched text (via 'manual')")
    p.add_argument("--delay", type=float, default=0.6)
    p.add_argument("--scrapling", choices=["auto", "always", "off"], default="auto")
    p.set_defaults(fn=cmd_add)

    p = sub.add_parser("build", help="lint data/findings.json and render the report, FAQ draft and CSV")
    p.add_argument("--run-dir", required=True)
    p.set_defaults(fn=cmd_build)

    p = sub.add_parser("workbook", help="fill the brand audit workbook (needs openpyxl)")
    p.add_argument("--run-dir", required=True)
    p.set_defaults(fn=cmd_workbook)

    args = ap.parse_args()
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(main())
