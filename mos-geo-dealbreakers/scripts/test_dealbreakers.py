"""Offline tests for dealbreakers.py. No network.

Run: python3 -m unittest scripts/test_dealbreakers.py  (from the skill folder)
The workbook tests need openpyxl: uv run --with openpyxl python -m unittest scripts/test_dealbreakers.py
"""

import copy
import json
import shutil
import sys
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).parent))
import dealbreakers as d  # noqa: E402

FIXTURE_RUN = Path(__file__).parent / "fixtures" / "run"


def load_findings() -> dict:
    return json.loads((FIXTURE_RUN / "data" / "findings.json").read_text(encoding="utf-8"))


def make_run(tmp: str, findings: dict | None = None) -> Path:
    run = Path(tmp) / "run"
    shutil.copytree(FIXTURE_RUN / "data", run / "data")
    if findings is not None:
        (run / "data" / "findings.json").write_text(json.dumps(findings), encoding="utf-8")
    return run


def run_build(tmp: str, findings: dict) -> tuple[int, str, Path]:
    run = make_run(tmp, findings)
    with mock.patch("builtins.print") as out:
        rc = d.cmd_build(Namespace(run_dir=str(run)))
    printed = "\n".join(str(c.args[0]) for c in out.call_args_list if c.args)
    return rc, printed, run


def find_db(findings: dict, fid: str) -> dict:
    return next(x for x in findings["dealbreakers"] if x["id"] == fid)


class QuoteNormalisation(unittest.TestCase):
    def test_smart_quotes_whitespace_and_case(self):
        self.assertEqual(d.normalize_text("  We’re  the   BEST—no contest.  "),
                         "were the best no contest")
        self.assertEqual(d.normalize_text('He said "hello world," loudly.'), "he said hello world loudly")
        self.assertEqual(d.normalize_text("Café Society"), "café society")  # nbsp acts like a space
        self.assertEqual(d.normalize_text("A"), d.normalize_text("a"))

    def test_substring_match_after_normalisation(self):
        page = d.normalize_text("Enterprise pricing is available on request.")
        self.assertIn(d.normalize_text("Enterprise pricing is available on request"), page)
        self.assertIn(d.normalize_text('“Enterprise pricing” is available — on request'), page)


class UrlTolerance(unittest.TestCase):
    def test_normalize_url_ignores_scheme_www_and_slash(self):
        variants = ["https://example.com/pricing/", "http://example.com/pricing", "https://www.example.com/pricing/",
                   "www.example.com/pricing", "EXAMPLE.COM/pricing/"]
        self.assertEqual(len({d.normalize_url(v) for v in variants}), 1)

    def test_url_index_keys_on_both_url_and_final_url(self):
        with tempfile.TemporaryDirectory() as tmp:
            run = make_run(tmp)
            pages = json.loads((run / "data" / "crawl.json").read_text())["pages"]
            idx = d.build_url_index(run, pages)
            self.assertIn(d.normalize_url("https://example.com/pricing/"), idx)
            self.assertIn(d.normalize_url("http://www.example.com/pricing"), idx)

    def test_build_accepts_url_variants_in_sources(self):
        f = load_findings()
        f["offer_sources"][1]["url"] = "http://www.example.com/pricing"  # variant of the saved URL
        with tempfile.TemporaryDirectory() as tmp:
            rc, out, _ = run_build(tmp, f)
        self.assertEqual(rc, 0, out)

    def test_unknown_url_fails_with_add_hint(self):
        f = load_findings()
        f["offer_sources"][0]["url"] = "https://example.com/never-saved/"
        with tempfile.TemporaryDirectory() as tmp:
            rc, out, _ = run_build(tmp, f)
        self.assertEqual(rc, 1)
        self.assertIn("was never saved", out)
        self.assertIn("dealbreakers.py add --run-dir ... https://example.com/never-saved/", out)


class CleanFixturePasses(unittest.TestCase):
    def test_clean_fixture_builds_with_no_warnings(self):
        with tempfile.TemporaryDirectory() as tmp:
            run = make_run(tmp)
            with mock.patch("builtins.print"):
                rc = d.cmd_build(Namespace(run_dir=str(run)))
            self.assertEqual(rc, 0)
            errors, warnings = d.lint(load_findings(), d.build_url_index(run, json.loads(
                (run / "data" / "crawl.json").read_text())["pages"]))
            self.assertEqual(errors, [])
            self.assertEqual(warnings, [])


class BuildRules(unittest.TestCase):
    """Each of build rules 1 to 8 in references/findings-schema.md, failing on a
    mutated copy of the (otherwise clean) fixture."""

    def test_rule1_enums_and_duplicate_ids(self):
        f = load_findings()
        f["mode"] = "bogus"
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("mode must be one of", out)

        f = load_findings()
        f["dealbreakers"][1]["id"] = f["dealbreakers"][0]["id"]
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("id is duplicated", out)

    def test_rule2_quote_must_appear_on_the_saved_page(self):
        f = load_findings()
        find_db(f, "pricing-01")["sources"][0]["quote"] = "this sentence is nowhere on the pricing page"
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("quote not found on https://example.com/pricing/", out)

    def test_rule3_stated_and_contradicted_source_counts(self):
        f = load_findings()
        find_db(f, "pricing-01")["sources"] = []
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("stated needs at least 1 source", out)

        f = load_findings()
        find_db(f, "product-01")["sources"] = find_db(f, "product-01")["sources"][:1]
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("contradicted needs 2+ sources with different URLs", out)

    def test_rule4_absent_needs_two_checked_saved_pages(self):
        f = load_findings()
        find_db(f, "proof-01")["checked"] = ["https://example.com/"]
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("absent needs 2+ entries in checked", out)

        f = load_findings()
        find_db(f, "proof-01")["checked"] = ["https://example.com/", "https://example.com/never-checked/"]
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("checked URL https://example.com/never-checked/ was never saved", out)

    def test_rule5_inferred_needs_reason_and_cannot_be_critical(self):
        f = load_findings()
        del find_db(f, "proof-02")["reason"]
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("inferred needs reason", out)

        f = load_findings()
        find_db(f, "proof-02")["severity"] = "critical"
        f["priorities"] = ["proof-02" if p == "product-01" else p for p in f["priorities"]]
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("inferred cannot be critical", out)

    def test_rule6_persona_ids_must_exist_and_be_used(self):
        f = load_findings()
        find_db(f, "proof-01")["personas"] = ["nobody-defined-this"]
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("persona 'nobody-defined-this' is not defined in personas", out)

        f = load_findings()
        f["personas"].append({"id": "unused-persona", "label": "Unused", "needs": "Never referenced.",
                              "source": "https://example.com/"})
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("persona 'unused-persona' is not used by any finding", out)

    def test_rule7_fix_answer_needs_answer_sources(self):
        f = load_findings()
        find_db(f, "support-01")["fix"]["answer"] = "We reply within one business day."
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("fix.answer without answer_sources", out)

    def test_rule8_priorities_count_and_no_minor(self):
        f = load_findings()
        f["priorities"] = f["priorities"][:2]
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("priorities: need 3 to 7 ids", out)

        f = load_findings()
        f["priorities"].append("risk-02")  # minor
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 1)
        self.assertIn("priorities: 'risk-02' is minor", out)


class Warnings(unittest.TestCase):
    def test_new_category_name_warns_not_fails(self):
        f = load_findings()
        find_db(f, "risk-02")["category"] = "Delivery Speed"
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 0, out)
        self.assertIn("new category name 'Delivery Speed'", out)

    def test_more_than_a_quarter_critical_warns(self):
        f = load_findings()
        for fid in ("pricing-02", "risk-01", "product-01"):
            find_db(f, fid)["severity"] = "critical"
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 0, out)
        self.assertIn("more than 25% critical", out)

    def test_category_with_one_finding_warns(self):
        f = load_findings()
        f["dealbreakers"] = [x for x in f["dealbreakers"] if x["id"] != "risk-02"]
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 0, out)
        self.assertIn("category 'Risk & Guarantees' has only one finding", out)

    def test_fewer_than_twelve_findings_warns(self):
        f = load_findings()
        keep_ids = {"pricing-01", "proof-01", "product-01", "support-01", "credibility-01"}
        f["dealbreakers"] = [x for x in f["dealbreakers"] if x["id"] in keep_ids]
        for p in f["personas"]:
            pass
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 0, out)
        self.assertIn("fewer than 12 findings", out)

    def test_fewer_than_six_categories_warns(self):
        f = load_findings()
        for x in f["dealbreakers"]:
            if x["category"] not in ("Proof & Results", "Pricing & Offer Clarity"):
                x["category"] = "Proof & Results" if x["severity"] != "minor" else "Pricing & Offer Clarity"
                x.setdefault("personas", f["personas"][0]["id"])
        rc, out, _ = run_build(tempfile.mkdtemp(), f)
        self.assertEqual(rc, 0, out)
        self.assertIn("fewer than 6 categories", out)


class Rendering(unittest.TestCase):
    def test_report_faq_and_csv_render(self):
        with tempfile.TemporaryDirectory() as tmp:
            run = make_run(tmp)
            with mock.patch("builtins.print"):
                rc = d.cmd_build(Namespace(run_dir=str(run)))
            self.assertEqual(rc, 0)
            report = (run / d.REPORT).read_text(encoding="utf-8")
            faq = (run / d.FAQ_DRAFT).read_text(encoding="utf-8")
            csv_text = (run / d.FINDINGS_CSV).read_text(encoding="utf-8")

        self.assertTrue(report.startswith("# Dealbreaker Detector: Acme Widgets"))
        self.assertIn("· own · 6 pages read (site 4 / extra 1 / added 0 / wayback 1) · https://example.com", report)
        self.assertIn("## What a buyer sees today", report)
        self.assertIn("## Who's buying", report)
        self.assertIn("- **Agency owner**: Needs proof that time tracking replaces manual timesheets before "
                      "switching. ([example.com](https://example.com/))", report)
        self.assertIn("## Fix these first", report)
        self.assertIn("1. **Enterprise pricing is hidden behind a sales call.** — 🔴 critical — "
                      "Fix: Publish an indicative Enterprise price band.", report)
        self.assertIn("*Generic 'is this another subscription' objections are left off", report)
        # Category ordering: Category Credibility is always last among the sections present.
        cat_idx = {name: report.index(f"## {emoji} {name}") for name, emoji in
                  (("Proof & Results", "🧪"), ("Category Credibility", "🎯"))}
        self.assertLess(cat_idx["Proof & Results"], cat_idx["Category Credibility"])
        self.assertIn('- 🔴 **"Show me the actual reduction in hours logged, not just \'automatic tracking\'."** '
                      "There's no quantified before/after for hours saved.", report)
        self.assertIn("_(not found on the 2 public pages checked)_", report)
        self.assertIn("_(inferred: No dated before/after data is shown anywhere on the public pages.)_", report)
        self.assertIn("## 🟢 Already handled", report)
        self.assertIn("- **Do they lock you into an annual contract?** — Plans are billed monthly with no "
                      "lock-in mentioned anywhere. ([example.com/pricing](https://example.com/pricing/))", report)
        self.assertIn("## Method and limits", report)
        self.assertIn("Dealbreaker Detector is a method coined by Steve Toth (Notebook Agency). This skill is "
                      "an independent implementation, not affiliated with or endorsed by them.", report)

        self.assertTrue(faq.startswith("# FAQ draft: Acme Widgets"))
        self.assertIn("[CLIENT TO ANSWER: Publish an indicative Enterprise price band.]", faq)
        # Ordered by priority then severity: pricing-01 (prio 1, faq-type fix) precedes support-01 (prio 4).
        self.assertLess(faq.index("What does the Enterprise plan actually cost?"),
                        faq.index("How fast do you respond to a support ticket?"))
        self.assertNotIn("Ship a CSV import", faq)  # fix.type 'product' never appears in the FAQ draft

        lines = csv_text.splitlines()
        self.assertEqual(lines[0], "id,priority,severity,category,claim,buyer_question,evidence,sources,"
                                   "personas,fix_type,fix_action")
        self.assertEqual(len(lines), 15)  # header + 14 findings
        self.assertIn("proof-01,2,critical,Proof & Results", csv_text)


class Paths(unittest.TestCase):
    def test_run_folder_rules(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp).resolve()  # run_dir_for resolves its start path (symlinked /tmp on macOS)
            self.assertEqual(d.run_dir_for("Acme Widgets", "2026-09-17", root),
                             root / "outputs/geo/2026-09/acme-widgets/dealbreakers")
            brain = root / "brain"
            (brain / ".mos").mkdir(parents=True)
            (brain / ".mos" / "config.yaml").write_text("mode: in-house\n")
            first = d.run_dir_for("Acme", "2026-09-17", brain)
            self.assertEqual(first, brain / "campaigns/geo/2026-09/dealbreakers")
            first.mkdir(parents=True)
            (first / "x").write_text("x")
            self.assertEqual(d.run_dir_for("Acme", "2026-09-17", brain), brain / "campaigns/geo/2026-09/dealbreakers-2")
            (brain / ".mos" / "config.yaml").write_text("mode: agency\n")
            self.assertEqual(d.run_dir_for("Acme", "2026-09-17", brain), brain / "campaigns/geo/2026-09/acme/dealbreakers")


class Add(unittest.TestCase):
    def test_text_file_is_saved_as_manual(self):
        with tempfile.TemporaryDirectory() as tmp:
            run = Path(tmp) / "run"
            (run / "data").mkdir(parents=True)
            text_file = Path(tmp) / "scraped.txt"
            text_file.write_text("This community post says Acme Widgets support replies within 24 hours.",
                                 encoding="utf-8")
            with mock.patch("builtins.print"):
                rc = d.cmd_add(Namespace(run_dir=str(run), urls=[], url="https://community.example/thread/1",
                                         text_file=str(text_file), delay=0, scrapling="off"))
            self.assertEqual(rc, 0)
            crawl = json.loads((run / "data" / "crawl.json").read_text())
            entry = next(e for e in crawl["pages"] if e["url"] == "https://community.example/thread/1")
            self.assertEqual(entry["kind"], "added")
            self.assertEqual(entry["via"], "manual")
            self.assertEqual(entry["status"], 200)
            saved = (run / entry["file"]).read_text(encoding="utf-8")
            self.assertIn("support replies within 24 hours", saved)

    def test_batch_of_urls_registers_every_page(self):
        # Regression: upsert filtered a stale snapshot, so only the last URL of a batch survived.
        html = b"<html><head><title>T</title></head><body><p>Some page text here.</p></body></html>"
        with tempfile.TemporaryDirectory() as tmp:
            run = Path(tmp) / "run"
            (run / "data").mkdir(parents=True)
            urls = ["https://a.example/one", "https://b.example/two", "https://c.example/three"]
            fake = lambda url, delay, scrapling: (200, html, url, "urllib")
            with mock.patch.object(d, "fetch_page", side_effect=fake), mock.patch("builtins.print"):
                d.cmd_add(Namespace(run_dir=str(run), urls=urls, url=None, text_file=None,
                                    delay=0, scrapling="off"))
            crawl = json.loads((run / "data" / "crawl.json").read_text())
            self.assertEqual(sorted(e["url"] for e in crawl["pages"]), sorted(urls))

    def test_missing_url_with_text_file_exits(self):
        with tempfile.TemporaryDirectory() as tmp:
            run = Path(tmp) / "run"
            (run / "data").mkdir(parents=True)
            with self.assertRaises(SystemExit):
                d.cmd_add(Namespace(run_dir=str(run), urls=[], url=None, text_file="whatever.txt",
                                    delay=0, scrapling="off"))


@unittest.skipUnless(__import__("importlib").util.find_spec("openpyxl"), "needs openpyxl")
class Workbook(unittest.TestCase):
    def test_workbook_ticks_checklist_and_writes_dealbreakers_tab(self):
        import openpyxl
        if not d.PACK_TEMPLATE.is_file():
            self.skipTest("pack template not built")
        with tempfile.TemporaryDirectory() as tmp:
            run = Path(tmp) / "dealbreakers"
            shutil.copytree(FIXTURE_RUN / "data", run / "data")
            with mock.patch("builtins.print"):
                self.assertEqual(d.cmd_build(Namespace(run_dir=str(run))), 0)
                self.assertEqual(d.cmd_workbook(Namespace(run_dir=str(run))), 0)
            book = Path(tmp) / d.WORKBOOK
            wb = openpyxl.load_workbook(book)
            rows = [r for r in wb["Checklist"].iter_rows(values_only=True) if r[2] == "mos-geo-dealbreakers"]
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0][7], "Client review")
            self.assertEqual(rows[0][8], "☑")
            tab = wb["Dealbreakers"]
            self.assertEqual(tab.cell(row=7, column=2).value, "proof-01")
            tab.cell(row=7, column=12, value="Agree")  # client verdict
            tab.cell(row=7, column=13, value="Confirmed with the client")
            wb.save(book)
            initiatives_before = [r[3] for r in wb["Initiatives"].iter_rows(min_row=6, values_only=True) if r[3]]

            # Re-run: verdict survives, and the Initiatives rows are not duplicated.
            with mock.patch("builtins.print"):
                self.assertEqual(d.cmd_workbook(Namespace(run_dir=str(run))), 0)
            wb = openpyxl.load_workbook(book)
            tab = wb["Dealbreakers"]
            self.assertEqual(tab.cell(row=7, column=12).value, "Agree")
            self.assertEqual(tab.cell(row=7, column=13).value, "Confirmed with the client")
            initiatives_after = [r[3] for r in wb["Initiatives"].iter_rows(min_row=6, values_only=True) if r[3]]
            self.assertEqual(sorted(initiatives_before), sorted(initiatives_after))
            self.assertIn("Answer dealbreaker: Enterprise pricing is hidden behind a sales call.", initiatives_after)
            self.assertEqual(len(initiatives_after), len(set(initiatives_after)))  # no duplicates

    def test_competitor_mode_writes_no_initiatives(self):
        import openpyxl
        if not d.PACK_TEMPLATE.is_file():
            self.skipTest("pack template not built")
        with tempfile.TemporaryDirectory() as tmp:
            run = Path(tmp) / "dealbreakers"
            findings = load_findings()
            findings["mode"] = "competitor"
            shutil.copytree(FIXTURE_RUN / "data", run / "data")
            (run / "data" / "findings.json").write_text(json.dumps(findings), encoding="utf-8")
            with mock.patch("builtins.print"):
                self.assertEqual(d.cmd_build(Namespace(run_dir=str(run))), 0)
                self.assertEqual(d.cmd_workbook(Namespace(run_dir=str(run))), 0)
            book = Path(tmp) / d.WORKBOOK
            wb = openpyxl.load_workbook(book)
            self.assertIn("Dealbreakers - Acme Widgets", wb.sheetnames)
            tasks = [r[3] for r in wb["Initiatives"].iter_rows(min_row=6, values_only=True) if r[3]]
            self.assertEqual(tasks, [])
            names = wb.sheetnames
            self.assertLess(names.index("Checklist"), names.index("Dealbreakers - Acme Widgets"))
            self.assertLess(names.index("Dealbreakers - Acme Widgets"), names.index("Initiatives"))


if __name__ == "__main__":
    unittest.main()
