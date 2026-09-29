"""Offline tests for fanmap.py. No network, no keys.

Run: uv run scripts/test_fanmap.py  (or: python3 -m unittest scripts/test_fanmap.py)
"""

import csv
import json
import sys
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).parent))
import fanmap as f  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"
SAMPLE_HTML = (FIXTURES / "sample-page.html").read_text()
PAGE_URL = "https://example-guides.test/guides/geo/llm-share-buttons/"
ENV = {"GEMINI_API_KEY": "test-fake-key-do-not-use-0000000000"}


def gemini_ok(payload: dict) -> tuple[int, dict]:
    return 200, {"candidates": [{"content": {"parts": [{"text": json.dumps(payload)}]}}]}


VALID_PREDICTION = {
    "primary_entity": "LLM share buttons",
    "prompts": ["What are LLM share buttons?", "How do LLM share buttons work?", "Best LLM share button setup"],
    "fan_outs": [
        {"query": "what do llm share buttons do", "type": "related", "coverage": "yes",
         "evidence": "What the buttons actually do"},
        {"query": "do llm share buttons raise citations", "type": "implicit", "coverage": "yes",
         "evidence": "Do they raise citations?"},
        {"query": "llm share button vs manual copy paste", "type": "comparative", "coverage": "no", "evidence": ""},
        {"query": "how to install llm share buttons", "type": "procedural", "coverage": "partial",
         "evidence": "chatgpt.com/?q="},
    ],
    "gaps": ["no install code sample"],
    "follow_ups": ["does this work on mobile", "is there a wordpress plugin", "what about perplexity"],
}


class FanmapTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        patches = [mock.patch.object(f, "load_env", return_value=dict(ENV)), mock.patch.object(f.time, "sleep")]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)
        self.addCleanup(self.tmp.cleanup)

    # ------------------------------------------------------------- pages --

    def test_chunk_extraction_on_fixture(self):
        with mock.patch.object(f, "fetch_page", return_value=(200, SAMPLE_HTML.encode(), PAGE_URL)):
            rec = f.extract_page(PAGE_URL)
        self.assertEqual(rec["status"], 200)
        self.assertIn("LLM Share Buttons", rec["title"])
        self.assertEqual(rec["h1"], "LLM share buttons for your GEO page")
        levels_texts = [(c["level"], c["text"]) for c in rec["chunks"]]
        self.assertIn(("h2", "What the buttons actually do"), levels_texts)
        self.assertIn(("h3", "What is not true"), levels_texts)
        # content following a heading is captured, capped, and heading text isn't duplicated into it
        h2 = next(c for c in rec["chunks"] if c["text"] == "What the buttons actually do")
        self.assertIn("chatgpt.com/?q=", h2["content"])
        self.assertNotIn("What the buttons actually do", h2["content"])
        self.assertLessEqual(len(h2["content"]), f.CHUNK_CONTENT_LIMIT)
        # nav/header/footer chrome never leaks into a chunk or a list
        self.assertTrue(rec["lists"])
        self.assertNotIn("Copyright", rec["lists"][0])
        self.assertIn("chatgpt.com/?q=", rec["lists"][0])
        # JSON-LD @type is collected even though it lives inside a <script> tag
        self.assertIn("Article", rec["json_ld_types"])
        self.assertIn("Person", rec["json_ld_types"])
        self.assertTrue(rec["content_hash"])

    def test_pages_command_skips_non_200_and_writes_jsonl(self):
        def fake_fetch(url, timeout=20):
            if "broken" in url:
                return 404, b"", url
            return 200, SAMPLE_HTML.encode(), url

        with mock.patch.object(f, "fetch_page", side_effect=fake_fetch):
            urls_file = self.dir / "urls.txt"
            urls_file.write_text(f"{PAGE_URL}\nhttps://example-guides.test/broken/\n")
            args = Namespace(url=None, urls=str(urls_file), sitemap=None, include=None, limit=None,
                             out=str(self.dir / "run"), workers=2, refresh=False)
            with mock.patch("builtins.print"):
                rc = f.cmd_pages(args)
        self.assertEqual(rc, 0)
        lines = (self.dir / "run" / "data" / "pages.jsonl").read_text().splitlines()
        records = [json.loads(l) for l in lines]
        by_status = {r["status"] for r in records}
        self.assertEqual(by_status, {200, 404})
        broken = next(r for r in records if r["status"] == 404)
        self.assertEqual(broken["error"], "HTTP 404")

    # ----------------------------------------------------------- predict --

    def _page_record(self) -> dict:
        with mock.patch.object(f, "fetch_page", return_value=(200, SAMPLE_HTML.encode(), PAGE_URL)):
            return f.extract_page(PAGE_URL)

    def _write_pages_jsonl(self, run_dir: Path, pages: list[dict]) -> None:
        (run_dir / "data").mkdir(parents=True, exist_ok=True)
        with (run_dir / "data" / "pages.jsonl").open("w", encoding="utf-8") as fh:
            for p in pages:
                fh.write(json.dumps(p) + "\n")

    def test_build_predict_prompt_includes_structure(self):
        page = self._page_record()
        prompt = f.build_predict_prompt(page)
        self.assertIn("SECTIONS", prompt)
        self.assertIn("What the buttons actually do", prompt)
        self.assertIn("Article", prompt)

    def test_validate_prediction(self):
        self.assertIsNone(f.validate_prediction(VALID_PREDICTION))
        self.assertIsNotNone(f.validate_prediction({}))
        self.assertIsNotNone(f.validate_prediction(dict(VALID_PREDICTION, fan_outs=[])))
        bad_type = json.loads(json.dumps(VALID_PREDICTION))
        bad_type["fan_outs"][0]["type"] = "not-a-real-type"
        self.assertIsNotNone(f.validate_prediction(bad_type))
        bad_coverage = json.loads(json.dumps(VALID_PREDICTION))
        bad_coverage["fan_outs"][0]["coverage"] = "maybe"
        self.assertIsNotNone(f.validate_prediction(bad_coverage))

    def test_predict_retries_once_on_malformed_json_then_succeeds(self):
        run_dir = self.dir / "run"
        page = self._page_record()
        self._write_pages_jsonl(run_dir, [page])
        calls = {"n": 0}

        def fake_http(method, url, headers=None, body=None, timeout=0):
            calls["n"] += 1
            if calls["n"] == 1:
                return 200, {"candidates": [{"content": {"parts": [{"text": "not json at all"}]}}]}
            return gemini_ok(VALID_PREDICTION)

        with mock.patch.object(f, "http_json", side_effect=fake_http):
            args = Namespace(out=str(run_dir), env_file=None, model="gemini-flash-latest", concurrency=1, rpm=0)
            with mock.patch("builtins.print"):
                rc = f.cmd_predict(args)
        self.assertEqual(rc, 0)
        self.assertEqual(calls["n"], 2)  # one bad response, one retry that validated
        rows = [json.loads(l) for l in (run_dir / "data" / "predictions.jsonl").read_text().splitlines()]
        self.assertEqual(rows[0]["status"], "ok")
        self.assertEqual(rows[0]["primary_entity"], "LLM share buttons")

    def test_predict_gives_up_after_retry_and_does_not_kill_the_run(self):
        run_dir = self.dir / "run"
        page = self._page_record()
        self._write_pages_jsonl(run_dir, [page])

        def fake_http(method, url, headers=None, body=None, timeout=0):
            return 200, {"candidates": [{"content": {"parts": [{"text": "still not json"}]}}]}

        with mock.patch.object(f, "http_json", side_effect=fake_http):
            args = Namespace(out=str(run_dir), env_file=None, model="gemini-flash-latest", concurrency=1, rpm=0)
            with mock.patch("builtins.print"):
                rc = f.cmd_predict(args)
        self.assertEqual(rc, 0)  # a bad page is logged, never fatal
        rows = [json.loads(l) for l in (run_dir / "data" / "predictions.jsonl").read_text().splitlines()]
        self.assertEqual(rows[0]["status"], "error")
        self.assertIn("invalid prediction JSON", rows[0]["error"])

    def test_predict_cache_hit_makes_zero_calls(self):
        run_dir = self.dir / "run"
        page = self._page_record()
        self._write_pages_jsonl(run_dir, [page])

        with mock.patch.object(f, "http_json", side_effect=lambda *a, **k: gemini_ok(VALID_PREDICTION)):
            args = Namespace(out=str(run_dir), env_file=None, model="gemini-flash-latest", concurrency=1, rpm=0)
            with mock.patch("builtins.print"):
                rc = f.cmd_predict(args)
        self.assertEqual(rc, 0)

        with mock.patch.object(f, "http_json", side_effect=AssertionError("must not call the API again")):
            args2 = Namespace(out=str(run_dir), env_file=None, model="gemini-flash-latest", concurrency=1, rpm=0)
            with mock.patch("builtins.print"):
                rc2 = f.cmd_predict(args2)
        self.assertEqual(rc2, 0)
        rows = [json.loads(l) for l in (run_dir / "data" / "predictions.jsonl").read_text().splitlines()]
        self.assertTrue(rows[0]["from_cache"])
        self.assertEqual(rows[0]["status"], "ok")

    def test_429_backoff_then_succeeds(self):
        run_dir = self.dir / "run"
        page = self._page_record()
        self._write_pages_jsonl(run_dir, [page])
        calls = {"n": 0}

        def fake_http(method, url, headers=None, body=None, timeout=0):
            calls["n"] += 1
            if calls["n"] < 3:
                return 429, {"error": "rate limited"}
            return gemini_ok(VALID_PREDICTION)

        with mock.patch.object(f, "http_json", side_effect=fake_http):
            with mock.patch.object(f.time, "sleep") as sleep_mock:
                args = Namespace(out=str(run_dir), env_file=None, model="gemini-flash-latest", concurrency=1, rpm=0)
                with mock.patch("builtins.print"):
                    rc = f.cmd_predict(args)
        self.assertEqual(rc, 0)
        self.assertEqual(calls["n"], 3)
        self.assertEqual(sleep_mock.call_count, 2)  # backed off twice before the 3rd call succeeded
        rows = [json.loads(l) for l in (run_dir / "data" / "predictions.jsonl").read_text().splitlines()]
        self.assertEqual(rows[0]["status"], "ok")

    def test_503_exhausts_retries_and_is_logged_not_fatal(self):
        run_dir = self.dir / "run"
        page = self._page_record()
        self._write_pages_jsonl(run_dir, [page])

        with mock.patch.object(f, "http_json", return_value=(503, {"error": "unavailable"})):
            args = Namespace(out=str(run_dir), env_file=None, model="gemini-flash-latest", concurrency=1, rpm=0)
            with mock.patch("builtins.print"):
                rc = f.cmd_predict(args)
        self.assertEqual(rc, 0)
        rows = [json.loads(l) for l in (run_dir / "data" / "predictions.jsonl").read_text().splitlines()]
        self.assertEqual(rows[0]["status"], "error")
        self.assertIn("HTTP 503", rows[0]["error"])

    # ------------------------------------------------------------ report --

    def test_cluster_fanouts_groups_near_duplicates(self):
        items = [
            ("https://a.test/x", "best AI SEO tools 2025", "yes"),
            ("https://b.test/y", "best AI SEO tools 2026", "no"),
            ("https://a.test/x", "site:vendor.example pricing", "no"),
        ]
        clusters = f.cluster_fanouts(items)
        self.assertEqual(len(clusters), 2)
        sizes = sorted(len(c["members"]) for c in clusters)
        self.assertEqual(sizes, [1, 2])

    def test_report_writes_site_wide_gaps_and_overlap(self):
        run_dir = self.dir / "run"
        page_a = dict(self._page_record(), url="https://example-guides.test/a/", slug="a")
        page_b = dict(self._page_record(), url="https://example-guides.test/b/", slug="b")
        self._write_pages_jsonl(run_dir, [page_a, page_b])

        predictions = [
            {"url": page_a["url"], "slug": "a", "status": "ok", "primary_entity": "Page A",
             "prompts": ["p1"], "gaps": ["missing install steps"], "follow_ups": ["f1"],
             "fan_outs": [
                 {"query": "shared unique topic only a covers", "type": "open", "coverage": "yes", "evidence": "e"},
                 {"query": "install steps for the widget", "type": "procedural", "coverage": "no", "evidence": ""},
             ]},
            {"url": page_b["url"], "slug": "b", "status": "ok", "primary_entity": "Page B",
             "prompts": ["p2"], "gaps": [], "follow_ups": [],
             "fan_outs": [
                 {"query": "shared unique topic only a covers", "type": "open", "coverage": "yes", "evidence": "e"},
                 {"query": "widget refund policy question", "type": "implicit", "coverage": "no", "evidence": ""},
             ]},
        ]
        with (run_dir / "data" / "predictions.jsonl").open("w", encoding="utf-8") as fh:
            for row in predictions:
                fh.write(json.dumps(row) + "\n")

        with mock.patch("builtins.print"):
            rc = f.cmd_report(Namespace(out=str(run_dir), top=10))
        self.assertEqual(rc, 0)

        report = (run_dir / "fan-out-map.md").read_text()
        self.assertIn("# Fan-out map", report)
        self.assertIn("PREDICTED by a model from the page's own content, not observed", report)
        # both pages claim "yes" on the near-duplicate topic -> overlap
        self.assertIn("## Possible overlap", report)
        self.assertIn("shared unique topic", report.split("## Possible overlap")[1].split("## Next step")[0])
        # install steps / refund policy are covered by nobody -> content gap candidates
        gaps_section = report.split("## Content gaps across the inventory")[1].split("## Possible overlap")[0]
        self.assertTrue("install steps" in gaps_section or "refund policy" in gaps_section)

        with (run_dir / "data" / "fan-out-map.csv").open(encoding="utf-8") as fh:
            csv_rows = list(csv.DictReader(fh))
        self.assertEqual(len(csv_rows), 4)  # 2 pages x 2 fan-outs each

        page_a_md = (run_dir / "pages" / "a.md").read_text()
        self.assertIn("Page A", page_a_md)
        self.assertIn("missing install steps", page_a_md)

    def test_top_optimise_ranks_lowest_coverage_first(self):
        run_dir = self.dir / "run"
        page_a = dict(self._page_record(), url="https://example-guides.test/a/", slug="a")
        page_b = dict(self._page_record(), url="https://example-guides.test/b/", slug="b")
        self._write_pages_jsonl(run_dir, [page_a, page_b])
        predictions = [
            {"url": page_a["url"], "slug": "a", "status": "ok", "primary_entity": "A", "prompts": [], "gaps": [],
             "follow_ups": [], "fan_outs": [{"query": "q1", "type": "open", "coverage": "no", "evidence": ""}]},
            {"url": page_b["url"], "slug": "b", "status": "ok", "primary_entity": "B", "prompts": [], "gaps": [],
             "follow_ups": [], "fan_outs": [{"query": "q2", "type": "open", "coverage": "yes", "evidence": ""}]},
        ]
        with (run_dir / "data" / "predictions.jsonl").open("w", encoding="utf-8") as fh:
            for row in predictions:
                fh.write(json.dumps(row) + "\n")
        with mock.patch("builtins.print"):
            f.cmd_report(Namespace(out=str(run_dir), top=10))
        report = (run_dir / "fan-out-map.md").read_text()
        optimise_section = report.split("## Top pages to optimise")[1].split("## Content gaps")[0]
        self.assertLess(optimise_section.index(page_a["url"]), optimise_section.index(page_b["url"]))

    # -------------------------------------------------------------- path --

    def test_run_dir_for_inside_and_outside_a_brain(self):
        brain = self.dir / "brain"
        (brain / ".mos").mkdir(parents=True)
        (brain / ".mos" / "config.yaml").write_text('{"mode": "in-house"}')
        (brain / "content").mkdir()
        inside = f.run_dir_for("example-guides.test", "2026-09-25", brain / "content")
        self.assertEqual(inside, (brain / "campaigns/geo/2026-09/mos-geo-fan-out-map").resolve())
        inside.mkdir(parents=True)
        (inside / "fan-out-map.md").write_text("x")
        again = f.run_dir_for("example-guides.test", "2026-09-30", brain)
        self.assertEqual(again.name, "mos-geo-fan-out-map-2")
        outside = f.run_dir_for("Acme & Co.", "2026-01-05", self.dir)
        self.assertEqual(outside, (self.dir / "outputs/geo/2026-01/acme-co/mos-geo-fan-out-map").resolve())
        with self.assertRaises(SystemExit):
            f.run_dir_for("Acme", "25-09-2026", self.dir)

    def test_cmd_path_derives_slug_from_url(self):
        with mock.patch("builtins.print") as p:
            f.cmd_path(Namespace(site=None, url=PAGE_URL, date="2026-09-25", start=str(self.dir)))
        printed = str(p.call_args[0][0]).replace("\\", "/")
        self.assertIn("outputs/geo/2026-09/example-guides-test/mos-geo-fan-out-map", printed)

    # ---------------------------------------------------------- preflight --

    def test_preflight_refuses_without_any_page_count(self):
        args = Namespace(env_file=None, model="gemini-flash-latest", pages=0, concurrency=4, rpm=0,
                         sitemap=None, include=None, limit=None, urls=None, url=None)
        with mock.patch.object(f, "http_json", side_effect=AssertionError("must not call the API")):
            with mock.patch("builtins.print"):
                rc = f.cmd_preflight(args)
        self.assertEqual(rc, 2)

    def test_preflight_ok_with_manual_page_count(self):
        def fake_http(method, url, headers=None, body=None, timeout=0):
            return 200, {"candidates": [{"content": {"parts": [{"text": "ok"}]}}]}

        args = Namespace(env_file=None, model="gemini-flash-latest", pages=3, concurrency=4, rpm=0,
                         sitemap=None, include=None, limit=None, urls=None, url=None)
        with mock.patch.object(f, "http_json", side_effect=fake_http):
            with mock.patch("builtins.print"):
                rc = f.cmd_preflight(args)
        self.assertEqual(rc, 0)

    # ---------------------------------------------------------- workbook --

    def _workbook_predictions(self, page_a_url: str, page_b_url: str) -> list[dict]:
        return [
            {"url": page_a_url, "slug": "a", "status": "ok", "primary_entity": "Page A",
             "prompts": ["p1"], "gaps": [], "follow_ups": [],
             "fan_outs": [
                 {"query": "widget install steps", "type": "procedural", "coverage": "no", "evidence": ""},
                 {"query": "widget refund policy", "type": "implicit", "coverage": "no", "evidence": ""},
             ]},
            {"url": page_b_url, "slug": "b", "status": "ok", "primary_entity": "Page B",
             "prompts": ["p2"], "gaps": [], "follow_ups": [],
             "fan_outs": [
                 {"query": "widget pricing plans", "type": "open", "coverage": "yes", "evidence": "e"},
             ]},
        ]

    def _write_predictions_jsonl(self, run_dir: Path, predictions: list[dict]) -> None:
        (run_dir / "data").mkdir(parents=True, exist_ok=True)
        with (run_dir / "data" / "predictions.jsonl").open("w", encoding="utf-8") as fh:
            for row in predictions:
                fh.write(json.dumps(row) + "\n")

    def _fanout_map_row_count(self, wb) -> int:
        ws = wb[f.FANOUT_MAP_TAB]
        return sum(1 for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=2).value)

    def _initiatives_row_count(self, wb) -> int:
        ws = wb["Initiatives"]
        return sum(1 for r in range(6, ws.max_row + 1) if ws.cell(row=r, column=4).value)

    @unittest.skipUnless(__import__("importlib").util.find_spec("openpyxl"), "openpyxl not installed")
    def test_workbook_creates_fanout_map_tab_and_initiatives(self):
        import openpyxl
        run_dir = self.dir / "run"
        self._write_predictions_jsonl(run_dir, self._workbook_predictions(
            "https://example-guides.test/a/", "https://example-guides.test/b/"))
        book_path = self.dir / "brand-audit-master.xlsx"

        # explicit --workbook: tick_checklist.py's own path math can't see this location, so
        # cmd_workbook must tick the Checklist row in-process instead of shelling out to it.
        with mock.patch.object(f.subprocess, "run") as run_mock:
            args = Namespace(out=str(run_dir), workbook=str(book_path), top=5)
            with mock.patch("builtins.print"):
                rc = f.cmd_workbook(args)
        self.assertEqual(rc, 0)
        run_mock.assert_not_called()

        self.assertTrue(book_path.is_file())
        wb = openpyxl.load_workbook(book_path)
        self.assertIn(f.FANOUT_MAP_TAB, wb.sheetnames)
        ws = wb[f.FANOUT_MAP_TAB]
        self.assertIn("PREDICTED", ws.cell(row=2, column=2).value)
        self.assertEqual(ws.cell(row=4, column=2).value, "Page URL")
        self.assertEqual(self._fanout_map_row_count(wb), 3)  # 2 + 1 fan-outs across 2 pages

        by_query = {ws.cell(row=r, column=5).value: ws.cell(row=r, column=10).value
                    for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=2).value}
        self.assertEqual(by_query["widget install steps"], "Y")            # coverage=no, nobody covers it -> gap
        self.assertIn(by_query["widget pricing plans"], (None, ""))        # coverage=yes -> not a gap

        self.assertIn("Initiatives", wb.sheetnames)
        init_ws = wb["Initiatives"]
        tasks = [init_ws.cell(row=r, column=4).value for r in range(6, init_ws.max_row + 1)
                if init_ws.cell(row=r, column=4).value]
        self.assertTrue(any(t.startswith("Optimise ") for t in tasks))
        self.assertTrue(any(t.startswith("New content candidate:") for t in tasks))

        # the real pack template registers this skill in Checklist - confirm it got ticked
        # in the SAME file cmd_workbook just wrote to, not some other default location.
        checklist_ws = wb["Checklist"]
        row = next(r for r in range(1, checklist_ws.max_row + 1)
                  if checklist_ws.cell(row=r, column=3).value == "mos-geo-fan-out-map")
        self.assertEqual(checklist_ws.cell(row=row, column=8).value, "Client review")
        self.assertEqual(checklist_ws.cell(row=row, column=9).value, "☑")

    @unittest.skipUnless(__import__("importlib").util.find_spec("openpyxl"), "openpyxl not installed")
    def test_workbook_default_location_shells_out_to_tick_checklist(self):
        run_dir = self.dir / "brain" / "campaigns" / "geo" / "2026-09" / "mos-geo-fan-out-map"
        self._write_predictions_jsonl(run_dir, self._workbook_predictions(
            "https://example-guides.test/a/", "https://example-guides.test/b/"))

        with mock.patch.object(f.subprocess, "run") as run_mock:
            args = Namespace(out=str(run_dir), workbook=None, top=5)
            with mock.patch("builtins.print"):
                rc = f.cmd_workbook(args)
        self.assertEqual(rc, 0)
        run_mock.assert_called_once()
        called_cmd = run_mock.call_args[0][0]
        self.assertIn("mos-geo-fan-out-map", called_cmd)
        self.assertIn(str(run_dir.resolve()), called_cmd)
        # the workbook itself still lands one level above the run folder, same as the sibling skill
        self.assertTrue((run_dir.parent / f.WORKBOOK).is_file())

    @unittest.skipUnless(__import__("importlib").util.find_spec("openpyxl"), "openpyxl not installed")
    def test_workbook_rerun_does_not_duplicate_rows_or_initiatives(self):
        import openpyxl
        run_dir = self.dir / "run"
        self._write_predictions_jsonl(run_dir, self._workbook_predictions(
            "https://example-guides.test/a/", "https://example-guides.test/b/"))
        book_path = self.dir / "brand-audit-master.xlsx"
        args = Namespace(out=str(run_dir), workbook=str(book_path), top=5)

        with mock.patch.object(f.subprocess, "run"):
            with mock.patch("builtins.print"):
                f.cmd_workbook(args)
        wb1 = openpyxl.load_workbook(book_path)
        fanout_n1, init_n1 = self._fanout_map_row_count(wb1), self._initiatives_row_count(wb1)

        with mock.patch.object(f.subprocess, "run"):
            with mock.patch("builtins.print"):
                rc2 = f.cmd_workbook(args)
        self.assertEqual(rc2, 0)
        wb2 = openpyxl.load_workbook(book_path)
        self.assertEqual(self._fanout_map_row_count(wb2), fanout_n1)
        self.assertEqual(self._initiatives_row_count(wb2), init_n1)


if __name__ == "__main__":
    unittest.main()
