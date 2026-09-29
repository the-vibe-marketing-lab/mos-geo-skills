#!/usr/bin/env python3
"""fanmap.py - the data layer for mos-geo-fan-out-map.

For a site's whole content inventory (a sitemap or a URL list), predicts each page's
query fan-out with Gemini Flash (free tier, no search grounding), scores each page's
own coverage of those fan-outs, and writes a site-wide map: which pages to optimise
and which topics across the inventory have no page at all. Free and fast, so it runs
BEFORE spending any time on real, observed verification (a live ChatGPT / Google AI
Mode check) of the pages it prioritises. Standard library only.

The prediction method (per-page prompt asking a model for buyer prompts + fan-out
queries + self-coverage) is adapted from Metehan Yesilyurt's Screaming Frog query
fan-out script: https://github.com/metehan777/screaming-frog-query-fan-out

Subcommands
  path        Print where this run's folder belongs (MarketingOS-aware).
  preflight   Check the Gemini key, that the model is reachable, the planned page
              count, and a rate-limit/time estimate.
  pages       Fetch and parse the page(s): title, H1, layout-aware section chunks
              (H2/H3 + up to 500 chars of following content), up to 5 lists, and
              every JSON-LD @type -> data/pages.jsonl.
  predict     One Gemini call per page: primary entity, 3 buyer prompts, 8-10
              fan-outs each typed and self-scored for coverage, gaps, follow-ups.
              Resumable cache keyed by (url, content hash, model) - a re-run with
              unchanged pages makes zero calls.
  report      Write fan-out-map.md (site summary, top pages to optimise, site-wide
              content gaps and overlaps, and a next-step note on confirming the
              pages that matter with an OBSERVED check), one pages/<slug>.md per
              page, and data/fan-out-map.csv (one row per page x fan-out).

Run folder layout
  data/pages.jsonl          one JSON object per fetched page (see `pages`)
  data/predictions.jsonl    one JSON object per page's Gemini prediction
  data/raw/<sha1>.json      the cached, validated Gemini response for one
                             (url, content hash, model)
  fan-out-map.md            the site-wide report
  pages/<slug>.md           per-page entity, prompts, fan-out table, gaps
  data/fan-out-map.csv      one row per page x fan-out, for spreadsheets

Credentials come from the environment, or from --env-file:
  GEMINI_API_KEY   https://aistudio.google.com/apikey (Gemini Flash free tier)

HONESTY NOTE (read this before showing a report to anyone): every fan-out and
coverage verdict in this pack is PREDICTED by a model from the page's own content -
it is not an observed search, and it is biased toward topics the page already
covers (a model reading a page tends to imagine searches that page would answer).
Use it to prioritise which pages are worth real, observed verification (a live
ChatGPT / Google AI Mode check of the page's buyer prompts) - never as the final
word on what an AI engine actually searches for or cites. See references/evidence.md.
"""

from __future__ import annotations

import argparse
import concurrent.futures as cf
import csv
import hashlib
import json
import os
import re
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from argparse import Namespace
from collections import deque
from html.parser import HTMLParser
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
API_UA = "mos-geo-fan-out-map/1.0"
# Some hosts 403 anything that doesn't look like a browser.
PAGE_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
          "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36")
ENV_KEYS = ("GEMINI_API_KEY",)
DATA_DIR = "data"
RAW_DIR = "raw"
SKILL_FOLDER = "mos-geo-fan-out-map"

GEMINI_BASE = "https://generativelanguage.googleapis.com/v1beta"
DEFAULT_MODEL = "gemini-flash-latest"
FALLBACK_MODEL = "gemini-flash-lite-latest"  # faster/cheaper, pass --model to use it

CHUNK_CONTENT_LIMIT = 500  # chars of body text kept per H2/H3 - see docstring
MAX_LISTS = 5
EST_SECONDS_PER_CALL = 4.0   # observed live, see references/evidence.md
EST_TOKENS_PER_CALL = 2500

COVERAGE_SCORE = {"yes": 1.0, "partial": 0.5, "no": 0.0}
VALID_TYPES = {"related", "implicit", "comparative", "procedural", "refinement"}
VALID_COVERAGE = {"yes", "partial", "no"}
REQUIRED_FANOUT_KEYS = {"query", "type", "coverage"}
CLUSTER_THRESHOLD = 0.6  # token Jaccard - same heuristic as the sibling skill

WORD_RE = re.compile(r"[a-z0-9][a-z0-9'\-]*")
STOPWORDS = {
    "the", "a", "an", "of", "for", "to", "in", "on", "and", "or", "is", "are", "what",
    "how", "do", "does", "you", "your", "i", "my", "we", "our", "it", "its", "with",
    "this", "that", "be", "can", "vs", "near", "me", "about", "at", "by", "from", "as",
}

MAX_JSON_RETRIES = 1     # extra attempts after the first, on invalid/malformed JSON
MAX_HTTP_RETRIES = 5     # attempts on 429/503 before giving up on a page
BACKOFF_BASE = 2.0       # seconds, doubles each retry

# the shared GEO brand-audit workbook - same conventions as the paid sibling's `workbook` step
WORKBOOK = "brand-audit-master.xlsx"
PACK_TEMPLATE = SKILL_DIR.parent / "_shared" / "brand-audit" / "brand-audit-template.xlsx"
TICK_SCRIPT = SKILL_DIR.parent / "_shared" / "brand-audit" / "tick_checklist.py"
FANOUT_MAP_TAB = "Fan-Out Map"
COVERAGE_FILL = {"yes": "DCFCE7", "partial": "FEF3C7", "no": "FEE2E2"}  # green / amber / red


# --------------------------------------------------------------- helpers --

def load_env(env_file: str | None) -> dict:
    env = {}
    if env_file:
        path = Path(env_file).expanduser()
        if not path.is_file():
            sys.exit(f"--env-file not found: {path}")
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, val = line.partition("=")
            env[key.strip().removeprefix("export ").strip()] = val.strip().strip("'\"")
    for key in ENV_KEYS:
        if os.environ.get(key):
            env[key] = os.environ[key]
    return env


def mask(value: str | None) -> str:
    return f"set (…{value[-4:]})" if value else "MISSING"


def http_json(method, url, headers=None, body=None, timeout=180):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    req.add_header("User-Agent", API_UA)
    if data is not None:
        req.add_header("Content-Type", "application/json")
    for k, v in (headers or {}).items():
        req.add_header(k, v)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            raw = resp.read().decode("utf-8", "replace")
            return resp.status, (json.loads(raw) if raw.strip() else None)
    except urllib.error.HTTPError as e:
        raw = e.read().decode("utf-8", "replace")
        try:
            return e.code, json.loads(raw)
        except ValueError:
            return e.code, {"error": raw[:500]}
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        return 0, {"error": f"{type(e).__name__}: {e}"}


def domain_of(url: str) -> str:
    host = urllib.parse.urlparse(url if "//" in url else f"//{url}").hostname or ""
    host = host.lower()
    return host[4:] if host.startswith("www.") else host


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")


def slug_of(url: str) -> str:
    p = urllib.parse.urlsplit(url)
    path = p.path.strip("/")
    raw = f"{domain_of(url)}-{path}" if path else domain_of(url) or "page"
    return re.sub(r"[^a-z0-9]+", "-", raw.lower()).strip("-") or "page"


# ------------------------------------------------------------------ path --

def _main_checkout(git_file: Path) -> Path | None:
    """The main checkout behind a linked worktree's `.git` file
    ("gitdir: <main>/.git/worktrees/<name>"), including a Windows path read from WSL."""
    m = re.match(r"gitdir:\s*(.+)", git_file.read_text(encoding="utf-8").strip())
    if not m:
        return None
    raw = m.group(1).strip()
    drive = re.match(r"^([A-Za-z]):[\\/](.*)$", raw)
    if drive and os.name != "nt":
        raw = f"/mnt/{drive.group(1).lower()}/{drive.group(2)}"
    gitdir = Path(raw.replace("\\", "/"))
    if not gitdir.is_absolute():
        gitdir = (git_file.parent / gitdir).resolve()
    for d in [gitdir, *gitdir.parents]:
        if d.name == ".git":
            return d.parent
    return None


def brain_config(root: Path) -> Path | None:
    own = root / ".mos" / "config.yaml"
    if own.is_file():
        return own
    git = root / ".git"
    if git.is_file():  # linked worktree: .mos/ is gitignored, so borrow the main checkout's
        main = _main_checkout(git)
        if main and (main / ".mos" / "config.yaml").is_file():
            return main / ".mos" / "config.yaml"
    return None


def find_brain(start: Path) -> Path | None:
    for d in [start, *start.parents]:
        if brain_config(d):
            return d
        if (d / ".git").exists():
            return None
    return None


def brain_mode(brain: Path) -> str:
    config = brain_config(brain)
    if not config:
        return "in-house"
    m = re.search(r'["\']?mode["\']?\s*:\s*["\']?([a-z-]+)', config.read_text(encoding="utf-8"))
    return m.group(1) if m else "in-house"


def run_dir_for(site: str, day: str, start: Path) -> Path:
    """campaigns/geo/YYYY-MM/mos-geo-fan-out-map/ in a one-brand brain, with a brand
    folder in an agency brain, outputs/geo/YYYY-MM/<site>/mos-geo-fan-out-map/ outside
    a brain. A repeat run in the same month gets -2, -3 ... so nothing is overwritten."""
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", day):
        sys.exit(f"--date must be YYYY-MM-DD, got {day}")
    slug = slugify(site) or "site"
    brain = find_brain(start.resolve())
    if brain and brain_mode(brain) != "agency":
        month = brain / "campaigns" / "geo" / day[:7]
    elif brain:
        month = brain / "campaigns" / "geo" / day[:7] / slug
    else:
        month = start.resolve() / "outputs" / "geo" / day[:7] / slug
    target, n = month / SKILL_FOLDER, 2
    while target.exists() and any(target.iterdir()):
        target = month / f"{SKILL_FOLDER}-{n}"
        n += 1
    return target


def cmd_path(args) -> int:
    day = args.date or time.strftime("%Y-%m-%d")
    site = args.site or (domain_of(args.url) if args.url else "site")
    print(run_dir_for(site, day, Path(args.start or ".")))
    return 0


# ----------------------------------------------------------------- pages --

# Tags whose descendant text never belongs to the article: chrome, not content.
# `script` is handled separately below so JSON-LD survives while other scripts are skipped.
BODY_SKIP = {"style", "noscript", "template", "svg", "iframe", "nav", "header", "footer", "aside"}
BLOCKY = {"p", "div", "li", "tr", "table", "ul", "ol", "br", "section", "article", "blockquote"}
LIST_TAGS = {"ul", "ol"}


class ChunkExtractor(HTMLParser):
    """Pulls layout-aware chunks out of one HTML page for the fan-out predictor:
    title, H1, each H2/H3 with up to CHUNK_CONTENT_LIMIT chars of the content that
    follows it (until the next heading), up to MAX_LISTS lists (their item text
    joined), and every JSON-LD @type declared on the page. Stdlib only - same
    skip-chrome approach as the sibling skill's PageExtractor, extended for chunks."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title = ""
        self.h1 = ""
        self.chunks: list[dict] = []       # [{"level": "h2", "text": ..., "content": [...]}]
        self.lists: list[str] = []
        self.json_ld_types: list[str] = []
        self._body: list[str] = []
        self._skip: list[str] = []
        self._in_title = False
        self._heading_tag: str | None = None
        self._heading_buf: list[str] = []
        self._current_chunk: dict | None = None
        self._in_ldjson = False
        self._ldjson_buf: list[str] = []
        self._list_depth = 0
        self._list_buf: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "script":
            if (dict(attrs).get("type") or "").strip().lower() == "application/ld+json":
                self._in_ldjson, self._ldjson_buf = True, []
            else:
                self._skip.append(tag)
            return
        if tag in BODY_SKIP:
            self._skip.append(tag)
            return
        if self._skip:
            return
        if tag == "title":
            self._in_title = True
        elif tag in ("h1", "h2", "h3"):
            self._close_heading()
            self._heading_tag, self._heading_buf = tag, []
        elif tag in LIST_TAGS:
            self._list_depth += 1
            if self._list_depth == 1 and len(self.lists) < MAX_LISTS:
                self._list_buf = []
        elif tag == "li" and self._list_depth >= 1 and self._list_buf is not None:
            self._list_buf.append("- ")
        elif tag in BLOCKY:
            self._body.append(" ")
            if self._current_chunk is not None:
                self._current_chunk["content"].append(" ")

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in BODY_SKIP and tag != "script":
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if tag == "script" and self._in_ldjson:
            self._in_ldjson = False
            self._parse_ldjson("".join(self._ldjson_buf))
            return
        if self._skip:
            if self._skip[-1] == tag:
                self._skip.pop()
            return
        if tag == "title":
            self._in_title = False
        elif self._heading_tag == tag:
            self._close_heading()
        elif tag in LIST_TAGS:
            if self._list_depth == 1 and self._list_buf is not None and len(self.lists) < MAX_LISTS:
                text = re.sub(r"\s+", " ", "".join(self._list_buf)).strip()
                if text:
                    self.lists.append(text[:CHUNK_CONTENT_LIMIT])
                self._list_buf = None
            self._list_depth = max(0, self._list_depth - 1)
        elif tag in BLOCKY:
            self._body.append(" ")
            if self._current_chunk is not None:
                self._current_chunk["content"].append(" ")

    def handle_data(self, data):
        if self._in_ldjson:
            self._ldjson_buf.append(data)
            return
        if self._skip or not data:
            return
        if self._in_title:
            self.title += data
        elif self._heading_tag:
            self._heading_buf.append(data)
        else:
            self._body.append(data)
            if self._list_depth >= 1 and self._list_buf is not None:
                self._list_buf.append(data)
            if self._current_chunk is not None:
                self._current_chunk["content"].append(data)

    def _close_heading(self):
        if self._heading_tag:
            text = " ".join("".join(self._heading_buf).split())
            if text:
                if self._heading_tag == "h1" and not self.h1:
                    self.h1 = text
                    self._current_chunk = None
                elif self._heading_tag in ("h2", "h3"):
                    self._current_chunk = {"level": self._heading_tag, "text": text, "content": []}
                    self.chunks.append(self._current_chunk)
            self._heading_tag, self._heading_buf = None, []

    def _parse_ldjson(self, raw: str) -> None:
        try:
            data = json.loads(raw)
        except (json.JSONDecodeError, TypeError):
            return

        def collect(obj):
            if isinstance(obj, dict):
                t = obj.get("@type")
                if isinstance(t, str):
                    self.json_ld_types.append(t)
                elif isinstance(t, list):
                    self.json_ld_types.extend(x for x in t if isinstance(x, str))
                for v in obj.values():
                    collect(v)
            elif isinstance(obj, list):
                for item in obj:
                    collect(item)

        collect(data)

    def finalize(self) -> "ChunkExtractor":
        for c in self.chunks:
            c["content"] = re.sub(r"\s+", " ", "".join(c["content"])).strip()[:CHUNK_CONTENT_LIMIT]
        seen, deduped = set(), []
        for t in self.json_ld_types:
            if t not in seen:
                seen.add(t)
                deduped.append(t)
        self.json_ld_types = deduped
        return self

    def main_text(self) -> str:
        return re.sub(r"\s+", " ", "".join(self._body)).strip()


def fetch_page(url: str, timeout: int = 20) -> tuple[int, bytes, str]:
    req = urllib.request.Request(url, headers={"User-Agent": PAGE_UA,
                                                "Accept": "text/html,application/xhtml+xml;q=0.9,*/*;q=0.8"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, r.read(), r.geturl()
    except urllib.error.HTTPError as e:
        return e.code, b"", url
    except Exception as e:  # noqa: BLE001 - network errors become a logged reason
        return 0, str(e).encode(), url


def fetch_sitemap_urls(url: str, include: str | None, limit: int | None, _seen=None) -> list[str]:
    seen = _seen if _seen is not None else set()
    if url in seen:
        return []
    seen.add(url)
    status, body, _ = fetch_page(url)
    if status != 200:
        sys.exit(f"could not fetch sitemap {url}: HTTP {status}")
    text = body.decode("utf-8", "replace")
    locs = re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", text)
    urls = []
    for loc in locs:
        if "sitemap" in loc.lower() and loc.lower().endswith((".xml", ".xml.gz")):
            urls += fetch_sitemap_urls(loc, None, None, seen)
        else:
            urls.append(loc)
    if include:
        pat = re.compile(include)
        urls = [u for u in urls if pat.search(u)]
    if limit:
        urls = urls[:limit]
    return urls


def page_content_hash(page: dict) -> str:
    """Fingerprints the extracted content this skill actually sends to Gemini, so a
    re-run only re-predicts pages whose relevant content changed."""
    basis = {"title": page.get("title", ""), "h1": page.get("h1", ""),
             "chunks": page.get("chunks", []), "lists": page.get("lists", []),
             "json_ld_types": page.get("json_ld_types", [])}
    return hashlib.sha1(json.dumps(basis, sort_keys=True).encode()).hexdigest()[:16]


def extract_page(url: str) -> dict:
    slug = slug_of(url)
    status, body, final_url = fetch_page(url)
    rec = {"url": url, "final_url": final_url, "status": status, "slug": slug}
    if status != 200:
        rec["error"] = f"HTTP {status}" if status else (body.decode("utf-8", "replace") or "fetch failed")
        rec.update(title="", h1="", chunks=[], lists=[], json_ld_types=[], word_count=0)
        return rec
    ext = ChunkExtractor()
    try:
        ext.feed(body.decode("utf-8", "replace"))
    except Exception as e:  # noqa: BLE001 - a malformed page still reports what it got
        rec["error"] = f"parse warning: {e}"
    ext.finalize()
    word_count = len(ext.main_text().split())
    rec.update(title=ext.title.strip(), h1=ext.h1, chunks=ext.chunks, lists=ext.lists,
               json_ld_types=ext.json_ld_types, word_count=word_count)
    rec["content_hash"] = page_content_hash(rec)
    return rec


def cmd_pages(args) -> int:
    if args.url:
        urls = [args.url]
    elif args.urls:
        urls = [l.strip() for l in Path(args.urls).read_text(encoding="utf-8").splitlines()
                if l.strip() and not l.strip().startswith("#")]
        if args.limit:
            urls = urls[: args.limit]
    elif args.sitemap:
        urls = fetch_sitemap_urls(args.sitemap, args.include, args.limit)
    else:
        sys.exit("give --url, --urls or --sitemap")
    if not urls:
        sys.exit("no URLs to fetch")

    out = Path(args.out)
    cache_dir = out / DATA_DIR / "pages"
    cache_dir.mkdir(parents=True, exist_ok=True)

    def work(url: str) -> dict:
        cache = cache_dir / f"{slug_of(url)}.json"
        if cache.is_file() and not args.refresh:
            return json.loads(cache.read_text(encoding="utf-8"))
        rec = extract_page(url)
        cache.write_text(json.dumps(rec, indent=2), encoding="utf-8")
        return rec

    with cf.ThreadPoolExecutor(max_workers=args.workers) as pool:
        records = list(pool.map(work, urls))

    data_dir = out / DATA_DIR
    data_dir.mkdir(parents=True, exist_ok=True)
    with (data_dir / "pages.jsonl").open("w", encoding="utf-8") as fh:
        for rec in records:
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")

    ok = [r for r in records if r.get("status") == 200]
    print(f"Fetched {len(records)} page(s): {len(ok)} ok, {len(records) - len(ok)} skipped.")
    for r in records:
        if r.get("status") != 200:
            print(f"  SKIP {r['url']}: {r.get('error')}")
    return 0 if ok else 2


def load_pages_jsonl(out: Path) -> list[dict]:
    path = out / DATA_DIR / "pages.jsonl"
    if not path.is_file():
        sys.exit(f"missing {path} - run `pages` first")
    return [json.loads(l) for l in path.read_text(encoding="utf-8").splitlines() if l.strip()]


# ------------------------------------------------------------- preflight --

def cmd_preflight(args) -> int:
    env = load_env(args.env_file)
    print("mos-geo-fan-out-map preflight")
    print("==============================")
    for key in ENV_KEYS:
        print(f"{key:<16} {mask(env.get(key))}")

    n_pages = args.pages
    if args.sitemap or args.urls or args.url:
        try:
            if args.url:
                urls = [args.url]
            elif args.urls:
                urls = [l.strip() for l in Path(args.urls).read_text(encoding="utf-8").splitlines()
                        if l.strip() and not l.strip().startswith("#")]
                if args.limit:
                    urls = urls[: args.limit]
            else:
                urls = fetch_sitemap_urls(args.sitemap, args.include, args.limit)
            n_pages = len(urls)
            print(f"\nDiscovered {n_pages} URL(s) from the given source.")
        except SystemExit as e:
            print(f"\nCould not discover URLs: {e}")
            return 2

    if not n_pages:
        print("\nNo page count known - pass --pages N, or --sitemap/--urls/--url to discover it.")
        return 2

    est_tokens = n_pages * EST_TOKENS_PER_CALL
    est_seconds = (n_pages / max(args.concurrency, 1)) * EST_SECONDS_PER_CALL
    if args.rpm:
        est_seconds = max(est_seconds, (n_pages / args.rpm) * 60.0)
    print(f"\nPlanned run: {n_pages} page(s), model {args.model!r}, concurrency {args.concurrency}"
          + (f", capped at {args.rpm} req/min" if args.rpm else "") + ".")
    print(f"Estimated: ~{est_tokens:,} tokens total, ~{est_seconds:.0f}s wall time "
          f"(Gemini Flash free tier - see references/evidence.md for the observed per-call cost).")

    if not env.get("GEMINI_API_KEY"):
        print("\nGEMINI_API_KEY missing.")
        return 2

    status, resp = http_json(
        "POST", f"{GEMINI_BASE}/models/{args.model}:generateContent",
        {"x-goog-api-key": env["GEMINI_API_KEY"]},
        {"contents": [{"parts": [{"text": "Reply with the single word: ok"}]}],
         "generationConfig": {"temperature": 0, "maxOutputTokens": 10}},
        timeout=30,
    )
    if status != 200:
        print(f"\n[FAIL] model {args.model!r} not reachable: HTTP {status} {resp}")
        return 2
    print(f"\n[ok] model {args.model!r} reachable.")
    print("\nReady.")
    return 0


# --------------------------------------------------------------- predict --

CREDIT_LINE = ("Method adapted from Metehan Yesilyurt's Screaming Frog query fan-out script: "
               "https://github.com/metehan777/screaming-frog-query-fan-out")

PROMPT_TEMPLATE = """You predict Google AI Mode / ChatGPT query fan-out for a web page, and score
how well the page's OWN content already covers each predicted fan-out.

URL: {url}
TITLE: {title}
H1: {h1}
SCHEMA TYPES: {schema_types}
SECTIONS (heading, then up to {content_limit} chars of the content that follows it):
{sections}
LISTS (up to {max_lists}, item text joined):
{lists}

Return JSON only, no markdown fences, matching exactly this shape:
{{"primary_entity": "the main entity/topic this page is about",
 "prompts": ["3 realistic buyer prompts this page should win"],
 "fan_outs": [{{"query": "a predicted fan-out search", "type": "related|implicit|comparative|procedural|refinement", "coverage": "yes|partial|no - does the page's OWN content above already answer this", "evidence": "a heading or short quote from the page proving coverage, or empty string if coverage is no"}}],
 "gaps": ["short strings describing what content is missing from this page"],
 "follow_ups": ["3 natural follow-up questions a reader would ask next"]}}
Return 8 to 10 fan_outs, a mix of every type. Base coverage and evidence ONLY on the page
content given above - never assume something is covered that is not shown here."""


def build_predict_prompt(page: dict) -> str:
    sections = "\n".join(f"- {c['level'].upper()} {c['text']}: {c['content']}" for c in page.get("chunks") or []) or "(none)"
    lists = "\n".join(f"- {l}" for l in page.get("lists") or []) or "(none)"
    schema_types = ", ".join(page.get("json_ld_types") or []) or "(none)"
    return PROMPT_TEMPLATE.format(
        url=page["url"], title=page.get("title") or "(none)", h1=page.get("h1") or "(none)",
        schema_types=schema_types, sections=sections, lists=lists,
        content_limit=CHUNK_CONTENT_LIMIT, max_lists=MAX_LISTS,
    )


def validate_prediction(data) -> str | None:
    """Returns None if `data` is a usable prediction, else a short reason it is not."""
    if not isinstance(data, dict):
        return "response is not a JSON object"
    if not isinstance(data.get("primary_entity"), str) or not data["primary_entity"].strip():
        return "missing primary_entity"
    if not isinstance(data.get("prompts"), list) or not data["prompts"]:
        return "missing prompts"
    fan_outs = data.get("fan_outs")
    if not isinstance(fan_outs, list) or not fan_outs:
        return "missing fan_outs"
    for fo in fan_outs:
        if not isinstance(fo, dict) or not REQUIRED_FANOUT_KEYS <= fo.keys():
            return "a fan_out is missing query/type/coverage"
        if not isinstance(fo.get("query"), str) or not fo["query"].strip():
            return "a fan_out has no query text"
        if fo.get("type") not in VALID_TYPES:
            return f"a fan_out has an invalid type: {fo.get('type')!r}"
        if fo.get("coverage") not in VALID_COVERAGE:
            return f"a fan_out has an invalid coverage: {fo.get('coverage')!r}"
    if not isinstance(data.get("gaps", []), list):
        return "gaps is not a list"
    if not isinstance(data.get("follow_ups", []), list):
        return "follow_ups is not a list"
    return None


def gemini_fetch_text(api_key: str, model: str, prompt: str, timeout: int = 120) -> tuple[str | None, str | None, int]:
    """The HTTP half only: one call, returns the raw model text (not yet parsed as JSON).
    Kept separate from JSON parsing/validation so a malformed-JSON response and a real
    HTTP/network failure are retried differently (see gemini_call_with_backoff / predict_one_page)."""
    body = {"contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {"temperature": 0.3, "responseMimeType": "application/json"}}
    status, resp = http_json("POST", f"{GEMINI_BASE}/models/{model}:generateContent",
                             {"x-goog-api-key": api_key}, body, timeout)
    if status != 200:
        return None, f"HTTP {status}: {(resp or {}).get('error')}", status
    try:
        text = resp["candidates"][0]["content"]["parts"][0]["text"]
    except (KeyError, IndexError, TypeError):
        return None, "unexpected response shape (no candidates[0].content.parts[0].text)", status
    return text, None, status


class RpmLimiter:
    """A simple, thread-safe requests-per-minute governor. Shared across the
    ThreadPoolExecutor's workers so `--rpm` bounds the whole run, not just one thread."""

    def __init__(self, rpm: int | None):
        self.rpm = rpm
        self.lock = threading.Lock()
        self.timestamps: deque = deque()

    def wait(self) -> None:
        if not self.rpm:
            return
        with self.lock:
            while True:
                now = time.time()
                while self.timestamps and now - self.timestamps[0] > 60:
                    self.timestamps.popleft()
                if len(self.timestamps) < self.rpm:
                    self.timestamps.append(now)
                    return
                time.sleep(max(0.0, 60 - (now - self.timestamps[0])))


def gemini_text_with_backoff(api_key: str, model: str, prompt: str, limiter: RpmLimiter) -> tuple[str | None, str | None]:
    """Retries 429/503 with exponential backoff. Any other HTTP-level error returns
    immediately - those aren't going to fix themselves on a retry."""
    delay = BACKOFF_BASE
    last_err = None
    for attempt in range(MAX_HTTP_RETRIES):
        limiter.wait()
        text, err, status = gemini_fetch_text(api_key, model, prompt)
        if err is None:
            return text, None
        if status in (429, 503) and attempt < MAX_HTTP_RETRIES - 1:
            last_err = err
            time.sleep(delay)
            delay *= 2
            continue
        return None, err
    return None, last_err


def predict_one_page(api_key: str, model: str, page: dict, limiter: RpmLimiter) -> tuple[dict | None, str | None]:
    """One prediction for one page. Up to MAX_JSON_RETRIES+1 full Gemini calls (each of
    which may itself retry 429/503 with backoff) if the model's response isn't usable
    JSON matching the expected shape - see gemini_text_with_backoff / validate_prediction."""
    prompt = build_predict_prompt(page)
    reason = None
    for json_attempt in range(MAX_JSON_RETRIES + 1):
        text, err = gemini_text_with_backoff(api_key, model, prompt, limiter)
        if err is not None:
            return None, err
        try:
            data = json.loads(text)
        except (json.JSONDecodeError, TypeError):
            reason = "response text was not valid JSON"
            continue
        reason = validate_prediction(data)
        if reason is None:
            return data, None
    return None, f"invalid prediction JSON after {MAX_JSON_RETRIES + 1} attempt(s): {reason}"


def raw_cache_key(url: str, content_hash: str, model: str) -> str:
    return hashlib.sha1(f"{url}|{content_hash}|{model}".encode()).hexdigest()


def cmd_predict(args) -> int:
    env = load_env(args.env_file)
    api_key = env.get("GEMINI_API_KEY")
    if not api_key:
        sys.exit("GEMINI_API_KEY missing (env or --env-file)")

    out = Path(args.out)
    pages = [p for p in load_pages_jsonl(out) if p.get("status") == 200]
    if not pages:
        sys.exit("no successfully fetched pages found in data/pages.jsonl - run `pages` first")

    raw_dir = out / DATA_DIR / RAW_DIR
    raw_dir.mkdir(parents=True, exist_ok=True)
    limiter = RpmLimiter(args.rpm)

    def work(page: dict) -> dict:
        content_hash = page.get("content_hash") or page_content_hash(page)
        cache_file = raw_dir / f"{raw_cache_key(page['url'], content_hash, args.model)}.json"
        if cache_file.is_file():
            cached = json.loads(cache_file.read_text(encoding="utf-8"))
            if not cached.get("error"):
                row = dict(cached["prediction"], url=page["url"], slug=page["slug"],
                          content_hash=content_hash, model=args.model, status="ok",
                          error=None, from_cache=True)
                return row
        try:
            data, err = predict_one_page(api_key, args.model, page, limiter)
        except Exception as e:  # noqa: BLE001 - one bad page must never kill the run
            data, err = None, f"{type(e).__name__}: {e}"
        cache_file.write_text(json.dumps({"prediction": data, "error": err}, indent=2), encoding="utf-8")
        if err:
            return {"url": page["url"], "slug": page["slug"], "content_hash": content_hash,
                    "model": args.model, "status": "error", "error": err, "from_cache": False}
        return dict(data, url=page["url"], slug=page["slug"], content_hash=content_hash,
                    model=args.model, status="ok", error=None, from_cache=False)

    with cf.ThreadPoolExecutor(max_workers=args.concurrency) as pool:
        rows = list(pool.map(work, pages))

    with (out / DATA_DIR / "predictions.jsonl").open("w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")

    n_calls = sum(1 for r in rows if not r.get("from_cache"))
    n_cached = len(rows) - n_calls
    n_errors = sum(1 for r in rows if r.get("status") == "error")
    print(f"Predicted {len(rows)} page(s): {n_calls} new Gemini call(s), {n_cached} served from cache, "
          f"{n_errors} failed (skipped, not fatal). {CREDIT_LINE}")
    for r in rows:
        if r.get("status") == "error":
            print(f"  SKIP {r['url']}: {r['error']}")
    return 0


# ---------------------------------------------------------------- report --

def raw_tokens(text: str) -> set[str]:
    return set(WORD_RE.findall((text or "").lower()))


def content_tokens(text: str) -> set[str]:
    return {t for t in raw_tokens(text) if t not in STOPWORDS and len(t) > 1}


def jaccard(a: set, b: set) -> float:
    if not a and not b:
        return 0.0
    return len(a & b) / len(a | b) if (a | b) else 0.0


def cluster_fanouts(items: list[tuple[str, str, str]], threshold: float = CLUSTER_THRESHOLD) -> list[dict]:
    """items: [(page_url, query, coverage), ...] across the WHOLE site inventory. Greedy
    token-Jaccard clustering, same heuristic the sibling skill uses for its per-page
    fan-out clusters - here applied site-wide to find content gaps and overlaps."""
    clusters: list[dict] = []
    for page_url, query, coverage in items:
        toks = content_tokens(query)
        placed = False
        for c in clusters:
            if jaccard(toks, c["tokens"]) >= threshold:
                c["members"].append((page_url, query, coverage))
                c["tokens"] |= toks
                placed = True
                break
        if not placed:
            clusters.append({"members": [(page_url, query, coverage)], "tokens": set(toks)})
    return clusters


def load_predictions_jsonl(out: Path) -> list[dict]:
    pred_path = out / DATA_DIR / "predictions.jsonl"
    if not pred_path.is_file():
        sys.exit(f"missing {pred_path} - run `predict` first")
    return [json.loads(l) for l in pred_path.read_text(encoding="utf-8").splitlines() if l.strip()]


def compute_site_analysis(predictions: list[dict]) -> dict:
    """The one place site-wide scoring and gap/overlap clustering happens, shared by
    `report` (the human-readable map) and `workbook` (the Fan-Out Map tab + Initiatives)
    so the two never disagree about which pages or clusters are the priority."""
    ok = [p for p in predictions if p.get("status") == "ok"]
    page_scores: dict[str, float] = {}
    all_fanouts: list[tuple[str, str, str]] = []  # (page_url, query, coverage) for site-wide clustering
    for p in ok:
        fan_outs = p.get("fan_outs") or []
        scores = [COVERAGE_SCORE.get(fo.get("coverage"), 0.0) for fo in fan_outs]
        page_scores[p["url"]] = round(sum(scores) / len(scores), 2) if scores else 0.0
        for fo in fan_outs:
            all_fanouts.append((p["url"], fo.get("query", ""), fo.get("coverage", "no")))

    clusters = cluster_fanouts(all_fanouts)
    gap_clusters, overlap_clusters, gap_queries = [], [], set()
    for c in clusters:
        rep = max((m[1] for m in c["members"]), key=len)
        covering_pages = {u for u, q, cov in c["members"] if cov == "yes"}
        entry = {"query": rep, "size": len(c["members"]), "pages": len({u for u, *_ in c["members"]}),
                 "covering_pages": sorted(covering_pages)}
        if not covering_pages:
            gap_clusters.append(entry)
            gap_queries.update(q for _, q, _ in c["members"])
        elif len(covering_pages) >= 2:
            overlap_clusters.append(entry)
    gap_clusters.sort(key=lambda e: -e["size"])
    overlap_clusters.sort(key=lambda e: -len(e["covering_pages"]))

    ranked_pages = sorted(page_scores.items(), key=lambda kv: kv[1])  # lowest coverage first
    avg_score = round(sum(page_scores.values()) / len(page_scores), 2) if page_scores else 0.0
    return {"ok": ok, "n_errors": len(predictions) - len(ok), "page_scores": page_scores,
            "ranked_pages": ranked_pages, "avg_score": avg_score, "gap_clusters": gap_clusters,
            "overlap_clusters": overlap_clusters, "gap_queries": gap_queries}


def cmd_report(args) -> int:
    out = Path(args.out)
    pages = load_pages_jsonl(out)
    pages_by_url = {p["url"]: p for p in pages}
    predictions = load_predictions_jsonl(out)
    analysis = compute_site_analysis(predictions)
    ok = analysis["ok"]

    # own-fan-out flat rows, and a fan-out-map.csv row per page x fan-out
    csv_rows: list[dict] = []
    for p in ok:
        for fo in p.get("fan_outs") or []:
            csv_rows.append({"page_url": p["url"], "slug": p["slug"], "primary_entity": p.get("primary_entity", ""),
                             "query": fo.get("query", ""), "type": fo.get("type", ""),
                             "coverage": fo.get("coverage", ""), "evidence": fo.get("evidence", "")})

    data_dir = out / DATA_DIR
    data_dir.mkdir(parents=True, exist_ok=True)
    with (data_dir / "fan-out-map.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["page_url", "slug", "primary_entity", "query", "type", "coverage", "evidence"])
        w.writeheader()
        w.writerows(csv_rows)

    gap_clusters, overlap_clusters = analysis["gap_clusters"], analysis["overlap_clusters"]
    ranked_pages = analysis["ranked_pages"]
    top_optimise = ranked_pages[: args.top]
    avg_score, n_errors = analysis["avg_score"], analysis["n_errors"]

    honesty = (
        "**These fan-outs are PREDICTED by a model from the page's own content, not observed "
        "searches.** They are biased toward topics already on the page - a model reading a page "
        "tends to imagine searches that page would answer, so it will systematically under-report "
        f"real content gaps. {CREDIT_LINE}. Confirm the pages this report prioritises with a real, "
        "OBSERVED check (a live ChatGPT / Google AI Mode query against the page's buyer prompts) "
        "before writing or optimising anything - see 'Next step' below."
    )

    L = ["# Fan-out map", "", honesty, "", "## Site summary", "",
         f"Pages predicted: {len(ok)} ok, {n_errors} failed | Average page coverage score: {avg_score} "
         "(yes=1, partial=0.5, no=0, averaged across each page's own fan-outs)", "",
         "## Top pages to optimise", "",
         "Lowest average coverage of their own predicted fan-outs - the pages most likely to be "
         "missing content an AI engine would search for.", "",
         "| Page | Coverage score |", "|---|---|"]
    for url, score in top_optimise:
        L.append(f"| {url} | {score} |")
    if not top_optimise:
        L.append("| - | no successfully predicted pages |")

    L += ["", "## Content gaps across the inventory", "",
          "Fan-out clusters (token Jaccard >= 0.6) that "
          f"**no page** in this inventory covers (`coverage: yes`). Showing the top {args.top}, "
          "ranked by how many predicted fan-outs landed in the cluster - these are new-content "
          "candidates, not confirmed demand.", "",
          "| Fan-out cluster | Seen across (fan-out mentions) | Pages that touched it |", "|---|---|---|"]
    for e in gap_clusters[: args.top]:
        L.append(f"| {e['query']} | {e['size']} | {e['pages']} |")
    if not gap_clusters:
        L.append("| - | no uncovered clusters found | - |")

    L += ["", "## Possible overlap", "",
          "Fan-out clusters that 2+ different pages both claim to cover (`coverage: yes`) - worth "
          "checking for cannibalisation or a page that should consolidate/redirect into another.",
          "", "| Fan-out cluster | Pages covering it |", "|---|---|"]
    for e in overlap_clusters[: args.top]:
        L.append(f"| {e['query']} | {', '.join(e['covering_pages'])} |")
    if not overlap_clusters:
        L.append("| - | no overlapping coverage found |")

    L += ["", "## Next step: confirm the priority pages with an OBSERVED check", "",
          "Predicted fan-out is a hypothesis, not an observed search. Before writing or optimising "
          "any page below, confirm it for real: open ChatGPT or Google AI Mode yourself and run each "
          "of the page's buyer prompts (see that page's `pages/<slug>.md`), then note whether the "
          "page - or a competitor - actually gets cited. A paid, DataForSEO-backed observed-fan-out "
          "skill is planned for this pack to automate that check; it hasn't shipped yet, so do this "
          "by hand for now.", "", "Priority pages (lowest predicted coverage first):", ""]
    L += [f"- {url}" for url, _ in top_optimise] or ["- (no priority pages)"]
    L += ["", "## Method and limits", "",
          "- Each page's fan-outs come from one Gemini Flash call (no search grounding) reading "
          "only that page's own title, H1, section headings + up to 500 chars of following "
          "content, up to 5 lists, and JSON-LD @types. It never sees another page on the site, "
          "competitor content, or a real search engine.",
          "- Coverage (`yes`/`partial`/`no`) is the model's own judgement of its own predictions "
          "against the page content it was given - not a citation guarantee and not verified "
          "against a live search.",
          "- Site-wide clustering is a token-Jaccard heuristic (>= 0.6): near-duplicate wording "
          "merges, differently-worded searches for the same real topic may not.",
          "- See `references/evidence.md` for what was verified live and when.", ""]
    (out / "fan-out-map.md").write_text("\n".join(L) + "\n", encoding="utf-8")

    pages_dir = out / "pages"
    pages_dir.mkdir(parents=True, exist_ok=True)
    for p in ok:
        url = p["url"]
        pl = [f"# {url}", "", honesty, "", f"**Primary entity:** {p.get('primary_entity', '')}", "",
              "## Buyer prompts this page should win", ""]
        pl += [f"- {q}" for q in p.get("prompts") or []] or ["(none)"]
        pl += ["", "## Predicted fan-out", "", "| Query | Type | Coverage | Evidence |", "|---|---|---|---|"]
        pl += [f"| {fo.get('query', '')} | {fo.get('type', '')} | {fo.get('coverage', '')} | "
              f"{fo.get('evidence', '') or '-'} |" for fo in p.get("fan_outs") or []] or ["| (none) | - | - | - |"]
        pl += ["", "## Gaps", ""]
        pl += [f"- {g}" for g in p.get("gaps") or []] or ["(none)"]
        pl += ["", "## Likely follow-up questions", ""]
        pl += [f"- {q}" for q in p.get("follow_ups") or []] or ["(none)"]
        (pages_dir / f"{p['slug']}.md").write_text("\n".join(pl) + "\n", encoding="utf-8")

    print(f"Wrote {out / 'fan-out-map.md'}, {len(ok)} page report(s) in {pages_dir}/, and "
          f"{data_dir / 'fan-out-map.csv'} ({len(csv_rows)} row(s)).")
    return 0


# --------------------------------------------------------------- workbook --

WORKBOOK_HONESTY = (
    "These fan-outs are PREDICTED by a model from the page's own content, not observed searches. "
    "They are biased toward topics already on the page. Confirm with a real, observed check (a live "
    "ChatGPT / Google AI Mode query against the page's buyer prompts) before writing or optimising."
)


def build_fanout_map_rows(analysis: dict) -> list[dict]:
    """One row per page x fan-out for the Fan-Out Map tab - the same shape as `report`'s
    fan-out-map.csv, plus the two columns a spreadsheet reviewer needs that the CSV
    doesn't carry: the page's buyer prompts and whether this fan-out's cluster is a
    site-wide gap (no page anywhere in the inventory covers it)."""
    rows = []
    for p in analysis["ok"]:
        prompts_joined = "; ".join(p.get("prompts") or [])
        score = analysis["page_scores"].get(p["url"], 0.0)
        for fo in p.get("fan_outs") or []:
            query = fo.get("query", "")
            rows.append({
                "page_url": p["url"], "primary_entity": p.get("primary_entity", ""), "page_score": score,
                "query": query, "type": fo.get("type", ""), "coverage": fo.get("coverage", ""),
                "evidence": fo.get("evidence", ""), "prompts": prompts_joined,
                "uncovered_site_wide": "Y" if query in analysis["gap_queries"] else "",
            })
    return rows


def build_fanout_map_tab(wb, rows: list[dict], styles: dict):
    if FANOUT_MAP_TAB in wb.sheetnames:
        del wb[FANOUT_MAP_TAB]
    ws = wb.create_sheet(FANOUT_MAP_TAB)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2

    ws.merge_cells("B2:J2")
    banner = ws.cell(row=2, column=2, value=WORKBOOK_HONESTY)
    banner.font, banner.fill, banner.alignment = styles["banner_font"], styles["banner_fill"], styles["wrap"]
    ws.row_dimensions[2].height = 30

    header_row = 4
    headers = ["Page URL", "Primary entity", "Page coverage score", "Fan-out query", "Type", "Coverage",
               "Evidence", "Buyer prompts", "Uncovered site-wide"]
    for i, h in enumerate(headers, start=2):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = styles["head"], styles["head_fill"], styles["wrap"], styles["edge"]
    ws.row_dimensions[header_row].height = 26
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)

    for r, row in enumerate(rows, start=header_row + 1):
        values = [row["page_url"], row["primary_entity"], row["page_score"], row["query"], row["type"],
                  row["coverage"], row["evidence"], row["prompts"], row["uncovered_site_wide"]]
        for i, v in enumerate(values, start=2):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment, c.border = styles["wrap"], styles["edge"]
        fill_colour = COVERAGE_FILL.get(row["coverage"])
        if fill_colour:
            ws.cell(row=r, column=7).fill = styles["fill"](fill_colour)  # column G = Coverage

    last_row = header_row + max(len(rows), 1)
    ws.auto_filter.ref = f"B{header_row}:J{last_row}"
    widths = {"B": 46, "C": 22, "D": 12, "E": 44, "F": 14, "G": 12, "H": 40, "I": 55, "J": 12}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    return ws


def ice_for_optimise(score: float) -> tuple[int, int, int]:
    """Worse predicted coverage -> higher confidence that spending observed-fan-out
    budget on this page is worth it."""
    impact, ease = 7, 5
    confidence = max(1, min(10, round(9 - score * 8)))
    return impact, confidence, ease


def ice_for_gap(size: int, max_size: int) -> tuple[int, int, int]:
    """More predicted mentions of a cluster across the inventory -> higher confidence
    it's a real content gap, not one page's one-off guess."""
    impact, ease = 6, 4
    confidence = max(1, min(10, round(3 + 6 * (size / max(max_size, 1)))))
    return impact, confidence, ease


def build_initiative_tasks(analysis: dict, top_n: int) -> list[dict]:
    tasks = []
    for url, score in analysis["ranked_pages"][:top_n]:
        impact, confidence, ease = ice_for_optimise(score)
        tasks.append({
            "task": f"Optimise {url}: confirm with observed fan-out before editing",
            "category": "Content Strategy",
            "instructions": f"Predicted coverage score {score} across its own fan-outs (Gemini, not "
                            "observed - see fan-out-map.md). Run its buyer prompts through a real "
                            "ChatGPT / Google AI Mode session for OBSERVED fan-out before writing or "
                            "optimising.",
            "from_skill": "mos-geo-fan-out-map", "impact": impact, "confidence": confidence, "ease": ease,
        })
    gap_clusters = analysis["gap_clusters"][:top_n]
    max_size = max((e["size"] for e in gap_clusters), default=1)
    for e in gap_clusters:
        impact, confidence, ease = ice_for_gap(e["size"], max_size)
        tasks.append({
            "task": f"New content candidate: {e['query']}",
            "category": "Content Strategy",
            "instructions": f"Predicted fan-out cluster with no page covering it ({e['size']} mention(s) "
                            f"across the inventory, {e['pages']} page(s) touched it). See fan-out-map.md - "
                            "Content gaps across the inventory.",
            "from_skill": "mos-geo-fan-out-map", "impact": impact, "confidence": confidence, "ease": ease,
        })
    return tasks


def add_initiatives(wb, tasks: list[dict]) -> int:
    """Idempotent: a (Task, From skill) pair already on the Initiatives tab is skipped,
    not re-added, so re-running `workbook` never duplicates a row."""
    if "Initiatives" not in wb.sheetnames:
        return 0
    ws = wb["Initiatives"]
    existing = {(ws.cell(row=r, column=4).value, ws.cell(row=r, column=6).value)
               for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=4).value}
    next_row, n = 6, 0  # template's Initiatives header is row 5; data starts row 6
    for t in tasks:
        key = (t["task"], t["from_skill"])
        if key in existing:
            continue
        while ws.cell(row=next_row, column=4).value:
            next_row += 1
        r = next_row
        ws.cell(row=r, column=2, value=f'=IFERROR(ROUND(J{r}*K{r}/L{r},1),"")')
        ws.cell(row=r, column=3, value=t["category"])
        ws.cell(row=r, column=4, value=t["task"])
        ws.cell(row=r, column=5, value=t["instructions"])
        ws.cell(row=r, column=6, value=t["from_skill"])
        ws.cell(row=r, column=7, value="Scheduled")
        ws.cell(row=r, column=9, value="[AGENCY]")
        ws.cell(row=r, column=10, value=t["impact"])
        ws.cell(row=r, column=11, value=t["confidence"])
        ws.cell(row=r, column=12, value=t["ease"])
        existing.add(key)
        next_row += 1
        n += 1
    return n


def tick_checklist_inline(wb, skill: str, status: str, note: str) -> bool:
    """Same column logic as _shared/brand-audit/tick_checklist.py, applied directly to an
    already-open workbook. Used instead of shelling out to that script when `--workbook`
    points somewhere other than tick_checklist.py's own default (run_dir.parent) location
    - otherwise the subprocess ticks a DIFFERENT file than the one this command just wrote."""
    if "Checklist" not in wb.sheetnames:
        return False
    ws = wb["Checklist"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=3).value == skill:
            ws.cell(row=r, column=8, value=status)
            ws.cell(row=r, column=9, value="☑" if status in ("Client review", "Completed") else "☐")
            ws.cell(row=r, column=10, value=time.strftime("%Y-%m-%d"))
            ws.cell(row=r, column=11, value=note)
            return True
    return False


def cmd_workbook(args) -> int:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        sys.exit("openpyxl is needed: uv run --with openpyxl python fanmap.py workbook ...")

    out = Path(args.out)
    predictions = load_predictions_jsonl(out)
    analysis = compute_site_analysis(predictions)
    rows = build_fanout_map_rows(analysis)

    book = Path(args.workbook) if args.workbook else out.resolve().parent / WORKBOOK
    if not book.is_file():
        if not PACK_TEMPLATE.is_file():
            sys.exit(f"no workbook at {book} and no pack template at {PACK_TEMPLATE}")
        book.parent.mkdir(parents=True, exist_ok=True)
        book.write_bytes(PACK_TEMPLATE.read_bytes())
    wb = openpyxl.load_workbook(book)

    styles = {"wrap": Alignment(wrap_text=True, vertical="top"), "head": Font(bold=True, color="FFFFFF"),
              "head_fill": PatternFill("solid", fgColor="1F2937"),
              "edge": Border(*(Side(style="thin", color="D1D5DB"),) * 4),
              "banner_font": Font(bold=True, italic=True, color="92400E"),
              "banner_fill": PatternFill("solid", fgColor="FEF3C7"),
              "fill": lambda hexcolor: PatternFill("solid", fgColor=hexcolor)}
    build_fanout_map_tab(wb, rows, styles)

    tasks = build_initiative_tasks(analysis, args.top)
    n_init = add_initiatives(wb, tasks)

    order = ["Checklist", "Brand Truth Review", "Brand 360 Report", "AI Visibility", "AI Info Page",
             "Fan-Out", FANOUT_MAP_TAB, "Initiatives"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in order]

    note = (f"{len(analysis['ok'])} page(s) predicted, {len(rows)} fan-out row(s); "
            f"see the '{FANOUT_MAP_TAB}' tab.")
    # tick_checklist.py always derives its target file as run_dir.parent/brand-audit-master.xlsx -
    # the same place `book` defaults to. If --workbook pointed somewhere else, shelling out to it
    # would silently tick a DIFFERENT file than the one we just wrote, so tick in-process instead.
    ticked_inline = False
    if args.workbook:
        ticked_inline = tick_checklist_inline(wb, "mos-geo-fan-out-map", "Client review", note)
    wb.save(book)
    print(f"Wrote {book}: {FANOUT_MAP_TAB} tab ({len(rows)} row(s)), {n_init} new Initiatives row(s) "
          f"({len(tasks) - n_init} already present, skipped).")

    if args.workbook:
        if not ticked_inline:
            print("(Checklist tick skipped: no Checklist row for mos-geo-fan-out-map - "
                  "add it to _shared/brand-audit/skills.json and rebuild the template)")
    else:
        try:
            subprocess.run([sys.executable, str(TICK_SCRIPT), "--skill", "mos-geo-fan-out-map",
                           "--run-dir", str(out.resolve()), "--status", "Client review", "--note", note],
                          check=True, capture_output=True, text=True)
        except Exception as e:  # noqa: BLE001 - the tabs above are already saved either way
            print(f"(Checklist tick skipped: {e})")
    return 0


# ------------------------------------------------------------------- cli --

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="fanmap.py", description=__doc__.splitlines()[0])
    sub = p.add_subparsers(dest="cmd", required=True)

    sp = sub.add_parser("path", help="print this run's folder (MarketingOS-aware)")
    sp.add_argument("--url"); sp.add_argument("--site"); sp.add_argument("--date")
    sp.add_argument("--start", default=".")
    sp.set_defaults(func=cmd_path)

    sp = sub.add_parser("preflight", help="check the Gemini key, model, page count and time estimate")
    sp.add_argument("--env-file")
    sp.add_argument("--model", default=DEFAULT_MODEL)
    sp.add_argument("--pages", type=int, default=0, help="planned page count (if not discovering below)")
    sp.add_argument("--concurrency", type=int, default=4)
    sp.add_argument("--rpm", type=int, default=0)
    sp.add_argument("--sitemap"); sp.add_argument("--include"); sp.add_argument("--limit", type=int)
    sp.add_argument("--urls"); sp.add_argument("--url")
    sp.set_defaults(func=cmd_preflight)

    sp = sub.add_parser("pages", help="fetch + extract layout-aware chunks -> data/pages.jsonl")
    sp.add_argument("--sitemap"); sp.add_argument("--include"); sp.add_argument("--limit", type=int)
    sp.add_argument("--urls"); sp.add_argument("--url")
    sp.add_argument("--out", required=True)
    sp.add_argument("--workers", type=int, default=10)
    sp.add_argument("--refresh", action="store_true")
    sp.set_defaults(func=cmd_pages)

    sp = sub.add_parser("predict", help="one Gemini call per page -> data/predictions.jsonl")
    sp.add_argument("--out", required=True)
    sp.add_argument("--env-file")
    sp.add_argument("--model", default=DEFAULT_MODEL)
    sp.add_argument("--concurrency", type=int, default=4)
    sp.add_argument("--rpm", type=int, default=0)
    sp.set_defaults(func=cmd_predict)

    sp = sub.add_parser("report", help="write fan-out-map.md, pages/<slug>.md, data/fan-out-map.csv")
    sp.add_argument("--out", required=True)
    sp.add_argument("--top", type=int, default=10)
    sp.set_defaults(func=cmd_report)

    sp = sub.add_parser("workbook", help=f"add the {FANOUT_MAP_TAB} tab + Initiatives and tick Checklist (needs openpyxl)")
    sp.add_argument("--out", required=True)
    sp.add_argument("--workbook", help="explicit path to brand-audit-master.xlsx (default: one level "
                    "above --out, created from the pack template if missing)")
    sp.add_argument("--top", type=int, default=5, help="max Initiatives rows per category (pages to "
                    "optimise, new-content candidates)")
    sp.set_defaults(func=cmd_workbook)

    return p


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
