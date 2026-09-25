"""Offline tests for fanout.py. No network, no keys.

Run: python3 scripts/test_fanout.py  (or: python3 -m unittest scripts/test_fanout.py)
"""

import json
import sys
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).parent))
import fanout as f  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"
FANOUT_FIXTURE = json.loads((FIXTURES / "llm_responses_live.json").read_text())
MENTIONS_FIXTURE = json.loads((FIXTURES / "llm_mentions_search_live.json").read_text())
SERP_FIXTURE = json.loads((FIXTURES / "serp_organic_live.json").read_text())
SAMPLE_HTML = (FIXTURES / "sample-page.html").read_text()

ENV = {"DATAFORSEO_LOGIN": "l", "DATAFORSEO_PASSWORD": "p"}
PAGE_URL = "https://example-guides.test/guides/geo/llm-share-buttons/"


def dfs_ok(result, cost=0.01):
    return 200, {"status_code": 20000, "tasks": [{"status_code": 20000, "cost": cost, "result": [result]}]}


def dfs_list_ok(items, cost=0.0):
    """For endpoints where task.result IS the flat array (locations, models lists) -
    not a single result object wrapped in a one-item list like dfs_ok()."""
    return 200, {"status_code": 20000, "tasks": [{"status_code": 20000, "cost": cost, "result": items}]}


LOCATIONS_RESULT = [{"location_code": 2036, "country_iso_code": "AU", "location_type": "Country"},
                    {"location_code": 2840, "country_iso_code": "US", "location_type": "Country"}]


class FanoutTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        f._LOCATION_CACHE.clear()
        patches = [mock.patch.object(f, "load_env", return_value=dict(ENV)), mock.patch.object(f.time, "sleep")]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)
        self.addCleanup(self.tmp.cleanup)

    # ------------------------------------------------------------- pages --

    def test_page_extractor_strips_chrome(self):
        rec = f.extract_page.__wrapped__ if False else None  # placeholder, real call below
        with mock.patch.object(f, "fetch_page", return_value=(200, SAMPLE_HTML.encode(), PAGE_URL)):
            rec = f.extract_page(PAGE_URL)
        self.assertEqual(rec["status"], 200)
        self.assertIn("LLM Share Buttons", rec["title"])
        self.assertEqual(rec["h1"], "LLM share buttons for your GEO page")
        h2s = [h["text"] for h in rec["headings"] if h["level"] == "h2"]
        self.assertIn("What the buttons actually do", h2s)
        self.assertIn("Do they raise citations?", h2s)
        # nav/header/footer text must not leak into the body
        self.assertNotIn("Copyright", rec["main_text"])
        self.assertNotIn("Home", rec["main_text"])
        self.assertNotIn("Contact", rec["main_text"])
        # heading text is captured separately, not duplicated into the body
        self.assertNotIn("What the buttons actually do", rec["main_text"])
        self.assertIn("chatgpt.com/?q=", rec["main_text"])
        self.assertGreater(rec["word_count"], 20)

    def test_pages_command_skips_non_200(self):
        def fake_fetch(url, timeout=20):
            if "broken" in url:
                return 404, b"", url
            return 200, SAMPLE_HTML.encode(), url

        with mock.patch.object(f, "fetch_page", side_effect=fake_fetch):
            urls_file = self.dir / "urls.txt"
            urls_file.write_text(f"{PAGE_URL}\nhttps://example-guides.test/broken/\n")
            args = Namespace(url=None, urls=str(urls_file), sitemap=None, include=None, limit=None,
                             out=str(self.dir / "run"), workers=2, refresh=False)
            with mock.patch("builtins.print"):
                rc = f.cmd_pages(args)
        self.assertEqual(rc, 0)
        index = json.loads((self.dir / "run" / "data" / "pages.json").read_text())
        by_status = {r["status"] for r in index["pages"]}
        self.assertEqual(by_status, {200, 404})
        broken = next(r for r in index["pages"] if r["status"] == 404)
        self.assertEqual(broken["error"], "HTTP 404")

    # ----------------------------------------------------------- prompts --

    def _write_page_index(self, run_dir: Path, page: dict) -> None:
        (run_dir / "data").mkdir(parents=True, exist_ok=True)
        (run_dir / "data" / "pages.json").write_text(json.dumps({"pages": [page]}))

    def _sample_page_record(self) -> dict:
        with mock.patch.object(f, "fetch_page", return_value=(200, SAMPLE_HTML.encode(), PAGE_URL)):
            return f.extract_page(PAGE_URL)

    def test_build_generator_prompt_respects_char_cap(self):
        page = dict(self._sample_page_record(), main_text="x" * 3000)
        prompt = f.build_generator_prompt(page)
        self.assertLessEqual(len(prompt), f.PROMPT_CHAR_LIMIT)
        self.assertIn("Headings:", prompt)

    def test_parse_prompt_json_handles_fences_and_junk(self):
        clean = '[{"text": "what is X", "intent": "what-is"}]'
        fenced = 'Sure, here you go:\n```json\n[{"text": "how to X", "intent": "how-to"}]\n```'
        self.assertEqual(f.parse_prompt_json(clean), [{"text": "what is X", "intent": "what-is"}])
        self.assertEqual(f.parse_prompt_json(fenced), [{"text": "how to X", "intent": "how-to"}])
        self.assertEqual(f.parse_prompt_json("not json at all"), [])

    def test_dedupe_prompts(self):
        items = [{"text": "What is X?", "intent": "a"}, {"text": "what is x?", "intent": "b"},
                 {"text": "How does X work?", "intent": "c"}]
        self.assertEqual(len(f.dedupe_prompts(items)), 2)

    def test_prompts_without_yes_makes_no_calls_and_writes_nothing(self):
        run_dir = self.dir / "run"
        self._write_page_index(run_dir, self._sample_page_record())
        with mock.patch.object(f, "http_json", side_effect=AssertionError("must not call the API")):
            args = Namespace(env_file=None, config=None, out=str(run_dir), country="AU", keyword=None,
                             keywords_csv=None, discover=False, max_generated=8, yes=False)
            with mock.patch("builtins.print"):
                rc = f.cmd_prompts(args)
        self.assertEqual(rc, 0)
        self.assertFalse((run_dir / "data" / "prompts.csv").is_file())

    def test_prompts_generated_and_discover(self):
        run_dir = self.dir / "run"
        self._write_page_index(run_dir, self._sample_page_record())
        generated = [{"text": "What are LLM share buttons?", "intent": "what-is"},
                     {"text": "What are LLM share buttons?", "intent": "what-is"},  # duplicate, dropped
                     {"text": "How do LLM share buttons work?", "intent": "how-to"}]

        def fake_http(method, url, headers=None, body=None, timeout=0):
            if "llm_scraper/locations" in url:
                return dfs_list_ok(LOCATIONS_RESULT)
            if "llm_responses/live" in url:
                self.assertEqual(body[0]["model_name"], "gpt-5.4-nano")
                self.assertFalse(body[0]["web_search"])
                return dfs_ok({"items": [{"type": "message", "sections": [
                    {"type": "text", "text": json.dumps(generated), "annotations": None}]}]}, cost=0.008)
            if "llm_mentions/search/live" in url:
                self.assertEqual(body[0]["target"][0]["keyword"], "LLM share buttons for your GEO page")
                return dfs_ok(MENTIONS_FIXTURE["tasks"][0]["result"][0], cost=0.105)
            raise AssertionError(f"unexpected call: {url}")

        with mock.patch.object(f, "http_json", side_effect=fake_http):
            args = Namespace(env_file=None, config=None, out=str(run_dir), country="AU", keyword=None,
                             keywords_csv=None, discover=True, max_generated=8, yes=True)
            with mock.patch("builtins.print"):
                rc = f.cmd_prompts(args)
        self.assertEqual(rc, 0)
        with open(run_dir / "data" / "prompts.csv", newline="", encoding="utf-8") as fh:
            import csv
            rows = list(csv.DictReader(fh))
        generated_rows = [r for r in rows if r["source"] == "generated"]
        observed_rows = [r for r in rows if r["source"] == "observed"]
        self.assertEqual(len(generated_rows), 2)  # deduped
        self.assertLessEqual(len(observed_rows), 3)
        self.assertTrue(all(r["page_url"] == PAGE_URL for r in rows))
        self.assertTrue(all(len(r["prompt"]) <= f.PROMPT_CHAR_LIMIT for r in rows))

    # --------------------------------------------------------------- run --

    def _prompts_csv(self, run_dir: Path, rows: list[dict]) -> Path:
        import csv
        (run_dir / "data").mkdir(parents=True, exist_ok=True)
        path = run_dir / "data" / "prompts.csv"
        with path.open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=["page_url", "prompt_id", "prompt", "source", "intent"])
            w.writeheader()
            w.writerows(rows)
        return path

    def test_run_is_resumable_with_zero_new_calls_on_replay(self):
        run_dir = self.dir / "run"
        prompts_path = self._prompts_csv(run_dir, [
            {"page_url": PAGE_URL, "prompt_id": "p1", "prompt": "What are LLM share buttons?",
             "source": "generated", "intent": "what-is"},
            {"page_url": PAGE_URL, "prompt_id": "p2", "prompt": "How do LLM share buttons work?",
             "source": "generated", "intent": "how-to"},
        ])
        calls = {"n": 0}

        def fake_http(method, url, headers=None, body=None, timeout=0):
            calls["n"] += 1
            self.assertIn("chat_gpt/llm_responses/live", url)
            self.assertEqual(body[0]["model_name"], "gpt-5.6-luna")  # the default fan-out model
            self.assertTrue(body[0]["web_search"])
            return dfs_ok(FANOUT_FIXTURE["tasks"][0]["result"][0], cost=0.05)

        args = lambda: Namespace(env_file=None, config=None, prompts=str(prompts_path), out=str(run_dir),
                                 country="AU", runs=2, only=None, workers=4, max_spend=25.0)
        with mock.patch.object(f, "http_json", side_effect=fake_http):
            with mock.patch("builtins.print"):
                rc = f.cmd_run(args())
        self.assertEqual(rc, 0)
        self.assertEqual(calls["n"], 4)  # 2 prompts x 1 engine x 2 runs
        lines = (run_dir / "data" / "results.jsonl").read_text().splitlines()
        self.assertEqual(len(lines), 4)
        rows = [json.loads(l) for l in lines]
        self.assertTrue(all(r["engine"] == "chat_gpt" for r in rows))
        self.assertTrue(all(r["fan_out_queries"] for r in rows))
        self.assertTrue(any(r["domain_cited"] for r in rows))

        # Replay: same prompts, same runs. Must make zero new API calls.
        with mock.patch.object(f, "http_json", side_effect=AssertionError("must not call the API again")):
            with mock.patch("builtins.print"):
                rc2 = f.cmd_run(args())
        self.assertEqual(rc2, 0)
        self.assertEqual(calls["n"], 4)  # unchanged
        lines_after = (run_dir / "data" / "results.jsonl").read_text().splitlines()
        self.assertEqual(len(lines_after), 4)  # nothing duplicated either

    def test_run_model_override_switches_model(self):
        run_dir = self.dir / "run"
        prompts_path = self._prompts_csv(run_dir, [
            {"page_url": PAGE_URL, "prompt_id": "p1", "prompt": "What are LLM share buttons?",
             "source": "generated", "intent": "what-is"},
        ])

        def fake_http(method, url, headers=None, body=None, timeout=0):
            self.assertEqual(body[0]["model_name"], "gpt-5.6-terra")  # overridden away from the default
            return dfs_ok(FANOUT_FIXTURE["tasks"][0]["result"][0], cost=0.055)

        args = Namespace(env_file=None, config=None, prompts=str(prompts_path), out=str(run_dir),
                         country="AU", runs=1, only=None, model="gpt-5.6-terra", workers=2, max_spend=25.0)
        with mock.patch.object(f, "http_json", side_effect=fake_http):
            with mock.patch("builtins.print"):
                rc = f.cmd_run(args)
        self.assertEqual(rc, 0)
        rows = [json.loads(l) for l in (run_dir / "data" / "results.jsonl").read_text().splitlines()]
        self.assertEqual(rows[0]["model"], "gpt-5.6-terra")

    def test_run_stops_mid_run_at_spend_cap(self):
        run_dir = self.dir / "run"
        prompts_path = self._prompts_csv(run_dir, [
            {"page_url": PAGE_URL, "prompt_id": "p1", "prompt": "What are LLM share buttons?",
             "source": "generated", "intent": "what-is"},
        ])

        def fake_http(method, url, headers=None, body=None, timeout=0):
            return dfs_ok(FANOUT_FIXTURE["tasks"][0]["result"][0], cost=0.5)

        args = Namespace(env_file=None, config=None, prompts=str(prompts_path), out=str(run_dir),
                         country="AU", runs=5, only=None, workers=1, max_spend=1.0)
        with mock.patch.object(f, "http_json", side_effect=fake_http):
            with mock.patch("builtins.print"):
                rc = f.cmd_run(args)
        self.assertEqual(rc, 0)
        lines = (run_dir / "data" / "results.jsonl").read_text().splitlines()
        # $0.50 per call, $1.00 cap -> exactly 2 calls succeed, the rest are stopped.
        self.assertEqual(len(lines), 2)

    def test_load_prompts_csv_refuses_long_prompt(self):
        run_dir = self.dir / "run"
        path = self._prompts_csv(run_dir, [
            {"page_url": PAGE_URL, "prompt_id": "p1", "prompt": "x" * 501, "source": "generated", "intent": "x"},
        ])
        with self.assertRaises(SystemExit):
            f.load_prompts_csv(str(path))

    # ----------------------------------------------------------- analyse --

    def test_classify_and_coverage(self):
        page = self._sample_page_record()
        own = f.domain_of(PAGE_URL)
        self.assertEqual(f.classify_query("site:platform.openai.com share params", own), "site_vendor")
        self.assertEqual(f.classify_query(f"site:{own} contact page", own), "site_own")
        self.assertEqual(f.classify_query('"chatgpt.com/?q=" prefill composer', own), "exact_string")
        self.assertEqual(f.classify_query("what do the buttons actually do", own), "open")

        self.assertEqual(f.coverage_of('"chatgpt.com/?q=" prefill composer', "exact_string", page), "exact")
        self.assertEqual(f.coverage_of('"not-on-this-page.example" widget', "exact_string", page), "missing")
        self.assertEqual(f.coverage_of("what do the buttons actually do", "open", page), "heading")
        self.assertEqual(f.coverage_of("share button engine cites the page", "open", page), "partial")
        self.assertEqual(f.coverage_of("widget maker installation guide", "open", page), "missing")

    def test_cluster_queries_groups_near_duplicates(self):
        queries = [("chat_gpt", 1, "best AI SEO tools 2025"), ("chat_gpt", 2, "best AI SEO tools 2026"),
                   ("chat_gpt", 3, "site:vendor.example pricing")]
        clusters = f.cluster_queries(queries)
        self.assertEqual(len(clusters), 2)
        sizes = sorted(len(c["members"]) for c in clusters)
        self.assertEqual(sizes, [1, 2])

    def _results_row(self, engine, prompt_id, run, fan_out, cited_urls, page_cited, domain_cited, cost=0.05):
        return {"engine": engine, "model": "gpt-5.6-luna", "page_url": PAGE_URL, "prompt_id": prompt_id,
                "run": run, "fan_out_queries": fan_out, "cited_urls": cited_urls, "page_cited": page_cited,
                "domain_cited": domain_cited, "cost": cost, "error": None, "from_cache": False}

    def test_analyse_and_report_pipeline(self):
        run_dir = self.dir / "run"
        page = self._sample_page_record()
        self._write_page_index(run_dir, page)
        self._prompts_csv(run_dir, [
            {"page_url": PAGE_URL, "prompt_id": "p1", "prompt": "What are LLM share buttons?",
             "source": "generated", "intent": "what-is"},
        ])
        own = f.domain_of(PAGE_URL)
        rows = [
            self._results_row("chat_gpt", "p1", 1,
                              ["site:platform.openai.com share params", '"chatgpt.com/?q=" prefill composer',
                               "what do the buttons actually do", "widget maker installation guide"],
                              [PAGE_URL, "https://vendor-docs.test/api/share-links"], True, True),
            self._results_row("chat_gpt", "p1", 2,
                              ["site:platform.openai.com share params", '"chatgpt.com/?q=" prefill composer',
                               "widget maker installation guide"],
                              ["https://vendor-docs.test/api/share-links"], False, False),
        ]
        with (run_dir / "data" / "results.jsonl").open("w", encoding="utf-8") as fh:
            for r in rows:
                fh.write(json.dumps(r) + "\n")

        args = Namespace(env_file=None, config=None, out=str(run_dir), country="AU", serp=False)
        with mock.patch("builtins.print"):
            rc = f.cmd_analyse(args)
        self.assertEqual(rc, 0)

        import csv
        with (run_dir / "data" / "clusters.csv").open(encoding="utf-8") as fh:
            clusters = list(csv.DictReader(fh))
        by_query = {c["cluster_query"]: c for c in clusters}
        self.assertEqual(by_query["site:platform.openai.com share params"]["type"], "site_vendor")
        self.assertEqual(by_query["site:platform.openai.com share params"]["stability"], "1.0")
        exact = by_query['"chatgpt.com/?q=" prefill composer']
        self.assertEqual(exact["type"], "exact_string")
        self.assertEqual(exact["coverage"], "exact")
        self.assertEqual(exact["stability"], "1.0")
        missing_open = by_query["widget maker installation guide"]
        self.assertEqual(missing_open["type"], "open")
        self.assertEqual(missing_open["coverage"], "missing")
        heading_hit = by_query["what do the buttons actually do"]
        self.assertEqual(heading_hit["coverage"], "heading")
        self.assertEqual(heading_hit["stability"], "0.5")  # only run 1 of 2

        with (run_dir / "data" / "pages-summary.csv").open(encoding="utf-8") as fh:
            summary = list(csv.DictReader(fh))
        self.assertEqual(len(summary), 1)
        self.assertEqual(summary[0]["page_citation_rate"], "0.5")
        self.assertEqual(summary[0]["domain_citation_rate"], "0.5")
        self.assertIn("vendor-docs.test", summary[0]["top_competitor_domains"])

        with mock.patch("builtins.print"):
            rc = f.cmd_report(Namespace(out=str(run_dir), top=20))
        self.assertEqual(rc, 0)
        report = (run_dir / "fan-out-report.md").read_text()
        self.assertIn("# Fan-out report", report)
        self.assertIn("hypothesis for citation lift, not a proven cause", report)
        self.assertIn("widget maker installation guide", report)
        self.assertNotIn("site:platform.openai.com share params", report.split("## Method")[0].split("## Top winnable gaps")[1])
        page_md = (run_dir / "pages" / page["slug"] / "..").resolve()  # sanity: pages dir exists
        page_file = run_dir / "pages" / f"{page['slug']}.md"
        self.assertTrue(page_file.is_file())
        page_text = page_file.read_text()
        self.assertIn("ADD_SECTION", page_text)
        self.assertNotIn("ADD_EXACT_STRING", page_text)  # the exact_string cluster is already covered

    def test_fix_for_all_branches(self):
        base = {"stability": "0.8"}
        self.assertEqual(f.fix_for(dict(base, type="site_vendor", coverage="missing", rank="")), "SKIP_VENDOR")
        self.assertEqual(f.fix_for(dict(base, type="site_own", coverage="missing", rank="")), "")
        self.assertEqual(f.fix_for(dict(base, type="exact_string", coverage="missing", rank="")), "ADD_EXACT_STRING")
        self.assertEqual(f.fix_for(dict(base, type="open", coverage="missing", rank="")), "ADD_SECTION")
        self.assertEqual(f.fix_for(dict(base, type="open", coverage="partial", rank="")), "")
        self.assertEqual(f.fix_for(dict(base, type="open", coverage="partial", rank="24")), "EARN_RETRIEVAL")
        self.assertEqual(f.fix_for(dict(base, type="open", coverage="partial", rank="none")), "EARN_RETRIEVAL")
        self.assertEqual(f.fix_for(dict(base, type="open", coverage="partial", rank="3")), "")

    def test_serp_rank(self):
        def fake_http(method, url, headers=None, body=None, timeout=0):
            if "llm_scraper/locations" in url:
                return dfs_list_ok(LOCATIONS_RESULT)
            self.assertIn("serp/google/organic/live/regular", url)
            return dfs_ok(SERP_FIXTURE["tasks"][0]["result"][0], cost=0.006)

        cfg = f.load_config(None)
        with mock.patch.object(f, "http_json", side_effect=fake_http):
            rank = f.dfs_serp_rank(cfg, dict(ENV), "share links api", "AU", "example-guides.test", {})
        self.assertEqual(rank, "24")

    # --------------------------------------------------------- preflight --

    def test_preflight_refuses_over_budget_without_any_call(self):
        args = Namespace(env_file=None, config=None, country="AU", pages=100, prompts=8, runs=5,
                         engines="chat_gpt", max_spend=25.0)
        with mock.patch.object(f, "http_json", side_effect=AssertionError("must not call the API")):
            with mock.patch("builtins.print"):
                rc = f.cmd_preflight(args)
        self.assertEqual(rc, 2)

    def test_preflight_ok(self):
        def fake_http(method, url, headers=None, body=None, timeout=0):
            if "llm_scraper/locations" in url:
                return dfs_list_ok(LOCATIONS_RESULT)
            if "llm_responses/models" in url:
                return dfs_list_ok([{"model_name": "gpt-5.6-luna"}, {"model_name": "gpt-5.6-terra"},
                                   {"model_name": "gpt-5.4-nano"}])
            raise AssertionError(f"unexpected call: {url}")

        args = Namespace(env_file=None, config=None, country="AU", pages=1, prompts=3, runs=1,
                         engines="chat_gpt", max_spend=25.0)
        with mock.patch.object(f, "http_json", side_effect=fake_http):
            with mock.patch("builtins.print"):
                rc = f.cmd_preflight(args)
        self.assertEqual(rc, 0)

    # -------------------------------------------------------------- path --

    def test_run_dir_for_inside_and_outside_a_brain(self):
        brain = self.dir / "brain"
        (brain / ".mos").mkdir(parents=True)
        (brain / ".mos" / "config.yaml").write_text('{"mode": "in-house"}')
        (brain / "content").mkdir()
        inside = f.run_dir_for("example-guides.test", "2026-09-25", brain / "content")
        self.assertEqual(inside, (brain / "campaigns/geo/2026-09/fan-out").resolve())
        inside.mkdir(parents=True)
        (inside / "fan-out-report.md").write_text("x")
        again = f.run_dir_for("example-guides.test", "2026-09-30", brain)
        self.assertEqual(again.name, "fan-out-2")
        (brain / ".mos" / "config.yaml").write_text("mode: agency\n")
        hq = f.run_dir_for("Acme & Co.", "2026-09-25", brain)
        self.assertEqual(hq, (brain / "campaigns/geo/2026-09/acme-co/fan-out").resolve())
        outside = f.run_dir_for("Acme & Co.", "2026-01-05", self.dir)
        self.assertEqual(outside, (self.dir / "outputs/geo/2026-01/acme-co/fan-out").resolve())
        with self.assertRaises(SystemExit):
            f.run_dir_for("Acme", "25-09-2026", self.dir)

    def test_cmd_path_derives_slug_from_url(self):
        with mock.patch("builtins.print") as p:
            f.cmd_path(Namespace(site=None, url=PAGE_URL, date="2026-09-25", start=str(self.dir)))
        printed = str(p.call_args[0][0]).replace("\\", "/")
        self.assertIn("outputs/geo/2026-09/example-guides-test/fan-out", printed)

    # ---------------------------------------------------------- retest ---

    def test_retest_diffs_citation_rate(self):
        baseline = self.dir / "baseline"
        self._prompts_csv(baseline, [
            {"page_url": PAGE_URL, "prompt_id": "p1", "prompt": "What are LLM share buttons?",
             "source": "generated", "intent": "what-is"},
        ])
        base_rows = [self._results_row("chat_gpt", "p1", 1, ["x"], [], False, False),
                    self._results_row("chat_gpt", "p1", 2, ["x"], [], False, False)]
        with (baseline / "data" / "results.jsonl").open("w", encoding="utf-8") as fh:
            for r in base_rows:
                fh.write(json.dumps(r) + "\n")

        out = self.dir / "retest"

        def fake_http(method, url, headers=None, body=None, timeout=0):
            return dfs_ok(FANOUT_FIXTURE["tasks"][0]["result"][0], cost=0.05)

        args = Namespace(env_file=None, config=None, baseline=str(baseline), out=str(out), country="AU",
                         runs=2, only=None, workers=2, max_spend=25.0)
        with mock.patch.object(f, "http_json", side_effect=fake_http):
            with mock.patch("builtins.print"):
                rc = f.cmd_retest(args)
        self.assertEqual(rc, 0)
        import csv
        with (out / "data" / "retest-diff.csv").open(encoding="utf-8") as fh:
            diff = list(csv.DictReader(fh))
        self.assertEqual(len(diff), 1)
        self.assertEqual(diff[0]["baseline_citation_rate"], "0.0")
        self.assertEqual(diff[0]["new_citation_rate"], "1.0")  # fixture citations include the page's own URL
        self.assertEqual(diff[0]["delta"], "1.0")

    # ---------------------------------------------------------- workbook --

    @unittest.skipUnless(__import__("importlib").util.find_spec("openpyxl"), "openpyxl not installed")
    def test_workbook_adds_fanout_tab(self):
        import openpyxl
        run_dir = self.dir / "2026-09" / "fan-out"
        (run_dir / "data").mkdir(parents=True)
        clusters = [
            {"page_url": PAGE_URL, "prompt_id": "p1", "prompt": "What are LLM share buttons?",
             "cluster_query": "widget maker installation guide", "variants": "1", "stability": "0.8",
             "type": "open", "coverage": "missing", "rank": ""},
            {"page_url": PAGE_URL, "prompt_id": "p1", "prompt": "What are LLM share buttons?",
             "cluster_query": "site:platform.openai.com share params", "variants": "2", "stability": "1.0",
             "type": "site_vendor", "coverage": "missing", "rank": ""},
        ]
        import csv
        with (run_dir / "data" / "clusters.csv").open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=f.CLUSTER_FIELDS)
            w.writeheader()
            w.writerows(clusters)

        with mock.patch("builtins.print"), mock.patch.object(f.subprocess, "run") as run_mock:
            rc = f.cmd_workbook(Namespace(out=str(run_dir)))
        self.assertEqual(rc, 0)
        run_mock.assert_called_once()
        book = self.dir / "2026-09" / f.WORKBOOK
        self.assertTrue(book.is_file())
        wb = openpyxl.load_workbook(book)
        self.assertIn("Fan-Out", wb.sheetnames)
        ws = wb["Fan-Out"]
        self.assertEqual(ws.cell(row=3, column=4).value, "widget maker installation guide")
        # the site_vendor row is unwinnable and gets no Initiative
        rows_with_init = [r for r in range(6, 20) if wb["Initiatives"].cell(row=r, column=4).value]
        self.assertEqual(len(rows_with_init), 1)


if __name__ == "__main__":
    unittest.main()
