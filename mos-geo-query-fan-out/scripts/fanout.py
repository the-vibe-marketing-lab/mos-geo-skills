#!/usr/bin/env python3
"""fanout.py - the data layer for mos-geo-query-fan-out.

For 1 to 500+ pages of a site, finds the searches AI engines actually run (query
fan-out) for the prompts each page should win, measures whether the page is cited
today, and produces a per-page fix list to raise citability. Re-runnable later
(`retest`) to measure lift. Standard library only, except `workbook` (openpyxl,
via uv).

Subcommands
  path        Print where this run's folder belongs (MarketingOS-aware).
  preflight   Check credentials, that the configured model exists at DataForSEO,
              resolve --country, and print a cost estimate for the planned run
              (refuses if it is over --max-spend).
  pages       Fetch and parse the page(s) to test: title, H1, H2/H3, main text.
  prompts     Draft 5-8 buyer prompts per page (generated, optionally + observed
              real questions via --discover). Stops for human review unless --yes.
  run         Send every prompt x --runs x engine live, once, with a hard spend
              cap. Caches every call, so a re-run never re-bills.
  analyse     Cluster fan-out queries, classify them, score whether the page
              already covers each one, and record the citation baseline.
  report      Write fan-out-report.md and one pages/<slug>.md per page.
  workbook    Add a Fan-Out tab and ICE-scored Initiatives to the shared audit
              workbook and tick the Checklist row. Needs openpyxl.
  retest      Re-run an earlier run's exact prompts.csv and diff citation rate.

Run folder layout
  data/pages.json       fetched pages + their extracted text (data/pages/<slug>.json each)
  data/prompts.csv      page_url, prompt_id, prompt, source, intent
  data/results.jsonl    one row per (engine, prompt, run): fan-out, citations, cost
  data/raw/<sha1>.json  the exact raw response for one (engine, model, prompt, run)
  data/clusters.csv     one row per (page, prompt, fan-out cluster)
  data/pages-summary.csv  one row per page: citation baseline, top gap
  fan-out-report.md     the client-facing report
  pages/<slug>.md       per-page prompts, fan-outs and fixes

Credentials come from the environment, or from --env-file:
  DATAFORSEO_LOGIN, DATAFORSEO_PASSWORD

Honesty note: fan-out coverage is a hypothesis for citation lift, not a proven
cause. Every fan-out run is a snapshot - see references/evidence.md.
"""

from __future__ import annotations

import argparse
import base64
import concurrent.futures as cf
import csv
import hashlib
import json
import os
import re
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from argparse import Namespace
from collections import Counter, defaultdict
from html.parser import HTMLParser
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_CONFIG = SKILL_DIR / "config" / "engines.json"
API_UA = "mos-geo-query-fan-out/1.0"
# Some hosts 403 anything that doesn't look like a browser.
PAGE_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
          "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36")
ENV_KEYS = ("DATAFORSEO_LOGIN", "DATAFORSEO_PASSWORD")
DATA_DIR = "data"
RAW_DIR = "raw"
PROMPT_CHAR_LIMIT = 500  # DataForSEO's user_prompt cap - see references/evidence.md
SKILL_FOLDER = "fan-out"
# Fallback when the locations endpoint is unreachable.
COUNTRY_CODES = {"AU": 2036, "US": 2840, "GB": 2826, "NZ": 2554, "CA": 2124, "IE": 2372, "SG": 2702}
# DataForSEO task codes worth one more try: search-engine side errors and timeouts.
DFS_RETRY_CODES = {40101, 40102, 50000, 50301}

WORD_RE = re.compile(r"[a-z0-9][a-z0-9'\-]*")
YEAR_RE = re.compile(r"\b(19|20)\d{2}\b")
STOPWORDS = {
    "the", "a", "an", "of", "for", "to", "in", "on", "and", "or", "is", "are", "what",
    "how", "do", "does", "you", "your", "i", "my", "we", "our", "it", "its", "with",
    "this", "that", "be", "can", "vs", "near", "me", "about", "at", "by", "from", "as",
}

FIX_LABELS = {
    "ADD_EXACT_STRING": "Add this exact phrase to the page, verbatim.",
    "ADD_SECTION": "Add a new H2 matching this fan-out plus a 40-60 word direct answer under it.",
    "EARN_RETRIEVAL": "The page is not in the top 20 organic results for this search - a content "
                      "fix will not help until it ranks.",
    "SKIP_VENDOR": "This search targets another domain's site (site:) and cannot be won by a "
                   "third-party page.",
}


# --------------------------------------------------------------- helpers --

def load_env(env_file: str | None) -> dict:
    env = {}
    if env_file:
        path = Path(env_file).expanduser()
        if not path.is_file():
            sys.exit(f"--env-file not found: {path}")
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, val = line.partition("=")
            env[key.strip().removeprefix("export ").strip()] = val.strip().strip("'\"")
    for key in ENV_KEYS:
        if os.environ.get(key):
            env[key] = os.environ[key]
    return env


def has_dataforseo(env: dict) -> bool:
    return bool(env.get("DATAFORSEO_LOGIN") and env.get("DATAFORSEO_PASSWORD"))


def load_config(path: str | None) -> dict:
    return json.loads(Path(path or DEFAULT_CONFIG).read_text(encoding="utf-8"))


def mask(value: str | None) -> str:
    return f"set (…{value[-4:]})" if value else "MISSING"


def http_json(method, url, headers=None, body=None, timeout=180):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    req.add_header("User-Agent", API_UA)
    if data is not None:
        req.add_header("Content-Type", "application/json")
    for k, v in (headers or {}).items():
        req.add_header(k, v)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            raw = resp.read().decode("utf-8", "replace")
            return resp.status, (json.loads(raw) if raw.strip() else None)
    except urllib.error.HTTPError as e:
        raw = e.read().decode("utf-8", "replace")
        try:
            return e.code, json.loads(raw)
        except ValueError:
            return e.code, {"error": raw[:500]}
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        return 0, {"error": f"{type(e).__name__}: {e}"}


def domain_of(url: str) -> str:
    host = urllib.parse.urlparse(url if "//" in url else f"//{url}").hostname or ""
    host = host.lower()
    return host[4:] if host.startswith("www.") else host


def norm_url_eq(a: str, b: str) -> bool:
    def norm(u: str) -> str:
        u = (u or "").split("#")[0].rstrip("/")
        return re.sub(r"^https?://(www\.)?", "", u, flags=re.I).lower()
    return norm(a) == norm(b)


def citation_list(items) -> list[dict]:
    """Normalise any list of URLs or {url|link|href, title} dicts, deduped."""
    out, seen = [], set()
    for item in items or []:
        if isinstance(item, str):
            url, title = item, ""
        elif isinstance(item, dict):
            url = item.get("url") or item.get("link") or item.get("href") or ""
            title = item.get("title") or item.get("name") or ""
        else:
            continue
        if url and url not in seen:
            seen.add(url)
            out.append({"url": url, "title": title})
    return out


# ------------------------------------------------------------ DataForSEO --

def dfs_request(cfg, env, method, path, body=None, timeout=180, retries=1):
    """Returns (task, result, cost, error). Retries transient failures once."""
    for attempt in range(retries + 1):
        task, result, cost, err, code = _dfs_once(cfg, env, method, path, body, timeout)
        if not err or code not in DFS_RETRY_CODES or attempt == retries:
            return task, result, cost, err
        time.sleep(5)


def _dfs_once(cfg, env, method, path, body, timeout):
    token = base64.b64encode(
        f"{env['DATAFORSEO_LOGIN']}:{env['DATAFORSEO_PASSWORD']}".encode()).decode()
    status, resp = http_json(method, cfg["providers"]["dataforseo"]["base_url"] + path,
                             {"Authorization": f"Basic {token}"}, body, timeout)
    resp = resp or {}
    task = (resp.get("tasks") or [{}])[0] or {}
    if status != 200 or resp.get("status_code") != 20000:
        code = 50000 if status in (0, 500, 502, 503, 504) else resp.get("status_code")
        return task, None, 0.0, f"HTTP {status} / {resp.get('status_code')}: {resp.get('status_message') or resp.get('error')}", code
    cost = float(task.get("cost") or 0)
    if task.get("status_code") != 20000:
        return task, None, cost, f"task {task.get('status_code')}: {task.get('status_message')}", task.get("status_code")
    return task, (task.get("result") or [None])[0], cost, None, 20000


_LOCATION_CACHE: dict = {}


def dfs_location_code(cfg, env, country: str) -> int:
    country = (country or "AU").upper()
    if not _LOCATION_CACHE:
        task, _, _, err = dfs_request(cfg, env, "GET", cfg["providers"]["dataforseo"]["locations_path"], timeout=60)
        for loc in (task.get("result") or []) if not err else []:
            if loc.get("location_type") == "Country" and loc.get("country_iso_code"):
                _LOCATION_CACHE[loc["country_iso_code"]] = loc["location_code"]
        _LOCATION_CACHE.setdefault("_loaded", True)
    code = _LOCATION_CACHE.get(country) or COUNTRY_CODES.get(country)
    if not code:
        sys.exit(f"Unknown --country {country} for DataForSEO")
    return code


def dfs_fanout_call(cfg, env, engine, prompt_text, country):
    """One live LLM Responses call, web_search on. Returns (answer, citations,
    fan_out_queries, cost, error)."""
    pcfg = cfg["providers"]["dataforseo"]
    body = {"user_prompt": prompt_text[:PROMPT_CHAR_LIMIT], "model_name": engine["model"],
            "max_output_tokens": pcfg.get("max_output_tokens", 1200)}
    if not engine.get("always_searches"):
        body["web_search"] = True
    if engine.get("country_field") and country:
        body["web_search_country_iso_code"] = country.upper()
    task, result, cost, err = dfs_request(
        cfg, env, "POST", pcfg["llm_responses_path"].format(se=engine["se"]), [body], timeout=240)
    if err:
        return "", [], [], cost, err
    text, cites = [], []
    for item in (result or {}).get("items") or []:
        if item.get("type") != "message":
            continue
        for sec in item.get("sections") or []:
            if sec.get("type") == "text" and sec.get("text"):
                text.append(sec["text"])
            cites += sec.get("annotations") or []
    fan_out = (result or {}).get("fan_out_queries") or []
    return "".join(text).strip(), citation_list(cites), fan_out, cost, None


def dfs_generator_call(cfg, env, page):
    """Closed-book call on the cheap generator model, asking for buyer prompts as JSON."""
    spec = cfg["generator_model"]
    pcfg = cfg["providers"]["dataforseo"]
    body = {"user_prompt": build_generator_prompt(page), "model_name": spec["model"],
            "web_search": False, "max_output_tokens": pcfg.get("max_output_tokens", 1200)}
    task, result, cost, err = dfs_request(
        cfg, env, "POST", pcfg["llm_responses_path"].format(se=spec["se"]), [body], timeout=240)
    if err:
        return "", cost, err
    text = []
    for item in (result or {}).get("items") or []:
        if item.get("type") != "message":
            continue
        for sec in item.get("sections") or []:
            if sec.get("type") == "text" and sec.get("text"):
                text.append(sec["text"])
    return "".join(text).strip(), cost, None


def dfs_llm_mentions_call(cfg, env, keyword, country, limit=10):
    """Real user questions whose fan-outs contain `keyword` (LLM Mentions, prompt discovery)."""
    pcfg = cfg["providers"]["dataforseo"]
    body = {"target": [{"keyword": keyword, "search_scope": ["fan_out_queries"]}],
            "platform": "chat_gpt", "location_code": dfs_location_code(cfg, env, country),
            "language_code": pcfg.get("language_code", "en"), "limit": limit}
    task, result, cost, err = dfs_request(cfg, env, "POST", pcfg["llm_mentions_path"], [body], timeout=180)
    if err:
        return [], cost, err
    return (result or {}).get("items") or [], cost, None


def dfs_serp_rank(cfg, env, query, country, own_domain, cache):
    """Cached Google organic rank (depth 20) for `own_domain` on `query`, or 'none'/'error'."""
    key = query.strip().lower()
    if key not in cache:
        pcfg = cfg["providers"]["dataforseo"]
        body = {"keyword": query[:PROMPT_CHAR_LIMIT], "location_code": dfs_location_code(cfg, env, country),
                "language_code": pcfg.get("language_code", "en"), "depth": 20}
        _, result, _, err = dfs_request(cfg, env, "POST", pcfg["serp_path"], [body], timeout=180)
        cache[key] = None if err else result
    result = cache[key]
    if result is None:
        return "error"
    for item in result.get("items") or []:
        if item.get("type") == "organic" and domain_of(item.get("url") or "") == own_domain:
            return str(item.get("rank_absolute") or item.get("rank_group") or "?")
    return "none"


def enabled_engines(cfg, only=None):
    if only:
        ids = {x.strip() for x in only.split(",") if x.strip()}
        unknown = ids - {e["id"] for e in cfg["engines"]}
        if unknown:
            sys.exit(f"--only: unknown engine id(s): {', '.join(sorted(unknown))}")
        return [e for e in cfg["engines"] if e["id"] in ids]
    return [e for e in cfg["engines"] if e.get("enabled", True)]


# ------------------------------------------------------------------ path --

def _main_checkout(git_file: Path) -> Path | None:
    """The main checkout behind a linked worktree's `.git` file
    ("gitdir: <main>/.git/worktrees/<name>"), including a Windows path read from WSL."""
    m = re.match(r"gitdir:\s*(.+)", git_file.read_text(encoding="utf-8").strip())
    if not m:
        return None
    raw = m.group(1).strip()
    drive = re.match(r"^([A-Za-z]):[\\/](.*)$", raw)
    if drive and os.name != "nt":
        raw = f"/mnt/{drive.group(1).lower()}/{drive.group(2)}"
    gitdir = Path(raw.replace("\\", "/"))
    if not gitdir.is_absolute():
        gitdir = (git_file.parent / gitdir).resolve()
    for d in [gitdir, *gitdir.parents]:
        if d.name == ".git":
            return d.parent
    return None


def brain_config(root: Path) -> Path | None:
    own = root / ".mos" / "config.yaml"
    if own.is_file():
        return own
    git = root / ".git"
    if git.is_file():  # linked worktree: .mos/ is gitignored, so borrow the main checkout's
        main = _main_checkout(git)
        if main and (main / ".mos" / "config.yaml").is_file():
            return main / ".mos" / "config.yaml"
    return None


def find_brain(start: Path) -> Path | None:
    for d in [start, *start.parents]:
        if brain_config(d):
            return d
        if (d / ".git").exists():
            return None
    return None


def brain_mode(brain: Path) -> str:
    config = brain_config(brain)
    if not config:
        return "in-house"
    m = re.search(r'["\']?mode["\']?\s*:\s*["\']?([a-z-]+)', config.read_text(encoding="utf-8"))
    return m.group(1) if m else "in-house"


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")


def run_dir_for(site: str, day: str, start: Path) -> Path:
    """campaigns/geo/YYYY-MM/fan-out/ in a one-brand brain, with a brand folder in an
    agency brain, outputs/geo/YYYY-MM/<site>/fan-out/ outside a brain. A repeat run in
    the same month gets fan-out-2, -3 ... so nothing is overwritten."""
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", day):
        sys.exit(f"--date must be YYYY-MM-DD, got {day}")
    slug = slugify(site) or "site"
    brain = find_brain(start.resolve())
    if brain and brain_mode(brain) != "agency":
        month = brain / "campaigns" / "geo" / day[:7]
    elif brain:
        month = brain / "campaigns" / "geo" / day[:7] / slug
    else:
        month = start.resolve() / "outputs" / "geo" / day[:7] / slug
    target, n = month / SKILL_FOLDER, 2
    while target.exists() and any(target.iterdir()):
        target = month / f"{SKILL_FOLDER}-{n}"
        n += 1
    return target


def cmd_path(args) -> int:
    day = args.date or time.strftime("%Y-%m-%d")
    site = args.site or (domain_of(args.url) if args.url else "site")
    print(run_dir_for(site, day, Path(args.start or ".")))
    return 0


# ------------------------------------------------------------- preflight --

def cmd_preflight(args) -> int:
    env = load_env(args.env_file)
    cfg = load_config(args.config)
    n_engines = len({x.strip() for x in args.engines.split(",") if x.strip()})

    print("mos-geo-query-fan-out preflight")
    print("==========================")
    for key in ENV_KEYS:
        print(f"{key:<21} {mask(env.get(key))}")

    est = args.pages * args.prompts * args.runs * n_engines * cfg["cost_estimate"]["engine_call"]
    print(f"\nPlanned run: {args.pages} page(s) x {args.prompts} prompt(s) x {args.runs} run(s) x "
          f"{n_engines} engine(s) @ ~${cfg['cost_estimate']['engine_call']:.3f}/call = ${est:.3f} estimated.")
    if est > args.max_spend:
        print(f"\nREFUSED: estimated ${est:.3f} is over --max-spend ${args.max_spend:.2f}. "
              "Lower --pages/--prompts/--runs, raise --max-spend, or plan a smaller run.")
        return 2

    if not has_dataforseo(env):
        print("\nDataForSEO credentials missing (DATAFORSEO_LOGIN / DATAFORSEO_PASSWORD).")
        return 2

    problems = 0
    print("\nDataForSEO models")
    checked_se = set()
    engine_specs = [dict(e) for e in cfg["engines"] if e.get("enabled", True)]
    engine_specs.append(dict(cfg["generator_model"], id="generator", label="generator (prompts)"))
    for e in engine_specs:
        se = e["se"]
        path = cfg["providers"]["dataforseo"]["models_path"].format(se=se)
        task, _, _, err = dfs_request(cfg, env, "GET", path, timeout=60)
        if err:
            print(f"  [FAIL] {e.get('label', se):<20} {err}")
            problems += 1
            continue
        live = {x["model_name"] for x in (task.get("result") or [])}
        ok = e["model"] in live
        problems += 0 if ok else 1
        print(f"  [{'ok' if ok else 'GONE'}] {e.get('label', se):<20} {e['model']}"
              + ("" if ok else f"  (available: {', '.join(list(live)[:8])} ...)"))
        checked_se.add(se)

    code = dfs_location_code(cfg, env, args.country)
    print(f"  [ok] country {args.country.upper()} -> location_code {code}")

    off = [e["label"] for e in cfg["engines"] if not e.get("enabled", True)]
    if off:
        print(f"  (switched off: {', '.join(off)})")

    print(f"\n{'Ready.' if problems == 0 else f'{problems} problem(s) - fix before running.'}")
    return 0 if problems == 0 else 2


# ----------------------------------------------------------------- pages --

# Tags whose descendant text never belongs to the article: chrome, not content.
BODY_SKIP = {"script", "style", "noscript", "template", "svg", "iframe", "nav", "header", "footer", "aside"}
BLOCKY = {"p", "div", "li", "tr", "table", "ul", "ol", "br", "section", "article", "blockquote"}


class PageExtractor(HTMLParser):
    """Pulls title, H1, H2/H3 and the main body text out of one HTML page, stdlib
    only. Skips nav/header/footer/aside/script/style/noscript entirely, and keeps
    heading text out of the body text so heading-coverage and body-coverage are
    two different signals."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title = ""
        self.h1 = ""
        self.headings: list[tuple[str, str]] = []
        self._body: list[str] = []
        self._skip: list[str] = []
        self._in_title = False
        self._heading_tag: str | None = None
        self._heading_buf: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in BODY_SKIP:
            self._skip.append(tag)
            return
        if self._skip:
            return
        if tag == "title":
            self._in_title = True
        elif tag in ("h1", "h2", "h3"):
            self._heading_tag, self._heading_buf = tag, []
        elif tag in BLOCKY:
            self._body.append(" ")

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in BODY_SKIP:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if self._skip:
            if self._skip[-1] == tag:
                self._skip.pop()
            return
        if tag == "title":
            self._in_title = False
        elif self._heading_tag == tag:
            text = " ".join("".join(self._heading_buf).split())
            if text:
                if tag == "h1" and not self.h1:
                    self.h1 = text
                elif tag in ("h2", "h3"):
                    self.headings.append((tag, text))
            self._heading_tag = None
        elif tag in BLOCKY:
            self._body.append(" ")

    def handle_data(self, data):
        if self._skip or not data:
            return
        if self._in_title:
            self.title += data
        elif self._heading_tag:
            self._heading_buf.append(data)
        else:
            self._body.append(data)

    def main_text(self) -> str:
        return re.sub(r"\s+", " ", "".join(self._body)).strip()


def slug_of(url: str) -> str:
    p = urllib.parse.urlsplit(url)
    path = p.path.strip("/")
    raw = f"{domain_of(url)}-{path}" if path else domain_of(url) or "page"
    return re.sub(r"[^a-z0-9]+", "-", raw.lower()).strip("-") or "page"


def fetch_page(url: str, timeout: int = 20) -> tuple[int, bytes, str]:
    req = urllib.request.Request(url, headers={"User-Agent": PAGE_UA,
                                                "Accept": "text/html,application/xhtml+xml;q=0.9,*/*;q=0.8"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, r.read(), r.geturl()
    except urllib.error.HTTPError as e:
        return e.code, b"", url
    except Exception as e:  # noqa: BLE001 - network errors become a logged reason
        return 0, str(e).encode(), url


def fetch_sitemap_urls(url: str, include: str | None, limit: int | None, _seen=None) -> list[str]:
    seen = _seen if _seen is not None else set()
    if url in seen:
        return []
    seen.add(url)
    status, body, _ = fetch_page(url)
    if status != 200:
        sys.exit(f"could not fetch sitemap {url}: HTTP {status}")
    text = body.decode("utf-8", "replace")
    locs = re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", text)
    urls = []
    for loc in locs:
        if "sitemap" in loc.lower() and loc.lower().endswith((".xml", ".xml.gz")):
            urls += fetch_sitemap_urls(loc, None, None, seen)
        else:
            urls.append(loc)
    if include:
        pat = re.compile(include)
        urls = [u for u in urls if pat.search(u)]
    if limit:
        urls = urls[:limit]
    return urls


def extract_page(url: str) -> dict:
    slug = slug_of(url)
    status, body, final_url = fetch_page(url)
    rec = {"url": url, "final_url": final_url, "status": status, "slug": slug}
    if status != 200:
        rec["error"] = f"HTTP {status}" if status else (body.decode("utf-8", "replace") or "fetch failed")
        rec.update(title="", h1="", headings=[], main_text="", word_count=0)
        return rec
    ext = PageExtractor()
    try:
        ext.feed(body.decode("utf-8", "replace"))
    except Exception as e:  # noqa: BLE001 - a malformed page still reports what it got
        rec["error"] = f"parse warning: {e}"
    main_text = ext.main_text()
    rec.update(title=ext.title.strip(), h1=ext.h1,
               headings=[{"level": lvl, "text": t} for lvl, t in ext.headings],
               main_text=main_text, word_count=len(main_text.split()))
    return rec


def cmd_pages(args) -> int:
    if args.url:
        urls = [args.url]
    elif args.urls:
        urls = [l.strip() for l in Path(args.urls).read_text(encoding="utf-8").splitlines()
                if l.strip() and not l.strip().startswith("#")]
        if args.limit:
            urls = urls[: args.limit]
    elif args.sitemap:
        urls = fetch_sitemap_urls(args.sitemap, args.include, args.limit)
    else:
        sys.exit("give --url, --urls or --sitemap")
    if not urls:
        sys.exit("no URLs to fetch")

    out = Path(args.out)
    pages_dir = out / DATA_DIR / "pages"
    pages_dir.mkdir(parents=True, exist_ok=True)

    def work(url: str) -> dict:
        cache = pages_dir / f"{slug_of(url)}.json"
        if cache.is_file() and not args.refresh:
            return json.loads(cache.read_text(encoding="utf-8"))
        rec = extract_page(url)
        cache.write_text(json.dumps(rec, indent=2), encoding="utf-8")
        return rec

    with cf.ThreadPoolExecutor(max_workers=args.workers) as pool:
        records = list(pool.map(work, urls))

    (out / DATA_DIR / "pages.json").write_text(json.dumps({"pages": records}, indent=2), encoding="utf-8")
    ok = [r for r in records if r.get("status") == 200]
    print(f"Fetched {len(records)} page(s): {len(ok)} ok, {len(records) - len(ok)} skipped.")
    for r in records:
        if r.get("status") != 200:
            print(f"  SKIP {r['url']}: {r.get('error')}")
    return 0 if ok else 2


# --------------------------------------------------------------- prompts --

def build_generator_prompt(page: dict) -> str:
    """Title + headings + as much of the page excerpt as fits under DataForSEO's
    500-char user_prompt cap. Headings are prioritised over the excerpt: they carry
    more signal per character than running body text."""
    intents = "what-is, how-to, best-compare, problem, cost"
    head = f"Buyer prompts for a page titled '{(page.get('title') or page.get('h1') or '').strip()}'."
    heads = "; ".join(h["text"] for h in (page.get("headings") or [])[:6])
    instr = (f" Headings: {heads}. Write 8 short buyer search prompts covering {intents}. "
             'JSON array only, no prose, no markdown fences: [{"text":"...","intent":"..."}]')
    budget = PROMPT_CHAR_LIMIT - len(head) - len(instr) - len(" Excerpt: .")
    excerpt = (page.get("main_text") or "")[: max(budget, 0)]
    prompt = f"{head} Excerpt: {excerpt}.{instr}" if excerpt else f"{head}{instr}"
    return prompt[:PROMPT_CHAR_LIMIT]


def parse_prompt_json(text: str) -> list[dict]:
    m = re.search(r"\[.*\]", text or "", re.S)
    raw = m.group(0) if m else (text or "")
    try:
        data = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return []
    out = []
    for item in data if isinstance(data, list) else []:
        if isinstance(item, dict) and item.get("text"):
            out.append({"text": str(item["text"]).strip()[:PROMPT_CHAR_LIMIT],
                        "intent": str(item.get("intent") or "open").strip()})
    return out


def dedupe_prompts(items: list[dict]) -> list[dict]:
    seen, out = set(), []
    for it in items:
        key = re.sub(r"\s+", " ", it["text"].lower()).strip()
        if key and key not in seen:
            seen.add(key)
            out.append(it)
    return out


def cmd_prompts(args) -> int:
    env = load_env(args.env_file)
    cfg = load_config(args.config)
    out = Path(args.out)
    index_path = out / DATA_DIR / "pages.json"
    if not index_path.is_file():
        sys.exit(f"missing {index_path} - run `pages` first")
    pages = [p for p in json.loads(index_path.read_text(encoding="utf-8"))["pages"] if p.get("status") == 200]
    if not pages:
        sys.exit("no successfully fetched pages found - run `pages` first")

    keywords: dict[str, str] = {}
    if args.keyword and len(pages) == 1:
        keywords[pages[0]["url"]] = args.keyword
    if args.keywords_csv:
        with open(args.keywords_csv, newline="", encoding="utf-8-sig") as fh:
            for row in csv.DictReader(fh):
                if row.get("url") and row.get("keyword"):
                    keywords[row["url"]] = row["keyword"]

    n_generated = len(pages)
    n_mentions = len(pages) if args.discover else 0
    est = (n_generated * cfg["cost_estimate"]["generator_call"]
           + n_mentions * cfg["cost_estimate"]["llm_mentions_call"])
    print(f"Plan: {n_generated} generator call(s)"
          + (f" + {n_mentions} LLM Mentions call(s)" if args.discover else "")
          + f". Estimated cost ${est:.3f}.")
    if not args.yes:
        print("Pass --yes to spend this and write data/prompts.csv.")
        return 0

    rows: list[dict] = []
    total_cost = 0.0
    for page in pages:
        gen_text, cost, err = dfs_generator_call(cfg, env, page)
        total_cost += cost
        if err:
            print(f"  [generated FAILED] {page['url']}: {err}")
            items = []
        else:
            items = dedupe_prompts(parse_prompt_json(gen_text))
        for i, item in enumerate(items[: args.max_generated], 1):
            rows.append({"page_url": page["url"], "prompt_id": f"{page['slug']}-g{i:02d}",
                        "prompt": item["text"], "source": "generated", "intent": item["intent"]})

        if args.discover:
            kw = (keywords.get(page["url"]) or page.get("h1") or page.get("title") or "").strip()[:80]
            if not kw:
                continue
            mitems, mcost, merr = dfs_llm_mentions_call(cfg, env, kw, args.country)
            total_cost += mcost
            if merr:
                print(f"  [observed FAILED] {page['url']}: {merr}")
                continue
            seen, j = set(), 0
            for m in mitems:
                q = (m.get("question") or "").strip()
                key = q.lower()
                if q and key not in seen and len(q) <= PROMPT_CHAR_LIMIT:
                    seen.add(key)
                    j += 1
                    rows.append({"page_url": page["url"], "prompt_id": f"{page['slug']}-o{j:02d}",
                                "prompt": q, "source": "observed", "intent": "observed"})
                if j >= 3:
                    break

    data_dir = out / DATA_DIR
    data_dir.mkdir(parents=True, exist_ok=True)
    path = data_dir / "prompts.csv"
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["page_url", "prompt_id", "prompt", "source", "intent"])
        w.writeheader()
        w.writerows(rows)
    print(f"Wrote {len(rows)} prompt(s) to {path} (spent ${total_cost:.3f}). Review it, then run `run`.")
    return 0


# ------------------------------------------------------------------- run --

def load_prompts_csv(path: str) -> list[dict]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"{path} has no prompt rows")
    for r in rows:
        if not r.get("prompt") or not r.get("prompt_id") or not r.get("page_url"):
            sys.exit(f"prompts.csv row missing page_url/prompt_id/prompt: {r}")
        if len(r["prompt"]) > PROMPT_CHAR_LIMIT:
            sys.exit(f"{r['prompt_id']} is over {PROMPT_CHAR_LIMIT} characters (DataForSEO limit)")
    return rows


def raw_key(engine_id: str, model: str, prompt_id: str, run_index: int) -> str:
    return hashlib.sha1(f"{engine_id}|{model}|{prompt_id}|{run_index}".encode()).hexdigest()


def build_record(engine: dict, prompt: dict, run_index: int, raw: dict, from_cache: bool) -> dict:
    page_url = prompt["page_url"]
    page_domain = domain_of(page_url)
    cites = raw.get("citations") or []
    cited_urls = [c["url"] for c in cites]
    return {
        "engine": engine["id"], "model": engine["model"], "page_url": page_url,
        "prompt_id": prompt["prompt_id"], "run": run_index,
        "fan_out_queries": raw.get("fan_out_queries") or [],
        "cited_urls": cited_urls,
        "page_cited": any(norm_url_eq(u, page_url) for u in cited_urls),
        "domain_cited": any(domain_of(u) == page_domain for u in cited_urls),
        "cost": round(raw.get("cost") or 0.0, 6), "error": raw.get("error"), "from_cache": from_cache,
    }


def cmd_run(args) -> int:
    env = load_env(args.env_file)
    cfg = load_config(args.config)
    prompts = load_prompts_csv(args.prompts)
    engines = [dict(e) for e in enabled_engines(cfg, args.only)]
    if not engines:
        sys.exit("no engines enabled (check config/engines.json or --only)")
    if getattr(args, "model", None):
        for e in engines:
            e["model"] = args.model

    out = Path(args.out)
    data = out / DATA_DIR
    raw_dir = data / RAW_DIR
    raw_dir.mkdir(parents=True, exist_ok=True)
    results_path = data / "results.jsonl"

    done: set[tuple[str, str, int]] = set()
    spent = 0.0
    if results_path.is_file():
        for line in results_path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            r = json.loads(line)
            spent += r.get("cost") or 0.0
            if not r.get("error"):
                done.add((r["engine"], r["prompt_id"], r["run"]))

    jobs = [(e, p, i) for e in engines for p in prompts for i in range(1, args.runs + 1)
            if (e["id"], p["prompt_id"], i) not in done]
    print(f"Plan: {len(jobs)} new call(s) ({len(prompts)} prompt(s) x {len(engines)} engine(s) x "
          f"{args.runs} run(s), {len(done)} already done). Spent so far: ${spent:.3f}. "
          f"Cap: ${args.max_spend:.2f}.")

    state = {"spent": spent, "stopped": spent >= args.max_spend}
    lock = threading.Lock()

    def worker(job) -> dict:
        engine, prompt, run_index = job
        key = raw_key(engine["id"], engine["model"], prompt["prompt_id"], run_index)
        cache_file = raw_dir / f"{key}.json"
        if cache_file.is_file():
            cached = json.loads(cache_file.read_text(encoding="utf-8"))
            if not cached.get("error"):
                return build_record(engine, prompt, run_index, cached, from_cache=True)
        with lock:
            if state["stopped"]:
                return {"skipped": True}
        answer, cites, fan_out, cost, err = dfs_fanout_call(cfg, env, engine, prompt["prompt"], args.country)
        raw = {"answer": answer, "citations": cites, "fan_out_queries": fan_out, "cost": cost, "error": err}
        cache_file.write_text(json.dumps(raw, indent=2), encoding="utf-8")
        with lock:
            state["spent"] += cost
            if state["spent"] >= args.max_spend:
                state["stopped"] = True
        return build_record(engine, prompt, run_index, raw, from_cache=False)

    new_records = []
    with cf.ThreadPoolExecutor(max_workers=args.workers) as pool:
        for rec in pool.map(worker, jobs):
            if not rec.get("skipped"):
                new_records.append(rec)

    with results_path.open("a", encoding="utf-8") as fh:
        for rec in new_records:
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")

    n_calls = sum(1 for r in new_records if not r.get("from_cache"))
    n_errors = sum(1 for r in new_records if r.get("error"))
    n_stopped = len(jobs) - len(new_records)
    print(f"Saved {len(new_records)} result(s): {n_calls} new API call(s), "
          f"{len(new_records) - n_calls} served from cache, {n_errors} failed, "
          f"{n_stopped} stopped by the spend cap. Total spend so far: ${state['spent']:.3f}.")
    return 0


# --------------------------------------------------------------- analyse --

def normalize_query(q: str) -> str:
    text = YEAR_RE.sub("", (q or "").lower())
    return re.sub(r"\s+", " ", text).strip()


def raw_tokens(text: str) -> set[str]:
    return set(WORD_RE.findall(normalize_query(text)))


def content_tokens(text: str) -> set[str]:
    return {t for t in raw_tokens(text) if t not in STOPWORDS and len(t) > 1}


def jaccard(a: set, b: set) -> float:
    if not a and not b:
        return 0.0
    return len(a & b) / len(a | b) if (a | b) else 0.0


def classify_query(q: str, own_domain: str) -> str:
    m = re.search(r"site:\s*([a-z0-9.\-]+)", (q or "").lower())
    if m:
        dom = m.group(1).lstrip(".")
        dom = dom[4:] if dom.startswith("www.") else dom
        if own_domain and (dom == own_domain or dom.endswith("." + own_domain)):
            return "site_own"
        return "site_vendor"
    if re.search(r'"[^"]{3,}"', q or "") or re.search(r"'[^']{3,}'", q or ""):
        return "exact_string"
    return "open"


def extract_quoted(q: str) -> list[str]:
    return re.findall(r'"([^"]{3,})"', q or "") + re.findall(r"'([^']{3,})'", q or "")


def cluster_queries(queries: list[tuple[str, int, str]], threshold: float = 0.6) -> list[dict]:
    """queries: [(engine, run, text), ...] for one (page, prompt). Greedy clustering:
    each query joins the first existing cluster whose tokens overlap >= threshold,
    else starts a new one."""
    clusters: list[dict] = []
    for engine, run, text in queries:
        toks = raw_tokens(text)
        placed = False
        for c in clusters:
            if jaccard(toks, c["tokens"]) >= threshold:
                c["members"].append(text)
                c["runs"].add((engine, run))
                c["tokens"] |= toks
                placed = True
                break
        if not placed:
            clusters.append({"members": [text], "runs": {(engine, run)}, "tokens": set(toks)})
    return clusters


def coverage_of(cluster_rep: str, cluster_type: str, page: dict) -> str:
    if cluster_type == "exact_string":
        quoted = extract_quoted(cluster_rep)
        haystack = ((page.get("main_text") or "") + " " + (page.get("title") or "")).lower()
        if quoted and any(q.lower() in haystack for q in quoted):
            return "exact"
    q_tokens = content_tokens(cluster_rep)
    if not q_tokens:
        return "missing"
    for h in page.get("headings") or []:
        h_text = h["text"] if isinstance(h, dict) else h[1]
        h_tokens = content_tokens(h_text)
        if h_tokens and len(q_tokens & h_tokens) / len(q_tokens) >= 0.6:
            return "heading"
    body_tokens = content_tokens(page.get("main_text") or "")
    if body_tokens and len(q_tokens & body_tokens) / len(q_tokens) >= 0.5:
        return "partial"
    return "missing"


CLUSTER_FIELDS = ["page_url", "prompt_id", "prompt", "cluster_query", "variants", "stability",
                  "type", "coverage", "rank"]
SUMMARY_FIELDS = ["page_url", "word_count", "num_prompts", "num_runs", "page_citation_rate",
                  "domain_citation_rate", "top_competitor_domains", "top_winnable_gap"]


def cmd_analyse(args) -> int:
    env = load_env(args.env_file) if args.serp else {}
    cfg = load_config(args.config)
    out = Path(args.out)
    data = out / DATA_DIR

    pages_index = json.loads((data / "pages.json").read_text(encoding="utf-8"))
    pages_by_url = {p["url"]: p for p in pages_index["pages"]}
    with open(data / "prompts.csv", newline="", encoding="utf-8-sig") as fh:
        prompts_by_id = {row["prompt_id"]: row for row in csv.DictReader(fh)}

    rows = []
    results_path = data / "results.jsonl"
    if results_path.is_file():
        rows = [json.loads(l) for l in results_path.read_text(encoding="utf-8").splitlines() if l.strip()]
    rows = [r for r in rows if not r.get("error")]

    by_prompt = defaultdict(list)
    for r in rows:
        by_prompt[r["prompt_id"]].append(r)

    serp_cache: dict = {}
    cluster_rows: list[dict] = []
    page_stats = defaultdict(lambda: {"runs": 0, "page_cited": 0, "domain_cited": 0,
                                      "competitors": Counter(), "clusters": []})

    for prompt_id, prows in by_prompt.items():
        prompt = prompts_by_id.get(prompt_id)
        if not prompt:
            continue
        page = pages_by_url.get(prompt["page_url"], {})
        own_domain = domain_of(prompt["page_url"])
        queries = [(r["engine"], r["run"], q) for r in prows for q in (r.get("fan_out_queries") or [])]
        total_runs = len({(r["engine"], r["run"]) for r in prows}) or 1
        clusters = cluster_queries(queries)

        ps = page_stats[prompt["page_url"]]
        ps["runs"] += total_runs
        ps["page_cited"] += sum(1 for r in prows if r.get("page_cited"))
        ps["domain_cited"] += sum(1 for r in prows if r.get("domain_cited"))
        for r in prows:
            for u in r.get("cited_urls") or []:
                d = domain_of(u)
                if d and d != own_domain:
                    ps["competitors"][d] += 1

        for c in clusters:
            rep = max(c["members"], key=len)
            ctype = classify_query(rep, own_domain)
            stability = round(len(c["runs"]) / total_runs, 2)
            coverage = coverage_of(rep, ctype, page)
            rank = ""
            if args.serp and ctype in ("open", "exact_string") and stability >= 0.5:
                rank = dfs_serp_rank(cfg, env, rep, args.country, own_domain, serp_cache)
            crow = {"page_url": prompt["page_url"], "prompt_id": prompt_id, "prompt": prompt["prompt"],
                    "cluster_query": rep, "variants": len(c["members"]), "stability": stability,
                    "type": ctype, "coverage": coverage, "rank": rank}
            cluster_rows.append(crow)
            ps["clusters"].append(crow)

    data.mkdir(parents=True, exist_ok=True)
    with (data / "clusters.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=CLUSTER_FIELDS)
        w.writeheader()
        w.writerows(cluster_rows)

    summary_rows = []
    for url, ps in page_stats.items():
        page = pages_by_url.get(url, {})
        top_comp = ", ".join(f"{d} ({n})" for d, n in ps["competitors"].most_common(5))
        gaps = sorted((c for c in ps["clusters"] if c["type"] != "site_vendor" and c["coverage"] == "missing"),
                     key=lambda c: -c["stability"])
        summary_rows.append({
            "page_url": url, "word_count": page.get("word_count", 0),
            "num_prompts": len({c["prompt_id"] for c in ps["clusters"]}), "num_runs": ps["runs"],
            "page_citation_rate": round(ps["page_cited"] / ps["runs"], 2) if ps["runs"] else 0.0,
            "domain_citation_rate": round(ps["domain_cited"] / ps["runs"], 2) if ps["runs"] else 0.0,
            "top_competitor_domains": top_comp, "top_winnable_gap": gaps[0]["cluster_query"] if gaps else "",
        })
    with (data / "pages-summary.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=SUMMARY_FIELDS)
        w.writeheader()
        w.writerows(summary_rows)

    print(f"Wrote {len(cluster_rows)} cluster row(s) to {data / 'clusters.csv'} and "
          f"{len(summary_rows)} page row(s) to {data / 'pages-summary.csv'}.")
    return 0


# ----------------------------------------------------------------- fixes --

def fix_for(cluster: dict) -> str:
    if cluster["type"] == "site_vendor":
        return "SKIP_VENDOR"
    if cluster["type"] == "site_own":
        return ""
    if cluster["coverage"] in ("exact", "heading", "partial"):
        rank = cluster.get("rank") or ""
        if cluster["type"] == "open" and rank and rank not in ("error",) and \
                (rank == "none" or (rank.isdigit() and int(rank) > 20)):
            return "EARN_RETRIEVAL"
        return ""
    return "ADD_EXACT_STRING" if cluster["type"] == "exact_string" else "ADD_SECTION"


# ---------------------------------------------------------------- report --

def cmd_report(args) -> int:
    out = Path(args.out)
    data = out / DATA_DIR
    for name in ("clusters.csv", "pages-summary.csv", "pages.json"):
        if not (data / name).is_file():
            sys.exit(f"missing {data / name} - run `analyse` first")

    with (data / "clusters.csv").open(encoding="utf-8-sig") as fh:
        clusters = list(csv.DictReader(fh))
    with (data / "pages-summary.csv").open(encoding="utf-8-sig") as fh:
        pages_summary = list(csv.DictReader(fh))
    pages_index = json.loads((data / "pages.json").read_text(encoding="utf-8"))
    prompts = []
    if (data / "prompts.csv").is_file():
        with (data / "prompts.csv").open(encoding="utf-8-sig") as fh:
            prompts = list(csv.DictReader(fh))
    results = ([json.loads(l) for l in (data / "results.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
               if (data / "results.jsonl").is_file() else [])

    for c in clusters:
        c["fix"] = fix_for(c)
    stable = [c for c in clusters if float(c["stability"] or 0) >= 0.5]
    winnable = [c for c in stable if c["fix"] and c["fix"] != "SKIP_VENDOR"]
    winnable.sort(key=lambda c: -(float(c["stability"]) * (1.0 if c["coverage"] == "missing" else 0.5)))

    spend = round(sum(r.get("cost") or 0 for r in results), 3)
    runs_seen = sorted({r["run"] for r in results}) if results else []

    L = ["# Fan-out report", "",
         f"Pages: {len(pages_index['pages'])} | Prompts: {len(prompts)} | Runs per prompt: "
         f"up to {max(runs_seen) if runs_seen else '-'} | DataForSEO spend so far: ${spend:.3f}", "",
         "**How to read this.** Query fan-out coverage is a hypothesis for citation lift, not a "
         "proven cause: covering a stable fan-out query on the page may raise the odds an engine's "
         "retrieval step finds it, but nothing in this pack has run a controlled test to prove that "
         "lift. Fan-out varies run to run, and every number below comes from a small number of live "
         "runs, so treat it as a snapshot, not a rate.", "",
         "## Citation baseline", "",
         "| Page | Word count | Page cited (share of runs) | Domain cited | Top competitor domains |",
         "|---|---|---|---|---|"]
    for p in pages_summary:
        L.append(f"| {p['page_url']} | {p['word_count']} | {p['page_citation_rate']} | "
                 f"{p['domain_citation_rate']} | {p['top_competitor_domains'] or '-'} |")

    L += ["", "## Top winnable gaps", "",
          "Ranked by stability (share of runs the fan-out showed up in), weighted toward gaps the "
          "page does not cover yet. Excludes `site:` fan-outs aimed at another domain - those cannot "
          "be won by this page.", "",
          "| Page | Fan-out query | Stability | Type | Coverage | Fix |", "|---|---|---|---|---|---|"]
    for c in winnable[: args.top]:
        L.append(f"| {c['page_url']} | {c['cluster_query']} | {c['stability']} | {c['type']} | "
                 f"{c['coverage']} | {c['fix']} |")
    if not winnable:
        L.append("| - | no stable, winnable gaps found in this run | - | - | - | - |")

    L += ["", "## Method and limits", "",
          "- Fan-out queries come from DataForSEO LLM Responses `fan_out_queries`, asked live with web "
          "search on, `--runs` times per prompt. Stability is the share of runs a query's cluster "
          "showed up in, not a guaranteed future rate.",
          "- `site:` fan-outs aimed at a vendor's own domain cannot be won by a third-party page and "
          "are excluded from the winnable list.",
          "- Coverage is a heuristic token match against the page's own headings and body text, not a "
          "citation guarantee, and clustering near-duplicate fan-outs is an approximation "
          "(token Jaccard >= 0.6).",
          "- This report ranks where the page's own content is the most likely blocker today. It does "
          "not prove that closing a gap raises citations - re-run `retest` after the fix ships to "
          "measure the actual change.",
          "- See `references/evidence.md` for what was verified live and when.", ""]
    (out / "fan-out-report.md").write_text("\n".join(L) + "\n", encoding="utf-8")

    pages_dir = out / "pages"
    pages_dir.mkdir(parents=True, exist_ok=True)
    by_page = defaultdict(list)
    for c in stable:
        by_page[c["page_url"]].append(c)
    prompts_by_page = defaultdict(list)
    for p in prompts:
        prompts_by_page[p["page_url"]].append(p)

    for page in pages_index["pages"]:
        url = page["url"]
        pl = [f"# {url}", "", "## Prompts asked", ""]
        pl += [f"- `{p['prompt_id']}` ({p['source']}): {p['prompt']}" for p in prompts_by_page.get(url, [])] or ["(none)"]
        pl += ["", "## Stable fan-outs (>= 50% of runs)", "",
               "| Fan-out query | Stability | Type | Coverage | Rank | Fix |", "|---|---|---|---|---|---|"]
        page_clusters = by_page.get(url, [])
        pl += [f"| {c['cluster_query']} | {c['stability']} | {c['type']} | {c['coverage']} | "
              f"{c['rank'] or '-'} | {c['fix'] or '-'} |" for c in page_clusters] or ["| (none) | - | - | - | - | - |"]
        (pages_dir / f"{page['slug']}.md").write_text("\n".join(pl) + "\n", encoding="utf-8")

    print(f"Wrote {out / 'fan-out-report.md'} and {len(pages_index['pages'])} page report(s) in {pages_dir}/")
    return 0


# -------------------------------------------------------------- workbook --

WORKBOOK = "brand-audit-master.xlsx"
PACK_TEMPLATE = SKILL_DIR.parent / "_shared" / "brand-audit" / "brand-audit-template.xlsx"
TICK_SCRIPT = SKILL_DIR.parent / "_shared" / "brand-audit" / "tick_checklist.py"
ICE_TABLE = {"ADD_EXACT_STRING": (8, 8, 2), "ADD_SECTION": (7, 6, 4), "EARN_RETRIEVAL": (6, 4, 7)}


def ice_for(fix: str, stability: float) -> tuple[int, int, int]:
    impact, base_conf, ease = ICE_TABLE.get(fix, (5, 5, 5))
    confidence = max(1, min(10, round(base_conf * max(stability, 0.1) / 0.6)))
    return impact, confidence, ease


def build_fanout_tab(wb, stable: list[dict], styles: dict):
    if "Fan-Out" in wb.sheetnames:
        del wb["Fan-Out"]
    ws = wb.create_sheet("Fan-Out")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    headers = ["Page", "Prompt", "Fan-out query", "Stability", "Type", "Coverage", "Rank", "Fix"]
    for i, h in enumerate(headers, start=2):
        c = ws.cell(row=2, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = styles["head"], styles["head_fill"], styles["wrap"], styles["edge"]
    for r, row in enumerate(stable, start=3):
        values = [row["page_url"], row["prompt"], row["cluster_query"], row["stability"], row["type"],
                  row["coverage"], row.get("rank") or "-", row["fix"] or "-"]
        for i, v in enumerate(values, start=2):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment, c.border = styles["wrap"], styles["edge"]
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 30
    return ws


def add_initiatives(wb, stable: list[dict]) -> int:
    if "Initiatives" not in wb.sheetnames:
        return 0
    ws = wb["Initiatives"]
    start_row = next((r for r in range(5, ws.max_row + 2) if not ws.cell(row=r, column=4).value), 6)
    n = 0
    for row in stable:
        if row["fix"] in ("", "SKIP_VENDOR"):
            continue
        impact, confidence, ease = ice_for(row["fix"], float(row["stability"] or 0))
        r = start_row + n
        ws.cell(row=r, column=2, value=f'=IFERROR(ROUND(J{r}*K{r}/L{r},1),"")')
        ws.cell(row=r, column=3, value="Content Strategy")
        ws.cell(row=r, column=4, value=f"Fix fan-out gap on {row['page_url']}")
        ws.cell(row=r, column=5, value=f"{FIX_LABELS.get(row['fix'], '')} Fan-out: {row['cluster_query']}")
        ws.cell(row=r, column=6, value="mos-geo-query-fan-out")
        ws.cell(row=r, column=7, value="Scheduled")
        ws.cell(row=r, column=9, value="[AGENCY]")
        ws.cell(row=r, column=10, value=impact)
        ws.cell(row=r, column=11, value=confidence)
        ws.cell(row=r, column=12, value=ease)
        n += 1
    return n


def cmd_workbook(args) -> int:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        sys.exit("openpyxl is needed: uv run --with openpyxl python fanout.py workbook ...")

    out = Path(args.out)
    data = out / DATA_DIR
    if not (data / "clusters.csv").is_file():
        sys.exit(f"missing {data / 'clusters.csv'} - run `analyse` first")
    with (data / "clusters.csv").open(encoding="utf-8-sig") as fh:
        clusters = list(csv.DictReader(fh))
    for c in clusters:
        c["fix"] = fix_for(c)
    stable = [c for c in clusters if float(c["stability"] or 0) >= 0.5]

    book = out.resolve().parent / WORKBOOK
    if not book.is_file():
        if not PACK_TEMPLATE.is_file():
            sys.exit(f"no workbook in the run folder and no pack template at {PACK_TEMPLATE}")
        book.write_bytes(PACK_TEMPLATE.read_bytes())
    wb = openpyxl.load_workbook(book)

    styles = {"wrap": Alignment(wrap_text=True, vertical="top"), "head": Font(bold=True, color="FFFFFF"),
              "head_fill": PatternFill("solid", fgColor="1F2937"),
              "edge": Border(*(Side(style="thin", color="D1D5DB"),) * 4)}
    build_fanout_tab(wb, stable, styles)
    n_init = add_initiatives(wb, stable)

    order = ["Checklist", "Brand Truth Review", "Brand 360 Report", "AI Visibility", "AI Info Page",
             "Fan-Out", "Initiatives"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in order]
    wb.save(book)
    print(f"Wrote {book}: Fan-Out tab ({len(stable)} row(s)), {n_init} Initiatives row(s) added.")

    note = f"{len(stable)} stable fan-out cluster(s); see the 'Fan-Out' tab. Files in {out.resolve().name}/"
    try:
        subprocess.run([sys.executable, str(TICK_SCRIPT), "--skill", "mos-geo-query-fan-out",
                       "--run-dir", str(out.resolve()), "--status", "Client review", "--note", note],
                      check=True, capture_output=True, text=True)
    except Exception as e:  # noqa: BLE001 - the tabs above are already saved either way
        print(f"(Checklist tick skipped: {e})")
    return 0


# ---------------------------------------------------------------- retest --

def cmd_retest(args) -> int:
    baseline = Path(args.baseline)
    src_prompts = baseline / DATA_DIR / "prompts.csv"
    if not src_prompts.is_file():
        sys.exit(f"missing {src_prompts}")
    out = Path(args.out)
    data = out / DATA_DIR
    data.mkdir(parents=True, exist_ok=True)
    dst_prompts = data / "prompts.csv"
    if not dst_prompts.is_file():
        dst_prompts.write_text(src_prompts.read_text(encoding="utf-8"), encoding="utf-8")

    rc = cmd_run(Namespace(env_file=args.env_file, config=args.config, prompts=str(dst_prompts),
                          out=str(out), country=args.country, only=args.only,
                          model=getattr(args, "model", None), workers=args.workers,
                          runs=args.runs, max_spend=args.max_spend))

    def citation_rates(run_dir: Path) -> dict[str, tuple[int, int]]:
        results = run_dir / DATA_DIR / "results.jsonl"
        if not results.is_file():
            return {}
        counts: dict[str, list[int]] = defaultdict(lambda: [0, 0])
        for line in results.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            r = json.loads(line)
            if r.get("error"):
                continue
            counts[r["page_url"]][1] += 1
            if r.get("page_cited"):
                counts[r["page_url"]][0] += 1
        return {u: (c[0], c[1]) for u, c in counts.items()}

    base_rates, new_rates = citation_rates(baseline), citation_rates(out)
    rows = []
    for url in sorted(set(base_rates) | set(new_rates)):
        bc, bn = base_rates.get(url, (0, 0))
        nc, nn = new_rates.get(url, (0, 0))
        b_rate = round(bc / bn, 2) if bn else 0.0
        n_rate = round(nc / nn, 2) if nn else 0.0
        rows.append({"page_url": url, "baseline_citation_rate": b_rate,
                    "new_citation_rate": n_rate, "delta": round(n_rate - b_rate, 2)})
    with (data / "retest-diff.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["page_url", "baseline_citation_rate", "new_citation_rate", "delta"])
        w.writeheader()
        w.writerows(rows)
    print(f"Wrote {data / 'retest-diff.csv'} ({len(rows)} page(s)).")
    for row in rows:
        sign = "+" if row["delta"] >= 0 else ""
        print(f"  {row['page_url']}: {row['baseline_citation_rate']} -> {row['new_citation_rate']} "
              f"({sign}{row['delta']})")
    return rc


# ------------------------------------------------------------------ main --

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    def common(p):
        p.add_argument("--config", help="engines config (default: config/engines.json)")
        p.add_argument("--env-file", help="a .env file holding DATAFORSEO_LOGIN / DATAFORSEO_PASSWORD")

    p = sub.add_parser("path", help="print the run folder for this run (MarketingOS-aware)")
    p.add_argument("--site", help="site/brand name, for the agency-mode folder slug")
    p.add_argument("--url", help="derive the slug from this URL's domain if --site is not given")
    p.add_argument("--date", help="YYYY-MM-DD (default today)")
    p.add_argument("--start", help="folder to resolve from (default: current folder)")
    p.set_defaults(fn=cmd_path)

    p = sub.add_parser("preflight", help="check credentials, model names, resolve --country, estimate cost")
    common(p)
    p.add_argument("--country", default="AU", help="2-letter market (default AU)")
    p.add_argument("--pages", type=int, default=1, help="planned page count, for the cost estimate")
    p.add_argument("--prompts", type=int, default=6, help="planned prompts per page")
    p.add_argument("--runs", type=int, default=3, help="planned runs per prompt")
    p.add_argument("--engines", default="chat_gpt", help="comma list of engine ids in the planned run")
    p.add_argument("--max-spend", type=float, default=25.0)
    p.set_defaults(fn=cmd_preflight)

    p = sub.add_parser("pages", help="fetch and parse the page(s) to test")
    p.add_argument("--url", help="a single page URL")
    p.add_argument("--urls", help="a text file, one URL per line")
    p.add_argument("--sitemap", help="a sitemap.xml URL")
    p.add_argument("--include", help="regex the sitemap URLs must match")
    p.add_argument("--limit", type=int, help="cap the number of URLs")
    p.add_argument("--out", required=True, help="run folder")
    p.add_argument("--workers", type=int, default=10)
    p.add_argument("--refresh", action="store_true", help="refetch pages already cached")
    p.set_defaults(fn=cmd_pages)

    p = sub.add_parser("prompts", help="draft 5-8 buyer prompts per page into data/prompts.csv")
    common(p)
    p.add_argument("--out", required=True, help="run folder")
    p.add_argument("--country", default="AU")
    p.add_argument("--keyword", help="primary keyword, for a single-page run")
    p.add_argument("--keywords-csv", help="CSV with url,keyword columns, for a multi-page run")
    p.add_argument("--discover", action="store_true",
                   help="also add up to 3 observed prompts per page via LLM Mentions")
    p.add_argument("--max-generated", type=int, default=8)
    p.add_argument("--yes", action="store_true", help="spend the estimated cost and write prompts.csv")
    p.set_defaults(fn=cmd_prompts)

    p = sub.add_parser("run", help="run every prompt through the engines, --runs times each, live")
    common(p)
    p.add_argument("--prompts", required=True, help="data/prompts.csv (see `prompts`)")
    p.add_argument("--out", required=True, help="run folder")
    p.add_argument("--country", default="AU")
    p.add_argument("--runs", type=int, default=3)
    p.add_argument("--only", help="comma list of engine ids (default: every enabled engine)")
    p.add_argument("--model", help="override the model for every enabled engine (default: each engine's "
                   "configured model - gpt-5.6-luna for chat_gpt; pass e.g. gpt-5.6-terra to compare)")
    p.add_argument("--workers", type=int, default=10)
    p.add_argument("--max-spend", type=float, default=25.0)
    p.set_defaults(fn=cmd_run)

    p = sub.add_parser("analyse", help="cluster fan-outs, classify them, score coverage and citations")
    common(p)
    p.add_argument("--out", required=True, help="run folder")
    p.add_argument("--country", default="AU")
    p.add_argument("--serp", action="store_true",
                   help="also check Google organic rank (depth 20) for open/exact_string clusters")
    p.set_defaults(fn=cmd_analyse)

    p = sub.add_parser("report", help="write fan-out-report.md and pages/<slug>.md")
    p.add_argument("--out", required=True, help="run folder")
    p.add_argument("--top", type=int, default=20, help="max rows in the top-winnable-gaps table")
    p.set_defaults(fn=cmd_report)

    p = sub.add_parser("workbook", help="add the Fan-Out tab + Initiatives and tick Checklist (needs openpyxl)")
    p.add_argument("--out", required=True, help="run folder")
    p.set_defaults(fn=cmd_workbook)

    p = sub.add_parser("retest", help="re-run an earlier run's exact prompts.csv and diff citation rate")
    common(p)
    p.add_argument("--baseline", required=True, help="an earlier run folder")
    p.add_argument("--out", required=True, help="the new run folder")
    p.add_argument("--country", default="AU")
    p.add_argument("--runs", type=int, default=3)
    p.add_argument("--only", help="comma list of engine ids")
    p.add_argument("--model", help="override the model for every enabled engine (see `run --model`)")
    p.add_argument("--workers", type=int, default=10)
    p.add_argument("--max-spend", type=float, default=25.0)
    p.set_defaults(fn=cmd_retest)

    args = ap.parse_args()
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(main())
